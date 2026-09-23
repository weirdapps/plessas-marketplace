"""Tests for inject_chart_data.py and inject_table_data.py (BUILD-CODE-5, ASSETS-CI-10).

Both tools update an existing deck in place of rebuilding it. Each defect below
produced a deck that looked updated and was not: charts reverted to the template's
numbers the first time someone clicked Edit Data, template series stayed plotted,
table cells mixed old and new text, and a config entry that matched nothing still
exited 0.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import inject_chart_data
import inject_table_data
import nbg_build
import pytest
from pptx import Presentation
from pptx.util import Inches
from testkit import deck, write_spec

HERE = Path(__file__).parent
SOURCE = {"name": "Management accounts", "as_of": "30 June 2026"}


@pytest.fixture
def chart_deck(tmp_path: Path) -> Path:
    """Slide 1 (0-based) holds a two-series bar chart, slide 2 a line chart."""
    spec = deck(
        [
            {
                "type": "chart",
                "content": {"title": "Volumes rose", "source": SOURCE},
                "chart": {
                    "type": "bar",
                    "data": {
                        "categories": ["Q1", "Q2"],
                        "series": [
                            {"name": "2025", "values": [1, 2]},
                            {"name": "2026", "values": [3, 4]},
                        ],
                    },
                },
            },
            {
                "type": "chart",
                "content": {"title": "Balances rose", "source": SOURCE},
                "chart": {
                    "type": "line",
                    "data": {
                        "categories": ["Jan", "Feb"],
                        "series": [{"name": "2026", "values": [5, 6]}],
                    },
                },
            },
        ]
    )
    out = tmp_path / "charts.pptx"
    nbg_build.build_presentation(write_spec(tmp_path, spec), out, validate=False)
    return out


def _config(tmp_path: Path, payload: dict) -> Path:
    path = tmp_path / "config.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _chart(pptx: Path, slide: int):
    return next(s for s in Presentation(str(pptx)).slides[slide].shapes if s.has_chart).chart


def _workbook_values(pptx: Path, slide: int) -> str:
    chart = _chart(pptx, slide)
    xlsx = chart.part.chart_workbook.xlsx_part.blob
    import io

    with zipfile.ZipFile(io.BytesIO(xlsx)) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        strings = (
            zf.read("xl/sharedStrings.xml").decode("utf-8")
            if "xl/sharedStrings.xml" in zf.namelist()
            else ""
        )
    return sheet + strings


# ---------------------------------------------------------------- charts


def test_chart_injection_rewrites_the_embedded_workbook(chart_deck, tmp_path):
    """Charts reverted to template numbers on Edit Data: only the cache was edited."""
    config = _config(
        tmp_path,
        {
            "charts": [
                {
                    "slide": 1,
                    "chart_index": 0,
                    "type": "bar",
                    "data": {
                        "categories": ["Q1", "Q2"],
                        "series": [
                            {"name": "2025", "values": [11, 22]},
                            {"name": "2026", "values": [33, 44]},
                        ],
                    },
                }
            ]
        },
    )
    out = tmp_path / "out.pptx"
    inject_chart_data.inject_chart_data(str(chart_deck), str(config), str(out))
    plot = _chart(out, 1).plots[0]
    assert [list(s.values) for s in plot.series] == [[11.0, 22.0], [33.0, 44.0]]
    workbook = _workbook_values(out, 1)
    # Numeric cells only: a shared-string cell (t="s") holds an index, not a value.
    numbers = {
        float(v) for v in re.findall(r'<c r="[A-Z]+\d+"(?: s="\d+")?><v>([^<]+)</v>', workbook)
    }
    assert numbers == {11.0, 22.0, 33.0, 44.0}, (
        "the embedded workbook still holds the template data"
    )


def test_chart_injection_adds_and_removes_series(chart_deck, tmp_path):
    three = {
        "categories": ["Q1", "Q2"],
        "series": [{"name": n, "values": [1, 2]} for n in ("a", "b", "c")],
    }
    out3 = tmp_path / "three.pptx"
    inject_chart_data.inject_chart_data(
        str(chart_deck),
        str(_config(tmp_path, {"charts": [{"slide": 1, "data": three}]})),
        str(out3),
    )
    assert [s.name for s in _chart(out3, 1).plots[0].series] == ["a", "b", "c"]

    one = {"categories": ["Q1", "Q2", "Q3"], "series": [{"name": "only", "values": [1, 2, 3]}]}
    out1 = tmp_path / "one.pptx"
    inject_chart_data.inject_chart_data(
        str(chart_deck), str(_config(tmp_path, {"charts": [{"slide": 1, "data": one}]})), str(out1)
    )
    plot = _chart(out1, 1).plots[0]
    assert [s.name for s in plot.series] == ["only"], "a template series is still plotted"
    assert list(plot.categories) == ["Q1", "Q2", "Q3"]


def test_a_line_chart_round_trips_and_keeps_its_brand_stroke(chart_deck, tmp_path):
    data = {
        "categories": ["Jan", "Feb", "Mar"],
        "series": [{"name": "2026", "values": [7, 8, 9]}, {"name": "2027", "values": [1, 2, 3]}],
    }
    out = tmp_path / "line.pptx"
    inject_chart_data.inject_chart_data(
        str(chart_deck), str(_config(tmp_path, {"charts": [{"slide": 2, "data": data}]})), str(out)
    )
    plot = _chart(out, 2).plots[0]
    assert len(plot.series) == 2 and list(plot.categories) == ["Jan", "Feb", "Mar"]
    xml = zipfile.ZipFile(out).read(_chart(out, 2).part.partname.lstrip("/")).decode("utf-8")
    assert (
        xml.count('<a:srgbClr val="00ADBF"/>') >= 1 and xml.count('<a:srgbClr val="003841"/>') >= 1
    )


def test_the_doughnut_shorthand_still_works(chart_deck, tmp_path):
    out = tmp_path / "labels.pptx"
    config = {
        "charts": [
            {"slide": 1, "type": "doughnut", "data": {"labels": ["A", "B"], "values": [60, 40]}}
        ]
    }
    inject_chart_data.inject_chart_data(str(chart_deck), str(_config(tmp_path, config)), str(out))
    plot = _chart(out, 1).plots[0]
    assert list(plot.categories) == ["A", "B"] and [list(s.values) for s in plot.series] == [
        [60.0, 40.0]
    ]


@pytest.mark.parametrize(
    "entry,problem",
    [
        (
            {"slide": 99, "data": {"categories": ["a"], "series": [{"name": "s", "values": [1]}]}},
            "slide 99",
        ),
        (
            {"slide": 0, "data": {"categories": ["a"], "series": [{"name": "s", "values": [1]}]}},
            "no chart",
        ),
        (
            {
                "slide": 1,
                "chart_index": 3,
                "data": {"categories": ["a"], "series": [{"name": "s", "values": [1]}]},
            },
            "chart_index 3",
        ),
        (
            {
                "slide": 1,
                "data": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]},
            },
            "1 value",
        ),
    ],
)
def test_a_chart_entry_that_matches_nothing_fails_and_writes_nothing(
    chart_deck, tmp_path, entry, problem
):
    out = tmp_path / "out.pptx"
    with pytest.raises(inject_chart_data.InjectionError, match=problem):
        inject_chart_data.inject_chart_data(
            str(chart_deck), str(_config(tmp_path, {"charts": [entry]})), str(out)
        )
    assert not out.exists()


def test_the_chart_cli_exits_1_on_an_unmatched_entry(chart_deck, tmp_path):
    config = _config(
        tmp_path,
        {
            "charts": [
                {
                    "slide": 99,
                    "data": {"categories": ["a"], "series": [{"name": "s", "values": [1]}]},
                }
            ]
        },
    )
    proc = subprocess.run(
        [
            sys.executable,
            str(HERE / "inject_chart_data.py"),
            str(chart_deck),
            str(config),
            str(tmp_path / "o.pptx"),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert proc.returncode == 1, proc.stdout + proc.stderr
    assert "slide 99" in proc.stderr


# ---------------------------------------------------------------- tables


@pytest.fixture
def table_deck(tmp_path: Path) -> Path:
    """Slide 0 holds a 3x3 table whose cells carry two runs each, as a PowerPoint
    template does after anyone has edited it."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(1.5)).table
    for r in range(3):
        for c in range(3):
            p = table.cell(r, c).text_frame.paragraphs[0]
            p.add_run().text = "Old"
            p.add_run().text = " template text"
    out = tmp_path / "table.pptx"
    prs.save(str(out))
    return out


def _cells(pptx: Path) -> list[list[str]]:
    table = next(s for s in Presentation(str(pptx)).slides[0].shapes if s.has_table).table
    return [[cell.text for cell in row.cells] for row in table.rows]


def test_table_injection_replaces_every_run(table_deck, tmp_path):
    data = {
        "headers": ["Metric", "2025", "2026"],
        "rows": [["Assets", "1", "2"], ["Deposits", "3", "4"]],
    }
    out = tmp_path / "out.pptx"
    inject_table_data.inject_table_data(
        str(table_deck), str(_config(tmp_path, {"tables": [{"slide": 0, "data": data}]})), str(out)
    )
    assert _cells(out) == [["Metric", "2025", "2026"], ["Assets", "1", "2"], ["Deposits", "3", "4"]]


def test_table_injection_applies_the_highlight_column(table_deck, tmp_path):
    data = {
        "headers": ["Metric", "2025", "2026"],
        "rows": [["Assets", "1", "2"], ["Deposits", "3", "4"]],
    }
    out = tmp_path / "out.pptx"
    config = {"tables": [{"slide": 0, "data": data, "highlight_column": 2}]}
    inject_table_data.inject_table_data(str(table_deck), str(_config(tmp_path, config)), str(out))
    table = next(s for s in Presentation(str(out)).slides[0].shapes if s.has_table).table
    run = table.cell(1, 2).text_frame.paragraphs[0].runs[0]
    assert run.font.bold is True and str(run.font.color.rgb) == "007B85"
    other = table.cell(1, 1).text_frame.paragraphs[0].runs[0]
    assert other.font.bold is not True


def test_table_injection_refuses_data_of_a_different_shape(table_deck, tmp_path):
    data = {"headers": ["A", "B"], "rows": [["1", "2"], ["3", "4"], ["5", "6"]]}
    out = tmp_path / "out.pptx"
    with pytest.raises(inject_table_data.InjectionError, match="4 x 2"):
        inject_table_data.inject_table_data(
            str(table_deck),
            str(_config(tmp_path, {"tables": [{"slide": 0, "data": data}]})),
            str(out),
        )
    assert not out.exists()


def test_the_table_cli_exits_1_on_an_unmatched_entry(table_deck, tmp_path):
    config = _config(tmp_path, {"tables": [{"slide": 5, "data": {"rows": [["x"]]}}]})
    proc = subprocess.run(
        [
            sys.executable,
            str(HERE / "inject_table_data.py"),
            str(table_deck),
            str(config),
            str(tmp_path / "o.pptx"),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert proc.returncode == 1 and "slide 5" in proc.stderr


# ---------------------------------------------------------------- namespaces

REAL_SLIDE_ROOT = (
    '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
    'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
    'mc:Ignorable="p14 a14">'
)


def test_table_injection_preserves_foreign_namespace_prefixes(tmp_path):
    """ElementTree renamed mc to ns1 and dropped p14, leaving mc:Ignorable dangling."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    plain = tmp_path / "plain.pptx"
    prs.save(str(plain))
    with zipfile.ZipFile(plain) as zf:
        blobs = {n: zf.read(n) for n in zf.namelist()}
    xml = blobs["ppt/slides/slide1.xml"].decode("utf-8")
    start = xml.index("<p:sld ")
    blobs["ppt/slides/slide1.xml"] = (
        xml[:start] + REAL_SLIDE_ROOT + xml[xml.index(">", start) + 1 :]
    ).encode("utf-8")
    real = tmp_path / "real.pptx"
    with zipfile.ZipFile(real, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in blobs.items():
            zf.writestr(name, blob)

    config = {
        "tables": [
            {"slide": 0, "data": {"headers": ["Metric", "Value"], "rows": [["Assets", "78.5"]]}}
        ]
    }
    out = tmp_path / "out.pptx"
    inject_table_data.inject_table_data(str(real), str(_config(tmp_path, config)), str(out))
    after = zipfile.ZipFile(out).read("ppt/slides/slide1.xml").decode("utf-8")
    for prefix in ("a", "r", "p", "mc", "p14", "a14"):
        assert f'xmlns:{prefix}="' in after, f"xmlns:{prefix} was dropped"
    assert 'mc:Ignorable="p14 a14"' in after
    assert not re.search(r"\bns\d+:", after), "a namespace prefix was renamed"
    assert ">Assets<" in after and ">78.5<" in after
