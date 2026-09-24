#!/usr/bin/env python3
"""Replace the data behind charts in an existing deck.

    inject_chart_data.py input.pptx chart_config.json output.pptx

chart_config.json:

    {
      "charts": [
        {
          "slide": 3,
          "chart_index": 0,
          "data": {
            "categories": ["Q1", "Q2"],
            "series": [{"name": "2025", "values": [12, 15]}]
          }
        },
        {"slide": 5, "data": {"labels": ["A", "B"], "values": [60, 40]}}
      ]
    }

"slide" counts from 0 in presentation order; "chart_index" counts from 0 in the
slide's shape order and defaults to 0. The second form, labels plus values, is a
single unnamed series (a doughnut's usual shape). An optional "title" replaces the
chart title when the chart has one.

Each chart's data is replaced through python-pptx's chart.replace_data, which
rewrites the cached values, the cell references and the embedded workbook, and adds
or removes series to match. The old tool edited the cached values only: the chart
reverted to its template numbers the moment anyone clicked Edit Data, and a
template series with no counterpart in the config stayed on the plot. Every series
is then coloured from the brand palette (tokens.yaml charts.palette): a line
through its stroke, anything else through its fill.

Nothing is written unless every entry matched a chart and its data fit: a config
that silently did nothing used to exit 0. Exit codes: 0 written; 1 an entry did
not match or its data was malformed (each named on stderr); 2 a file could not be
read or written.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import nbg_tokens  # noqa: E402


class InjectionError(ValueError):
    """At least one config entry did not match the deck; nothing was written."""


def _series_of(entry: dict[str, Any], where: str) -> tuple[list[str], list[dict[str, Any]]]:
    data = entry.get("data")
    if not isinstance(data, dict):
        raise InjectionError(f"{where}: no data mapping")
    categories = data.get("categories") or data.get("labels")
    series = data.get("series")
    if series is None and "values" in data:
        series = [{"name": str(entry.get("title") or ""), "values": data["values"]}]
    if not isinstance(categories, list) or not categories:
        raise InjectionError(f"{where}: data.categories (or data.labels) is missing or empty")
    if not isinstance(series, list) or not series:
        raise InjectionError(f"{where}: data.series (or data.values) is missing or empty")
    for s, item in enumerate(series):
        values = item.get("values") if isinstance(item, dict) else None
        if not isinstance(values, list) or len(values) != len(categories):
            count = len(values) if isinstance(values, list) else 0
            raise InjectionError(
                f"{where}: series {s} has {count} value(s) for {len(categories)} categories"
            )
    return [str(c) for c in categories], series


def _recolour(chart: Any) -> None:
    palette = nbg_tokens.chart_palette()
    for plot in chart.plots:
        is_line = plot._element.tag.endswith("lineChart")
        for i, series in enumerate(plot.series):
            colour = RGBColor.from_string(palette[i % len(palette)])
            if is_line:
                series.format.line.color.rgb = colour
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = colour


def inject_chart_data(pptx_path: str, config_path: str, output_path: str) -> int:
    """Apply every entry, then save. Returns the number of charts updated."""
    prs = Presentation(str(Path(pptx_path).expanduser()))
    with open(Path(config_path).expanduser(), encoding="utf-8") as f:
        config = json.load(f)
    problems: list[str] = []
    updated = 0
    for n, entry in enumerate(config.get("charts", [])):
        where = f"charts[{n}]"
        slide_index = entry.get("slide")
        if not isinstance(slide_index, int) or not 0 <= slide_index < len(prs.slides):
            problems.append(
                f"{where}: slide {slide_index} does not exist (the deck has {len(prs.slides)} "
                "slides, numbered from 0)"
            )
            continue
        charts = [s for s in prs.slides[slide_index].shapes if getattr(s, "has_chart", False)]
        if not charts:
            problems.append(f"{where}: slide {slide_index} has no chart")
            continue
        chart_index = entry.get("chart_index", 0)
        if not isinstance(chart_index, int) or not 0 <= chart_index < len(charts):
            problems.append(
                f"{where}: chart_index {chart_index} does not exist (slide {slide_index} has "
                f"{len(charts)} chart(s), numbered from 0)"
            )
            continue
        try:
            categories, series = _series_of(entry, where)
        except InjectionError as e:
            problems.append(str(e))
            continue
        chart = charts[chart_index].chart
        data = CategoryChartData()
        data.categories = categories
        for s in series:
            data.add_series(str(s.get("name", "")), list(s["values"]))
        chart.replace_data(data)
        _recolour(chart)
        if entry.get("title") and chart.has_title:
            chart.chart_title.text_frame.text = str(entry["title"])
        updated += 1
    if problems:
        raise InjectionError("\n".join(problems))
    prs.save(str(Path(output_path).expanduser()))
    return updated


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) != 3:
        print(
            "usage: inject_chart_data.py <input.pptx> <config.json> <output.pptx>", file=sys.stderr
        )
        return 2
    try:
        updated = inject_chart_data(*args)
    except InjectionError as e:
        print(f"inject_chart_data.py: nothing written.\n{e}", file=sys.stderr)
        return 1
    except (OSError, ValueError, KeyError) as e:
        print(f"inject_chart_data.py: {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    print(f"Updated {updated} chart(s): {args[2]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
