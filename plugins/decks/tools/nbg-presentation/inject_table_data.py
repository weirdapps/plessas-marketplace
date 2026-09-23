#!/usr/bin/env python3
"""Replace the text in tables of an existing deck.

    inject_table_data.py input.pptx table_config.json output.pptx

table_config.json:

    {
      "tables": [
        {
          "slide": 2,
          "table_index": 0,
          "data": {
            "headers": ["Metric", "2025", "2026"],
            "rows": [["Assets (EUR B)", "78.5", "82.1"], ["CET1 ratio", "17.2%", "15.8%"]]
          },
          "highlight_column": 2
        }
      ]
    }

"slide" counts from 0 in presentation order; "table_index" counts from 0 in the
slide's shape order and defaults to 0. "headers" is optional and, when given, fills
the first row. "highlight_column" (0-based) sets that column's body cells in bold
NBG Teal, as the builder does.

The data must be exactly the table's shape, rows and columns: a template with more
rows than the data used to keep its old rows, and one with fewer silently dropped
data. Each cell keeps the character formatting of its first run and loses every
other run, so no cell mixes old and new text. Nothing is written unless every entry
matched. Exit codes: 0 written; 1 an entry did not match (each named on stderr);
2 a file could not be read or written.
"""

from __future__ import annotations

import copy
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import nbg_tokens  # noqa: E402


class InjectionError(ValueError):
    """At least one config entry did not match the deck; nothing was written."""


def set_cell_text(cell: Any, text: str, *, highlight: bool = False) -> None:
    """Replace a cell's text with one run that keeps the first run's formatting."""
    txbody = cell.text_frame._txBody
    paragraphs = txbody.findall(qn("a:p"))
    first = paragraphs[0]
    for extra in paragraphs[1:]:
        txbody.remove(extra)
    runs = first.findall(qn("a:r"))
    template = runs[0].find(qn("a:rPr")) if runs else None
    for child in list(first):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            first.remove(child)
    paragraph = cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    if template is not None:
        old = run._r.find(qn("a:rPr"))
        if old is not None:
            run._r.remove(old)
        run._r.insert(0, copy.deepcopy(template))
    if highlight:
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(nbg_tokens.color("teal"))


def inject_table_data(pptx_path: str, config_path: str, output_path: str) -> int:
    """Apply every entry, then save. Returns the number of tables updated."""
    prs = Presentation(str(Path(pptx_path).expanduser()))
    with open(Path(config_path).expanduser(), encoding="utf-8") as f:
        config = json.load(f)
    problems: list[str] = []
    updated = 0
    for n, entry in enumerate(config.get("tables", [])):
        where = f"tables[{n}]"
        slide_index = entry.get("slide")
        if not isinstance(slide_index, int) or not 0 <= slide_index < len(prs.slides):
            problems.append(
                f"{where}: slide {slide_index} does not exist (the deck has {len(prs.slides)} "
                "slides, numbered from 0)"
            )
            continue
        tables = [s for s in prs.slides[slide_index].shapes if getattr(s, "has_table", False)]
        table_index = entry.get("table_index", 0)
        if not isinstance(table_index, int) or not 0 <= table_index < len(tables):
            problems.append(
                f"{where}: table_index {table_index} does not exist (slide {slide_index} has "
                f"{len(tables)} table(s), numbered from 0)"
            )
            continue
        data = entry.get("data") or {}
        headers = data.get("headers")
        rows = ([headers] if headers else []) + list(data.get("rows") or [])
        table = tables[table_index].table
        n_rows, n_cols = len(table.rows), len(table.columns)
        widths = {len(r) for r in rows if isinstance(r, list)}
        if len(rows) != n_rows or widths != {n_cols} or not all(isinstance(r, list) for r in rows):
            got_cols = max(widths) if widths else 0
            problems.append(
                f"{where}: the data is {len(rows)} x {got_cols}, the table on slide {slide_index} "
                f"is {n_rows} x {n_cols} (rows x columns, headers included)"
            )
            continue
        highlight = entry.get("highlight_column")
        body_start = 1 if headers else 0
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                set_cell_text(
                    table.cell(r, c),
                    "" if value is None else str(value),
                    highlight=highlight == c and r >= body_start,
                )
        updated += 1
    if problems:
        raise InjectionError("\n".join(problems))
    prs.save(str(Path(output_path).expanduser()))
    return updated


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) != 3:
        print(
            "usage: inject_table_data.py <input.pptx> <config.json> <output.pptx>", file=sys.stderr
        )
        return 2
    try:
        updated = inject_table_data(*args)
    except InjectionError as e:
        print(f"inject_table_data.py: nothing written.\n{e}", file=sys.stderr)
        return 1
    except (OSError, ValueError, KeyError) as e:
        print(f"inject_table_data.py: {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    print(f"Updated {updated} table(s): {args[2]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
