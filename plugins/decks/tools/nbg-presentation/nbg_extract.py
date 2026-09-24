#!/usr/bin/env python3
"""A deck's content as markdown, for /redesign-deck to rebuild from.

    decks-py extract deck.pptx      (or a .pdf)

For a .pptx: every slide's title, its text in reading order (bullets as "- ",
indented by level), its tables as markdown tables, its charts as a table of
categories by series, its pictures by their alt text, and its speaker notes. For
a .pdf: the text of each page. Replaces markitdown, which nothing installed.

A .docx is not supported: python-docx is not one of this tool's dependencies. Save
the document as PDF and extract that.

A .pptx passes nbg_package's limits before anything opens it: a third-party deck can
be a zip bomb.

Exit codes: 0 markdown on stdout; 2 the file could not be read.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import nbg_package  # noqa: E402

EXIT_OK = 0
EXIT_ERROR = 2

_CHART_NAMES = {
    "COLUMN_CLUSTERED": "bar",
    "COLUMN_STACKED": "bar_stacked",
    "COLUMN_STACKED_100": "bar_stacked",
    "BAR_CLUSTERED": "bar_horizontal",
    "BAR_STACKED": "bar_horizontal",
    "LINE": "line",
    "LINE_MARKERS": "line",
    "AREA": "area",
    "DOUGHNUT": "doughnut",
    "PIE": "pie",
}


class ExtractError(Exception):
    pass


def _cell(text: Any) -> str:
    return str(text if text is not None else "").replace("|", "\\|").replace("\n", " ").strip()


def _md_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(_cell(c) for c in rows[0]) + " |", "|" + " --- |" * width]
    out += ["| " + " | ".join(_cell(c) for c in r) + " |" for r in rows[1:]]
    return out


def _is_bullet(paragraph: Any) -> bool:
    ppr = paragraph._p.pPr
    if ppr is None:
        return False
    tags = {child.tag.split("}")[1] for child in ppr}
    return bool(tags & {"buChar", "buAutoNum", "buBlip"})


def _text_lines(shape: Any) -> list[str]:
    lines = []
    for p in shape.text_frame.paragraphs:
        text = "".join(r.text for r in p.runs) or p.text
        text = text.strip()
        if not text:
            continue
        if _is_bullet(p):
            lines.append("  " * p.level + f"- {text}")
        else:
            lines.append(text)
    return lines


def _waterfall_rows(chart: Any) -> list[list[str]] | None:
    """A bridge nbg_build drew: categories and the signed labels it wrote as text."""
    plot = chart.plots[0]
    names = [s.name for s in plot.series]
    if not names or names[0] != "Base" or "Labels" not in names:
        return None
    categories = list(plot.categories)
    labels = [""] * len(categories)
    for series in plot.series:
        for i in range(len(categories)):
            dlbl = series.points[i].data_label
            if dlbl.has_text_frame and dlbl.text_frame.text.strip():
                labels[i] = dlbl.text_frame.text.strip()
    return [["Step", "Value"]] + [[str(c), labels[i]] for i, c in enumerate(categories)]


def _chart_lines(chart: Any) -> list[str]:
    kinds = []
    for plot in chart.plots:
        tag = plot._element.tag.split("}")[1]
        kinds.append(tag)
    kind = (
        _CHART_NAMES.get(chart.chart_type.name, chart.chart_type.name.lower())
        if chart.chart_type
        else "chart"
    )
    if "areaChart" in kinds and "lineChart" in kinds:
        kind = "area_line"
    waterfall = _waterfall_rows(chart)
    if waterfall is not None:
        return ["Chart (waterfall):", "", *_md_table(waterfall)]
    plot = next((p for p in chart.plots if p._element.tag.endswith("lineChart")), chart.plots[0])
    categories = [str(c) for c in plot.categories]
    series = list(plot.series)
    rows = [["Category", *[s.name for s in series]]]
    for i, category in enumerate(categories):
        row = [category]
        for s in series:
            values = list(s.values)
            value = values[i] if i < len(values) else None
            row.append("" if value is None else f"{value:g}")
        rows.append(row)
    return [f"Chart ({kind}):", "", *_md_table(rows)]


_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_DIAGRAM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_DECORATIVE_EXT = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
# Charts python-pptx cannot read, by their plot element, for the cached-values reader.
_XML_CHART_NAMES = {
    "pie3DChart": "3-D pie",
    "bar3DChart": "3-D bar",
    "line3DChart": "3-D line",
    "area3DChart": "3-D area",
    "surface3DChart": "3-D surface",
    "surfaceChart": "surface",
    "stockChart": "stock",
    "ofPieChart": "pie of pie",
}


def _cached_chart_lines(chart: Any) -> list[str]:
    """A chart python-pptx cannot read (a 3-D, stock or surface chart), from the values
    its XML caches: each series' name, the categories and the values. A 3-D pie in a
    colleague's old deck crashed extract with a traceback (BUILDER-CODE-09)."""
    plot_area = chart._chartSpace.find(f"{_C}chart/{_C}plotArea")
    plots = [el for el in plot_area if el.tag.endswith("Chart")] if plot_area is not None else []
    if not plots:
        return ["[chart: data not read]"]
    tag = plots[0].tag.split("}")[1]
    kind = _XML_CHART_NAMES.get(tag, tag)
    series = plots[0].findall(f"{_C}ser")
    if not series:
        return [f"[chart: {kind}, data not read]"]

    def cached(el: Any) -> dict[int, str]:
        if el is None:
            return {}
        return {int(pt.get("idx", 0)): pt.findtext(f"{_C}v") or "" for pt in el.iter(f"{_C}pt")}

    def figure(text: str) -> str:
        try:
            return f"{float(text):g}"
        except ValueError:
            return text

    categories = cached(series[0].find(f"{_C}cat"))
    names, values = [], []
    for i, ser in enumerate(series):
        tx = ser.find(f"{_C}tx")
        names.append("".join(v.text or "" for v in tx.iter(f"{_C}v")) if tx is not None else "")
        names[-1] = names[-1] or f"Series {i + 1}"
        val = ser.find(f"{_C}val")
        values.append(cached(val if val is not None else ser.find(f"{_C}yVal")))
    count = max([len(categories)] + [max(v, default=-1) + 1 for v in values])
    rows = [["Category", *names]]
    rows += [
        [categories.get(i, str(i + 1)), *[figure(v.get(i, "")) for v in values]]
        for i in range(count)
    ]
    return [f"Chart ({kind}):", "", *_md_table(rows)]


def _smartart_text(shape: Any) -> str:
    """The words of a SmartArt graphic, one entry per node, from its diagram data part."""
    from lxml import etree

    rel = shape._element.find(f".//{{{_DIAGRAM}}}relIds")
    rid = rel.get(f"{_R}dm") if rel is not None else None
    if not rid:
        return ""
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    root = etree.fromstring(shape.part.related_part(rid).blob, parser)
    words = []
    for pt in root.iter(f"{{{_DIAGRAM}}}pt"):
        text = " ".join("".join(t.text or "" for t in pt.iter(f"{_A}t")).split())
        if text:
            words.append(text)
    return "; ".join(words)


def _alt_text(shape: Any) -> str:
    nv = shape._element.find(f".//{_P}cNvPr")
    return (nv.get("descr") or "").strip() if nv is not None else ""


def _is_picture(shape: Any) -> bool:
    """A picture, in a placeholder or not: both are p:pic in the slide XML."""
    return bool(shape._element.tag == f"{_P}pic")


def _decorative(shape: Any) -> bool:
    """Marked decorative by its author (PowerPoint's flag, as the builder sets on logos)."""
    nv = shape._element.find(f".//{_P}cNvPr")
    if nv is None:
        return False
    for ext in nv.iter(f"{_A}ext"):
        if ext.get("uri", "").upper() == _DECORATIVE_EXT:
            return any(child.get("val", "1") in ("1", "true") for child in ext)
    return False


def _is_page_number(shape: Any) -> bool:
    """A box holding only a slide-number field is chrome, not content."""
    if not shape.has_text_frame:
        return False
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    fields = shape._element.findall(f".//{a}fld")
    return (
        bool(fields)
        and all(f.get("type") == "slidenum" for f in fields)
        and not shape._element.findall(f".//{a}r")
    )


def _iter_shapes(shapes: Any):
    for shape in shapes:
        if shape.shape_type is not None and shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _slide_title(slide: Any) -> tuple[str, Any]:
    title = slide.shapes.title
    if title is not None and title.has_text_frame and title.text_frame.text.strip():
        return " ".join(title.text_frame.text.split()), title
    # No placeholder: the largest text nearest the top is the title.
    best = None
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip() or shape.top is None:
            continue
        sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
        size = max(sizes, default=0)
        key = (-size, shape.top)
        if best is None or key < best[0]:
            best = (key, shape)
    if best is None:
        return "", None
    return " ".join(best[1].text_frame.text.split()), best[1]


def _shape_lines(shape: Any) -> list[str]:
    """One shape as markdown lines: text, a table, a chart's data, or a marker for what
    has no text of its own (a picture, an embedded object, a SmartArt graphic)."""
    if _is_page_number(shape):
        return []
    if getattr(shape, "has_chart", False) and shape.has_chart:
        try:
            return _chart_lines(shape.chart) + [""]
        except Exception:  # noqa: BLE001 - python-pptx reads only some chart types
            return _cached_chart_lines(shape.chart) + [""]
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return _md_table(rows) + [""]
    if _is_picture(shape):
        # Every picture gets a marker (the redesign asks for its file by it), unless
        # its author marked it decorative, as the builder does its logos.
        if _decorative(shape):
            return []
        alt = _alt_text(shape)
        return [f"[image: {alt}]" if alt else f"[image: no alt text, {shape.name}]", ""]
    if shape.shape_type in (7, 10):  # MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, LINKED_OLE_OBJECT
        return [f"[embedded object: {shape.ole_format.prog_id or 'unknown'}]", ""]
    graphic = shape._element.find(f".//{_A}graphicData")
    if graphic is not None and graphic.get("uri") == _DIAGRAM:
        words = _smartart_text(shape)
        return [f"[SmartArt: {words}]" if words else "[SmartArt]", ""]
    if shape.has_text_frame:
        lines = _text_lines(shape)
        return lines + [""] if lines else []
    return []


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    try:
        nbg_package.check_package(path)
    except nbg_package.PackageError as e:
        raise ExtractError(f"{path.name} was not opened: {e}") from e
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001 - python-pptx raises several types for a bad file
        raise ExtractError(f"{path.name} is not a readable .pptx ({type(e).__name__}: {e})") from e
    out = [f"# {path.name}", ""]
    for number, slide in enumerate(prs.slides, 1):
        title, title_shape = _slide_title(slide)
        out.append(f"## Slide {number}: {title}" if title else f"## Slide {number}")
        out.append("")
        for shape in _iter_shapes(slide.shapes):
            # python-pptx hands out a new proxy per access; compare the XML elements.
            if title_shape is not None and shape._element is title_shape._element:
                continue
            try:
                out += _shape_lines(shape)
            except Exception as e:  # noqa: BLE001 - one odd shape must not lose the deck
                out += [f"[shape: {shape.name}, not read ({type(e).__name__})]", ""]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out += [f"Notes: {notes}", ""]
    return "\n".join(out).rstrip() + "\n"


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"{path.name} is not a readable PDF ({type(e).__name__}: {e})") from e
    out = [f"# {path.name}", ""]
    for number, page in enumerate(reader.pages, 1):
        out += [f"## Page {number}", "", (page.extract_text() or "").strip(), ""]
    return "\n".join(out).rstrip() + "\n"


def extract(path: Path) -> str:
    if not path.is_file():
        raise ExtractError(f"file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        raise ExtractError(
            ".docx is not supported: python-docx is not a dependency of the decks tools. "
            "Save the document as PDF and extract that"
        )
    raise ExtractError(
        f"cannot extract {suffix or 'a file with no extension'}: use a .pptx or a .pdf"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="nbg_extract.py", description="A deck or PDF as markdown."
    )
    parser.add_argument("file", help=".pptx or .pdf")
    args = parser.parse_args(argv)
    try:
        text = extract(Path(args.file).expanduser().resolve())
    except ExtractError as e:
        print(f"nbg_extract.py: {e}", file=sys.stderr)
        return EXIT_ERROR
    reconfigure = getattr(sys.stdout, "reconfigure", None)
    if reconfigure is not None:
        reconfigure(encoding="utf-8")  # Greek decks on a Windows console
    sys.stdout.write(text)
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
