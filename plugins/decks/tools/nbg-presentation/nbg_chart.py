#!/usr/bin/env python3
"""Native PowerPoint charts in the NBG chart style, for nbg_build.py.

Every value comes from tokens.yaml `charts`: the palette, the 12pt axis, label and
legend type, the 3.5pt line with hollow circle markers, the 55% doughnut hole, the
waterfall colours. Two things python-pptx gets wrong by default are set explicitly
on every chart, because PowerPoint shows the theme colour whenever the chart part
does not name one:

- a line series takes its colour from `c:spPr/a:ln`, not from a fill, so writing
  the palette into the fill left every line chart in Office blue and red;
- a stacked or in-bar label needs its own colour, measured against the fill it
  sits on, or it inherits #202020 and reads at 1.27:1 on a dark teal segment.

`area_line` is a line chart with an area chart under it on the same axes: python-
pptx cannot write a combination chart, so the area series are added to the plot
area as XML after the line chart is built.

A chart comparing the systemic banks (tokens.yaml `banks`) draws each bank in its
brand colour. When logos go under the category axis, the plot area is placed
exactly (a manual layout) so that nbg_build can put each logo under its own bar.
"""

from __future__ import annotations

import copy
import itertools
import math
import sys
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_LABEL_POSITION,
    XL_TICK_MARK,
)
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import nbg_spec  # noqa: E402
import nbg_tokens  # noqa: E402
from nbg_color import AA_NORMAL, contrast_ratio  # noqa: E402
from nbg_text import format_number, localise_number, metrics  # noqa: E402

CHARTS = nbg_tokens.load()["charts"]
FONT = str(CHARTS["font"])
PALETTE = nbg_tokens.chart_palette()
BANKS = nbg_tokens.load()["banks"]
BANK_COLOURS = nbg_tokens.get("extended_palettes.peer_banks")
LOGOS = nbg_tokens.load()["components"]["bank_logos"]
FIT = 0.97  # the measuring margin nbg_build uses: 97% of a width counts as full

XL_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar_horizontal": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area_line": XL_CHART_TYPE.LINE_MARKERS,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}

# The dark candidate for a label sitting on a coloured fill. PURE BLACK, not the
# #202020 body text: see label_text_color for why the difference decides AA.
LABEL_DARK_HEX = "000000"


def hexcolor(token: str) -> str:
    return str(nbg_tokens.color(token))


def rgb(token: str) -> RGBColor:
    return RGBColor.from_string(hexcolor(token))


def label_text_color(fill_hex: str) -> RGBColor:
    """White or black, whichever contrasts more with the fill a label sits on.

    colors.md offers an R+G+B > 400 rule of thumb, which misreads the cyan #00ADBF
    used for chart series: dark text clears 6:1 on it where white manages 2.7:1.
    Measure the contrast and take the better of the two.

    The dark candidate has to be pure black. For any candidate pair the worst case
    is the fill luminance where the two curves cross, and only black puts that
    crossing above the 4.5:1 AA floor:

        black   crossover at L=0.1791, best available ratio 4.583:1  clears AA
        #202020 crossover at L=0.2101, best available ratio 4.036:1  fails AA

    These labels are 12pt bold, which is not WCAG large text (that needs 14pt bold),
    so the floor is 4.5:1. Hand-specified dark text stays #202020 everywhere else.
    """
    white = contrast_ratio(fill_hex, "FFFFFF")
    dark = contrast_ratio(fill_hex, LABEL_DARK_HEX)
    if max(white, dark) < AA_NORMAL:
        raise ValueError(
            f"no AA-compliant label colour for fill #{fill_hex}: white {white:.2f}:1, "
            f"black {dark:.2f}:1, both under {AA_NORMAL}:1"
        )
    return RGBColor.from_string("FFFFFF" if white > dark else LABEL_DARK_HEX)


def default_number_format(values: list[Any]) -> str:
    """Thousands separator and the data's own precision, at most two decimals."""
    decimals = 0
    for v in values:
        if isinstance(v, float) and not v.is_integer():
            text = f"{v:.6f}".rstrip("0")
            decimals = max(decimals, min(2, len(text.split(".")[1])))
    return "#,##0" + ("." + "0" * decimals if decimals else "")


def _all_values(series: list[dict[str, Any]]) -> list[Any]:
    return [v for s in series for v in (s.get("values") or []) if v is not None]


# ---------------------------------------------------------------- peer banks


def bank_colours(spec: dict[str, Any]) -> tuple[list[str], list[str] | None]:
    """(colour per series, colour per point or None). A bank comparison paints each
    bank in its brand colour, per series or per point, and mutes whatever else shares
    the chart (a market average). Any other chart takes the palette in order."""
    count = len(spec["data"]["series"])
    palette = [PALETTE[i % len(PALETTE)] for i in range(count)]
    plan = nbg_spec.bank_plan(spec)
    if plan is None:
        return palette, None
    mode, banks = plan
    muted = hexcolor(CHARTS["highlight"]["muted"])
    colours = [str(BANK_COLOURS[b]).upper() if b else muted for b in banks]
    return (colours, None) if mode == "series" else (palette, colours)


def logo_path(bank: str) -> Path:
    return Path(nbg_spec.ASSETS_DIR) / str(BANKS[bank]["logo"])


def logo_width(bank: str, height: float) -> float:
    """The width a bank's logo takes at `height`, keeping its own aspect (Standard #4)."""
    from PIL import Image

    with Image.open(logo_path(bank)) as img:
        w_px, h_px = img.size
    return float(height * w_px / h_px)


@dataclass(frozen=True)
class AxisLayout:
    """Where a bank bar chart's plot sits in its frame, and where each logo goes."""

    inner: tuple[float, float, float, float]  # x, y, w, h as fractions of the frame
    anchors: list[tuple[float, float]]  # each category's logo centre, in slide inches
    logo_h: float


_PLACEHOLDERS = set("0#?.,")


def format_value(value: float, number_format: str | None, lang: str = "en") -> str:
    """value as an Excel number format shows it, for the formats decks use: 0, 0.0,
    #,##0.00, 0% and 0.0% (times 100), text in quotes or after a backslash ("EUR "#,##0,
    #,##0" m", 0.0\\x), a currency tag ([$€-408]) and scaling commas (#,##0,, is
    millions). Only the first ;-section is read, a negative takes a leading minus, and
    General writes the number as it is. Greek separators when lang is el."""
    section = (number_format or "General").split(";")[0]
    sign = "-" if value < 0 else ""
    magnitude = abs(float(value))
    if section.strip().casefold() in ("general", ""):
        return sign + str(localise_number(f"{magnitude:.10g}", lang))
    parts: list[tuple[str, str]] = []
    percent = 0
    i = 0
    while i < len(section):
        ch = section[i]
        if ch == '"':
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            parts.append(("text", section[i + 1 : end]))
            i = end + 1
        elif ch == "\\" and i + 1 < len(section):
            parts.append(("text", section[i + 1]))
            i += 2
        elif ch in "_*" and i + 1 < len(section):
            i += 2  # padding and fill print nothing
        elif ch == "[":
            end = section.find("]", i)
            end = len(section) if end < 0 else end
            tag = section[i + 1 : end]
            if tag.startswith("$"):
                parts.append(("text", tag[1:].split("-")[0]))
            i = end + 1
        elif ch in _PLACEHOLDERS:
            end = i
            while end < len(section) and section[end] in _PLACEHOLDERS:
                end += 1
            parts.append(("number", section[i:end]))
            i = end
        else:
            percent += ch == "%"
            parts.append(("text", ch))
            i += 1
    core = "".join(p for kind, p in parts if kind == "number")
    whole, _, fraction = core.partition(".")
    decimals = sum(ch in "0#?" for ch in fraction)
    scale = 1000 ** (len(core) - len(core.rstrip(",")))
    number = format_number(
        magnitude / scale * 100**percent, decimals, lang, group="," in whole.rstrip(",")
    )
    out: list[str] = []
    for kind, p in parts:
        if kind == "text":
            out.append(p)
        elif number:
            out.append(number)
            number = ""  # one number, however many placeholder runs
    return sign + "".join(out)


def _value_text(value: Any, number_format: str) -> str:
    """What a data label shows, so its width can be measured."""
    return "" if value is None else format_value(float(value), number_format)


def axis_layout(spec: dict[str, Any], box: tuple[float, float, float, float]) -> AxisLayout:
    """The plot area of a bank bar chart, leaving room for the logos: a band under the
    category labels of a column chart, a column left of them on a horizontal one.
    Raises ValueError when the frame cannot keep a readable plot beside the logos."""
    x, y, w, h = box
    categories = [str(c) for c in spec["data"]["categories"]]
    n = len(categories)
    plan = nbg_spec.bank_plan(spec)
    banks = plan[1] if plan else [None] * n
    logo_h, gap = float(LOGOS["h"]), float(LOGOS["gap"])
    m = metrics()
    axis_pt = float(nbg_tokens.get("type.chart_axis.size"))
    label_pt = float(nbg_tokens.get("type.chart_data_label.size"))
    values = spec["data"]["series"][0]["values"]
    number_format = str(spec.get("number_format") or default_number_format(values))
    if spec["type"] == "bar_horizontal":
        column = max(logo_width(b, logo_h) for b in banks if b)
        labels_w = max(m.width(c, axis_pt, False) for c in categories) / FIT + gap
        values_w = (
            max(m.width(_value_text(v, number_format), label_pt, True) for v in values) / FIT
            + 2 * gap
        )
        plot_w = w - column - gap - labels_w - values_w
        if plot_w < 1.5:
            raise ValueError(f"the bars get {plot_w:.2f} in beside the logos and labels")
        pad = 0.05
        inner = ((column + gap + labels_w) / w, pad / h, plot_w / w, (h - 2 * pad) / h)
        anchors = [
            (x + column / 2, y + h * (inner[1] + inner[3] * (j + 0.5) / n)) for j in range(n)
        ]
        return AxisLayout(inner, anchors, logo_h)
    top = m.line_height(label_pt, 1.0) + gap
    labels_h = m.line_height(axis_pt, 1.0) + gap
    plot_h = h - top - labels_h - gap - logo_h
    if plot_h < 1.0:
        raise ValueError(f"the bars get {plot_h:.2f} in above the logos")
    logo_y = y + top + plot_h + labels_h + gap + logo_h / 2
    anchors = [(x + w * (j + 0.5) / n, logo_y) for j in range(n)]
    return AxisLayout((0.0, top / h, 1.0, plot_h / h), anchors, logo_h)


def _manual_layout(chart: Any, inner: tuple[float, float, float, float]) -> None:
    """Place the plot area's inner rectangle (the bars, not their labels) exactly."""
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        plot_area.insert(0, layout)  # CT_PlotArea: layout? comes first
    for child in list(layout):
        layout.remove(child)
    manual = etree.SubElement(layout, qn("c:manualLayout"))
    etree.SubElement(manual, qn("c:layoutTarget")).set("val", "inner")
    etree.SubElement(manual, qn("c:xMode")).set("val", "edge")
    etree.SubElement(manual, qn("c:yMode")).set("val", "edge")
    for tag, value in zip(("c:x", "c:y", "c:w", "c:h"), inner, strict=True):
        etree.SubElement(manual, qn(tag)).set("val", f"{value:.4f}")


# ---------------------------------------------------------------- alt text

ALT = {
    "en": {
        "bar": "Column chart",
        "bar_stacked": "Stacked column chart",
        "bar_horizontal": "Horizontal bar chart",
        "line": "Line chart",
        "area_line": "Area line chart",
        "doughnut": "Doughnut chart",
        "waterfall": "Waterfall chart",
        "series": "series",
        "categories": "categories",
        "to": "to",
        "from": "from",
        "steps": "steps",
    },
    "el": {
        "bar": "Ραβδόγραμμα στηλών",
        "bar_stacked": "Σωρευτικό ραβδόγραμμα",
        "bar_horizontal": "Οριζόντιο ραβδόγραμμα",
        "line": "Γράφημα γραμμής",
        "area_line": "Γράφημα γραμμής με περιοχή",
        "doughnut": "Γράφημα δακτυλίου",
        "waterfall": "Γράφημα καταρράκτη",
        "series": "σειρές",
        "categories": "κατηγορίες",
        "to": "έως",
        "from": "από",
        "steps": "βήματα",
    },
}


def chart_alt_text(
    chart_type: str, categories: list[Any], series: list[dict[str, Any]], lang: str
) -> str:
    """What the chart plots, in how many series, over which span. Never the title."""
    words = ALT["el" if lang == "el" else "en"]
    names = ", ".join(
        str(s.get("name", "")).strip() for s in series if str(s.get("name", "")).strip()
    )
    parts = [words.get(chart_type, chart_type)]
    parts.append(f"{len(series)} {words['series']}" + (f" ({names})" if names else ""))
    if categories:
        span = (
            f"{categories[0]} {words['to']} {categories[-1]}"
            if len(categories) > 1
            else str(categories[0])
        )
        parts.append(f"{len(categories)} {words['categories']}, {span}")
    return ". ".join(parts) + "."


# ---------------------------------------------------------------- shared styling


def _set_font(
    font: Any, role: str, *, color_hex: str | None = None, bold: bool | None = None
) -> None:
    style = nbg_tokens.type_style(role)
    font.name = FONT
    font.size = Pt(style["size"])
    font.bold = style["bold"] if bold is None else bold
    font.color.rgb = RGBColor.from_string(color_hex or style["color"])


def _no_line(fmt: Any) -> None:
    fmt.line.fill.background()


def _chart_space(chart: Any) -> None:
    """No title, square corners, transparent chart area, Aptos for every chart text."""
    chart.has_title = False
    cs = chart._chartSpace
    rounded = cs.find(qn("c:roundedCorners"))
    if rounded is None:
        # CT_ChartSpace: date1904?, lang?, roundedCorners?, AlternateContent?,
        # clrMapOvr?, pivotSource?, protection?, chart. Absent means rounded.
        rounded = etree.Element(qn("c:roundedCorners"))
        successors = (
            "c:AlternateContent",
            "c:clrMapOvr",
            "c:pivotSource",
            "c:protection",
            "c:chart",
        )
        anchor = next(el for tag in successors if (el := cs.find(qn(tag))) is not None)
        anchor.addprevious(rounded)
    rounded.set("val", "0")
    chart.font.name = FONT
    chart.font.size = Pt(nbg_tokens.get("type.chart_axis.size"))
    chart.font.color.rgb = rgb(str(nbg_tokens.get("type.chart_axis.color")))
    # Chart area: no fill, no border. spPr follows c:chart in CT_ChartSpace.
    sppr = cs.find(qn("c:spPr"))
    if sppr is None:
        sppr = etree.Element(qn("c:spPr"))
        cs.find(qn("c:chart")).addnext(sppr)
    for child in list(sppr):
        sppr.remove(child)
    etree.SubElement(sppr, qn("a:noFill"))
    ln = etree.SubElement(sppr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    _plot_area_clear(plot_area)
    _positive_axis_ids(cs)


def _positive_axis_ids(chart_space: Any) -> None:
    """python-pptx's chart templates hard-code negative axis ids (-2068027336), which
    ST_UnsignedIntHex rejects: every such chart failed schema validation. Flip the
    sign on every id and on every reference to it."""
    for tag in ("c:axId", "c:crossAx"):
        for el in chart_space.iter(qn(tag)):
            value = int(el.get("val", "0"))
            if value < 0:
                el.set("val", str(-value))


def _plot_area_clear(plot_area: Any) -> None:
    """A transparent, borderless plot area. spPr sits after the axes in CT_PlotArea."""
    sppr = plot_area.find(qn("c:spPr"))
    if sppr is None:
        sppr = etree.SubElement(plot_area, qn("c:spPr"))
        ext = plot_area.find(qn("c:extLst"))
        if ext is not None:
            plot_area.remove(sppr)
            ext.addprevious(sppr)
    for child in list(sppr):
        sppr.remove(child)
    etree.SubElement(sppr, qn("a:noFill"))
    ln = etree.SubElement(sppr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def _category_axis(chart: Any, *, reverse: bool = False) -> None:
    axis = chart.category_axis
    style = CHARTS["category_axis"]
    _set_font(axis.tick_labels.font, "chart_axis", color_hex=hexcolor(style["color"]))
    axis.has_major_gridlines = False
    axis.has_minor_gridlines = False
    axis.major_tick_mark = XL_TICK_MARK.NONE
    axis.minor_tick_mark = XL_TICK_MARK.NONE
    axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    axis.format.line.color.rgb = rgb(style["line"])
    axis.format.line.width = Pt(style["line_pt"])
    if reverse:
        axis.reverse_order = True


def _value_axis(
    chart: Any,
    *,
    visible: bool,
    number_format: str,
    from_zero: bool = False,
    top: float | None = None,
) -> None:
    """from_zero pins the axis minimum at 0: PowerPoint's automatic scale starts close
    values (93, 91, 90, 88) near 86, so a bar's length stops meaning its value. top
    fixes the maximum, where the builder needs to know how tall a value draws."""
    axis = chart.value_axis
    axis.has_major_gridlines = False
    axis.has_minor_gridlines = False
    axis.major_tick_mark = XL_TICK_MARK.NONE
    axis.minor_tick_mark = XL_TICK_MARK.NONE
    if from_zero:
        axis.minimum_scale = 0
    if top is not None:
        axis.maximum_scale = top
    if not visible:
        axis.visible = False
        return
    muted = hexcolor(CHARTS["value_axis"]["muted_color"])
    font = axis.tick_labels.font
    font.name = FONT
    font.size = Pt(CHARTS["value_axis"]["size"])
    font.bold = False
    font.color.rgb = RGBColor.from_string(muted)
    axis.tick_labels.number_format = number_format
    axis.tick_labels.number_format_is_linked = False
    axis.format.line.fill.background()


def _legend(chart: Any, show: bool) -> None:
    chart.has_legend = show
    if not show:
        return
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    _set_font(chart.legend.font, "chart_legend")


def _labels(
    labels: Any, number_format: str, *, color_hex: str | None = None, position: Any = None
) -> None:
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.show_value = True
    labels.show_category_name = False
    labels.show_series_name = False
    labels.show_legend_key = False
    _set_font(labels.font, "chart_data_label", color_hex=color_hex)
    if position is not None:
        labels.position = position


def _delete_point_label(point: Any) -> None:
    """Hide one point's label. Blanking the text is not enough: showVal still wins."""
    dlbl = point.data_label._get_or_add_dLbl()
    for child in list(dlbl):
        if child.tag != qn("c:idx"):
            dlbl.remove(child)
    etree.SubElement(dlbl, qn("c:delete")).set("val", "1")


# ---------------------------------------------------------------- chart types


@dataclass(frozen=True)
class StackPlan:
    """Where a stacked column chart's plot sits, where its value axis ends, and which
    segment labels have no room."""

    inner: tuple[float, float, float, float]  # x, y, w, h as fractions of the frame
    top: float | None  # the axis maximum, fixed when every stack is at or above zero
    dropped: list[tuple[int, int]]  # (series, category) of each label left off


def stack_plan(
    spec: dict[str, Any], box: tuple[float, float, float, float], legend: bool
) -> StackPlan:
    """Lay a stacked column chart out exactly, so each segment's height is known: the
    plot fills the frame above its category labels (and the legend), and the value
    axis runs from 0 to the tallest stack, which needs no headroom because every label
    sits inside its segment. A segment shorter than one label line, or narrower than
    its label, leaves its label off (E2E-OUTPUT-08: they overprinted each other). A
    stack below zero keeps the automatic axis, since the validator fails any bar axis
    minimum but 0, so its heights are estimates."""
    _, _, w, h = box
    m = metrics()
    series = spec["data"]["series"]
    categories = [str(c) for c in spec["data"]["categories"]]
    number_format = str(spec.get("number_format") or default_number_format(_all_values(series)))
    label_pt = float(nbg_tokens.get("type.chart_data_label.size"))
    axis_pt = float(CHARTS["category_axis"]["size"])

    def stack(j: int, sign: float) -> float:
        cells = [float(s["values"][j] or 0) for s in series if j < len(s["values"])]
        return sum(v for v in cells if v * sign > 0)

    up = max((stack(j, 1) for j in range(len(categories))), default=0.0)
    down = min((stack(j, -1) for j in range(len(categories))), default=0.0)
    span = (up - down) or 1.0
    slot = 0.96 * w / max(1, len(categories))
    lines = max((len(m.wrap(c, slot, axis_pt)) for c in categories), default=1)
    band = lines * m.line_height(axis_pt) + 0.15
    legend_h = _legend_band([str(s.get("name", "")) for s in series], w) if legend else 0.0
    plot_h = max(0.5, h - 0.1 - band - legend_h)
    bar_w = slot / (1 + float(CHARTS["bar"]["gap_width"]) / 100)
    line = m.line_height(label_pt) + 0.02
    dropped = []
    for i, s in enumerate(series):
        for j, v in enumerate(s["values"][: len(categories)]):
            if v in (None, 0):
                continue
            label_w = m.width(_value_text(v, number_format), label_pt, True) / FIT + 0.04
            if abs(float(v)) / span * plot_h < line or label_w > bar_w:
                dropped.append((i, j))
    top = up if down == 0 and up > 0 else None
    return StackPlan((0.02, 0.1 / h, 0.96, plot_h / h), top, dropped)


def _warn_dropped(
    warn: Callable[[str, str, str], None], spec: dict[str, Any], stack: StackPlan
) -> None:
    series, categories = spec["data"]["series"], spec["data"]["categories"]
    named = [f"'{series[i].get('name', '')}' in {categories[j]}" for i, j in stack.dropped]
    listed = ", ".join(named[:6]) + (f" and {len(named) - 6} more" if len(named) > 6 else "")
    warn(
        "data.series",
        f"{len(named)} segment label(s) left off, the segments too thin to hold them: {listed}",
        "group the small series into 'Other', or give the chart more room; the values "
        "stay in the chart data",
    )


def _style_bars(
    chart: Any,
    spec: dict[str, Any],
    number_format: str,
    series_colours: list[str],
    point_colours: list[str] | None,
    stack: StackPlan | None = None,
) -> None:
    ctype = spec["type"]
    plot = chart.plots[0]
    plot.gap_width = int(CHARTS["bar"]["gap_width"])
    stacked = ctype == "bar_stacked"
    if stacked:
        plot.overlap = int(CHARTS["bar"]["overlap_stacked"])
    series_list = list(plot.series)
    highlight = spec.get("highlight_category")
    categories = [str(c) for c in spec["data"]["categories"]]
    for i, series in enumerate(series_list):
        fill = series_colours[i]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(fill)
        _no_line(series.format)
        series.invert_if_negative = False
        values = spec["data"]["series"][i]["values"]
        if stacked:
            _labels(
                series.data_labels,
                number_format,
                color_hex=str(label_text_color(fill)),
                position=XL_LABEL_POSITION.CENTER,
            )
            dropped = {j for s, j in stack.dropped if s == i} if stack else set()
            for j, (point, value) in enumerate(zip(series.points, values, strict=False)):
                if value in (None, 0) or j in dropped:
                    _delete_point_label(point)
        if point_colours is not None:
            # Each bank's bar in its brand colour; the labels sit outside, on white.
            for point, colour in zip(series.points, point_colours, strict=False):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(colour)
                _no_line(point.format)
        elif highlight is not None and len(series_list) == 1:
            accent = hexcolor(CHARTS["highlight"]["accent"])
            muted = hexcolor(CHARTS["highlight"]["muted"])
            for point, category in zip(series.points, categories, strict=False):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(
                    accent if category == str(highlight) else muted
                )
                _no_line(point.format)
    if not stacked:
        plot.has_data_labels = True
        _labels(plot.data_labels, number_format, position=XL_LABEL_POSITION.OUTSIDE_END)
    _category_axis(chart, reverse=ctype == "bar_horizontal")
    values = _all_values(spec["data"]["series"])
    _value_axis(
        chart,
        visible=CHARTS["value_axis"]["bar"] != "hidden",
        number_format=number_format,
        from_zero=all(float(v) >= 0 for v in values),
        top=stack.top if stack else None,
    )


def _style_line_series(series: Any, colour: str) -> None:
    line_cfg = CHARTS["line"]
    width = Pt(line_cfg["width_pt"])
    series.format.line.color.rgb = RGBColor.from_string(colour)
    series.format.line.width = width
    series.smooth = bool(line_cfg["smooth"])
    marker = series.marker
    marker.style = XL_MARKER_STYLE.CIRCLE
    marker.size = int(line_cfg["marker_size"])
    marker.format.fill.solid()
    marker.format.fill.fore_color.rgb = rgb(line_cfg["marker_fill"])
    marker.format.line.color.rgb = RGBColor.from_string(colour)
    marker.format.line.width = Pt(line_cfg["marker_line_pt"])


def _style_lines(
    chart: Any, spec: dict[str, Any], number_format: str, series_colours: list[str]
) -> None:
    plot = chart.plots[0]
    plot.has_data_labels = False
    for i, series in enumerate(plot.series):
        _style_line_series(series, series_colours[i])
    if len(spec["data"]["series"]) > 1:
        _name_line_ends(chart, spec)
    _category_axis(chart)
    _value_axis(
        chart, visible=CHARTS["value_axis"]["line"] != "hidden", number_format=number_format
    )
    if spec["type"] == "area_line":
        _add_area_under_lines(chart, series_colours)


def _name_line_ends(chart: Any, spec: dict[str, Any]) -> None:
    """Each line named at its last point, right of it: Standard #22 wants a series
    labelled where it is, not keyed by colour to a legend (E2E-OUTPUT-05)."""
    for series, data in zip(chart.plots[0].series, spec["data"]["series"], strict=True):
        points = [i for i, v in enumerate(data["values"]) if v is not None]
        if not points:
            continue
        label = series.points[points[-1]].data_label
        _point_label_shows(label, "General", series_name=True)
        _set_font(label.font, "chart_legend")
        label.position = XL_LABEL_POSITION.RIGHT


def end_label_layout(spec: dict[str, Any], box: tuple[float, float, float, float]) -> AxisLayout:
    """The plot area of a multi-series line chart, leaving the value axis its labels on
    the left and the series names their room on the right of the last points."""
    _, _, w, h = box
    m = metrics()
    values = [abs(float(v)) for v in _all_values(spec["data"]["series"])] or [0.0]
    number_format = str(spec.get("number_format") or default_number_format(values))
    axis_pt = float(CHARTS["value_axis"]["size"])
    left = m.width(_value_text(max(values) * 1.25, number_format), axis_pt, False) / FIT + 0.15
    names = [str(s.get("name", "")) for s in spec["data"]["series"]]
    legend_pt = float(nbg_tokens.get("type.chart_legend.size"))
    right = max(m.width(n, legend_pt, False) for n in names) / FIT + 0.25
    top, bottom = 0.1, m.line_height(float(CHARTS["category_axis"]["size"]), 1.0) + 0.12
    inner = (left / w, top / h, max(0.2, (w - left - right) / w), max(0.2, (h - top - bottom) / h))
    return AxisLayout(inner, [], 0.0)


def _add_area_under_lines(chart: Any, colours: list[str]) -> None:
    """Put an area chart under the first line, on the same axes: the 15% fill. One
    area only: under three or more lines the 15% fills blended into a grey-brown wash
    (E2E-OUTPUT-03), so the others stay plain lines."""
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    line_chart = plot_area.find(qn("c:lineChart"))
    series = line_chart.findall(qn("c:ser"))[:1]
    alpha = int(round(float(CHARTS["area_line"]["fill_alpha"]) * 100000))
    area = etree.Element(qn("c:areaChart"))
    etree.SubElement(area, qn("c:grouping")).set("val", "standard")
    etree.SubElement(area, qn("c:varyColors")).set("val", "0")
    offset = len(line_chart.findall(qn("c:ser")))
    for i, ser in enumerate(series):
        colour = colours[i]
        new = etree.SubElement(area, qn("c:ser"))
        etree.SubElement(new, qn("c:idx")).set("val", str(offset + i))
        etree.SubElement(new, qn("c:order")).set("val", str(offset + i))
        tx = ser.find(qn("c:tx"))
        if tx is not None:
            new.append(copy.deepcopy(tx))
        sppr = etree.SubElement(new, qn("c:spPr"))
        fill = etree.SubElement(sppr, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", colour)
        etree.SubElement(clr, qn("a:alpha")).set("val", str(alpha))
        ln = etree.SubElement(sppr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
        for tag in ("c:cat", "c:val"):
            el = ser.find(qn(tag))
            if el is not None:
                new.append(copy.deepcopy(el))
    for ax_id in line_chart.findall(qn("c:axId")):
        area.append(copy.deepcopy(ax_id))
    line_chart.addprevious(area)
    legend = chart._chartSpace.find(qn("c:chart")).find(qn("c:legend"))
    if legend is not None:
        # Legend entries follow plot order, and the area chart comes first so it draws
        # under the lines: the first entries are the area copies.
        _hide_legend_entries(legend, range(len(series)))


def _hide_legend_entries(legend: Any, indices: Any) -> None:
    """The area copies share the lines' names; one legend entry per series is enough."""
    anchor = legend.find(qn("c:legendPos"))
    for idx in indices:
        entry = etree.Element(qn("c:legendEntry"))
        etree.SubElement(entry, qn("c:idx")).set("val", str(idx))
        etree.SubElement(entry, qn("c:delete")).set("val", "1")
        anchor.addnext(entry)
        anchor = entry


def _point_label_shows(
    label: Any,
    number_format: str,
    *,
    percent: bool = False,
    category: bool = False,
    series_name: bool = False,
) -> None:
    """A point's own c:dLbl overrides the series labels entirely, flags included, so
    it has to say what to show (and in which format) itself. The value shows unless
    the share or the series name stands in for it; a category name goes on its own
    line above the number."""
    value = not (percent or series_name)
    dlbl = label._get_or_add_dLbl()
    flags = {
        "c:showLegendKey": "0",
        "c:showVal": "1" if value else "0",
        "c:showCatName": "1" if category else "0",
        "c:showSerName": "1" if series_name else "0",
        "c:showPercent": "1" if percent else "0",
        "c:showBubbleSize": "0",
    }
    for tag, flag in flags.items():
        el = dlbl.find(qn(tag))
        if el is not None:
            el.set("val", flag)
    if category and dlbl.find(qn("c:separator")) is None:
        separator = etree.Element(qn("c:separator"))
        separator.text = "\n"
        dlbl.find(qn("c:showBubbleSize")).addnext(
            separator
        )  # CT_DLbl: ...showBubbleSize, separator
    if dlbl.find(qn("c:numFmt")) is None:
        fmt = etree.Element(qn("c:numFmt"))
        fmt.set("formatCode", number_format)
        fmt.set("sourceLinked", "0")
        # CT_DLbl: idx, layout?, tx?, numFmt?, spPr?, txPr?, dLblPos?, show*...
        anchor = next(
            (
                el
                for tag in ("c:spPr", "c:txPr", "c:dLblPos", "c:showLegendKey")
                if (el := dlbl.find(qn(tag))) is not None
            ),
            None,
        )
        if anchor is not None:
            anchor.addprevious(fmt)
        else:
            dlbl.append(fmt)


@dataclass(frozen=True)
class RingPlan:
    """Where a doughnut sits in its frame and what each slice's label says."""

    inner: tuple[float, float, float, float]  # the plot square, as fractions of the frame
    names: list[str | None]  # each slice's name as drawn, line breaks included, or None
    values: list[bool]  # whether each slice shows its share (or value)
    legend: bool  # the chart's own legend names the slices
    unnamed: list[str]  # the slices whose name has no room inside the ring


RING_MARGIN = 0.05  # inches kept clear round the ring inside its frame
RING_TOLERANCE = 0.02  # how far a label may cross its slice's outline, in inches


def _slice_spans(values: list[Any]) -> list[tuple[float, float]]:
    """Each slice's start and end angle, clockwise from 12 o'clock (firstSliceAng 0)."""
    shares = [max(0.0, float(v or 0)) for v in values]
    total = sum(shares) or 1.0
    spans, start = [], 0.0
    for share in shares:
        end = start + share / total * 2 * math.pi
        spans.append((start, end))
        start = end
    return spans


def _off_ray(x: float, y: float, angle: float) -> float:
    """How far (x, y) lies from the ray leaving the ring's centre at angle."""
    ux, uy = math.sin(angle), math.cos(angle)
    if x * ux + y * uy <= 0:
        return math.hypot(x, y)
    return abs(x * uy - y * ux)


def _label_fits(w: float, h: float, span: tuple[float, float], outer: float, hole: float) -> bool:
    """Whether a w x h label, centred on its slice at mid-angle and mid-ring as both
    PowerPoint and LibreOffice centre it, stays on the slice: every point of its
    outline between the hole and the rim, and between the slice's two edges."""
    start, end = span
    if end <= start:
        return False
    mid, radius = (start + end) / 2, (outer + hole) / 2
    cx, cy = radius * math.sin(mid), radius * math.cos(mid)
    whole = end - start >= 2 * math.pi - 1e-9
    steps = [i / 8 for i in range(9)]
    outline = [(cx - w / 2 + w * t, cy + s * h / 2) for t in steps for s in (-1, 1)]
    outline += [(cx + s * w / 2, cy - h / 2 + h * t) for t in steps for s in (-1, 1)]
    for x, y in outline:
        r = math.hypot(x, y)
        if r < hole - RING_TOLERANCE or r > outer + RING_TOLERANCE:
            return False
        if whole or (math.atan2(x, y) - start) % (2 * math.pi) <= end - start:
            continue
        if min(_off_ray(x, y, start), _off_ray(x, y, end)) > RING_TOLERANCE:
            return False
    return True


def _name_options(name: str) -> list[list[str]]:
    """A slice name on one line, then broken between words into two and three lines,
    each time at the break whose widest line is narrowest."""
    m, size = metrics(), float(nbg_tokens.get("type.chart_data_label.size"))
    words = name.split()
    options = [[name]]
    for count in (2, 3):
        if len(words) < count:
            break
        splits = [
            [" ".join(words[a:b]) for a, b in zip((0, *cuts), (*cuts, len(words)), strict=True)]
            for cuts in itertools.combinations(range(1, len(words)), count - 1)
        ]
        options.append(min(splits, key=lambda lines: max(m.width(t, size, True) for t in lines)))
    return options


def _label_box(lines: list[str]) -> tuple[float, float]:
    m, size = metrics(), float(nbg_tokens.get("type.chart_data_label.size"))
    w = max(m.width(t, size, True) for t in lines) / FIT + 0.04
    return w, len(lines) * m.line_height(size) + 0.02


def _legend_band(names: list[str], width: float) -> float:
    """The height a bottom legend of these names takes, row wrapping included."""
    m, size = metrics(), float(nbg_tokens.get("type.chart_legend.size"))
    rows, used = 1, 0.0
    for name in names:
        entry = m.width(name, size) / FIT + 0.35  # the key, its gap, the space after
        if used and used + entry > width - 0.2:
            rows, used = rows + 1, 0.0
        used += entry
    return float(rows * m.line_height(size) + 0.15)


def _ring_square(
    box: tuple[float, float, float, float], band: float
) -> tuple[tuple[float, float, float, float], float]:
    """The plot square centred in the frame above a legend band, and the ring's radius."""
    _, _, w, h = box
    side = max(0.5, min(w, h - band) - 2 * RING_MARGIN)
    return ((w - side) / 2 / w, (h - band - side) / 2 / h, side / w, side / h), side / 2


def ring_plan(
    spec: dict[str, Any], box: tuple[float, float, float, float], legend: bool | None
) -> RingPlan:
    """Lay out a doughnut's labels. Each slice carries its name and share inside the
    ring (Standard #22), the name broken between words until the label fits its slice.
    When a slice has no room for its name at any break, a legend at the bottom names
    the slices instead (charts.md allows either) and each keeps only its share.

    legend is the caller's or the spec's choice: None decides by fit; False means no
    chart legend, because something else already names the slices (a peer-bank logo
    row) or the author said so, so the names that fit stay; True asks for the legend."""
    names = [str(c) for c in spec["data"]["categories"]]
    values = list(spec["data"]["series"][0]["values"])
    spans = _slice_spans(values)
    total = sum(max(0.0, float(v or 0)) for v in values) or 1.0
    if spec.get("number_format"):
        texts = [_value_text(v, str(spec["number_format"])) for v in values]
    else:
        texts = [f"{round(max(0.0, float(v or 0)) / total * 100)}%" for v in values]
    hole_share = float(CHARTS["doughnut"]["hole_size"]) / 100

    def fits(lines: list[str], i: int, outer: float) -> bool:
        return _label_fits(*_label_box(lines), spans[i], outer, outer * hole_share)

    unnamed: list[str] = []
    if legend is not True:
        inner, outer = _ring_square(box, 0.0)
        chosen = [
            next((o for o in _name_options(n) if fits([*o, texts[i]], i, outer)), None)
            for i, n in enumerate(names)
        ]
        unnamed = [n for n, c in zip(names, chosen, strict=True) if c is None]
        if not unnamed or legend is False:
            shown = [c is not None or fits([texts[i]], i, outer) for i, c in enumerate(chosen)]
            drawn = ["\n".join(c) if c else None for c in chosen]
            return RingPlan(inner, drawn, shown, False, unnamed)
    inner, outer = _ring_square(box, _legend_band(names, box[2]))
    shown = [fits([texts[i]], i, outer) for i in range(len(names))]
    return RingPlan(inner, [None] * len(names), shown, True, unnamed)


def _warn_unnamed(warn: Callable[[str, str, str], None], ring: RingPlan) -> None:
    quoted = ", ".join(f"'{n}'" for n in ring.unnamed)
    room = "has no room for its name" if len(ring.unnamed) == 1 else "have no room for their names"
    if ring.legend:
        warn(
            "data.categories",
            f"{quoted} {room} inside the ring, so the doughnut names its slices in a legend"
            " instead of on the slices",
            "merge the smallest slices into one, shorten the names, or give the chart more room",
        )
    else:
        warn(
            "data.categories",
            f"{quoted} {room} inside the ring and show_legend is false, so nothing names "
            + ("it" if len(ring.unnamed) == 1 else "them"),
            "drop show_legend: false, merge the smallest slices into one, or shorten the names",
        )


def _style_doughnut(
    chart: Any, spec: dict[str, Any], point_colours: list[str] | None, ring: RingPlan
) -> None:
    plot = chart.plots[0]
    doughnut = (
        chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:doughnutChart"))
    )
    hole = doughnut.find(qn("c:holeSize"))
    if hole is None:
        hole = etree.SubElement(doughnut, qn("c:holeSize"))
    hole.set("val", str(int(CHARTS["doughnut"]["hole_size"])))
    first = doughnut.find(qn("c:firstSliceAng"))
    if first is None:
        first = etree.Element(qn("c:firstSliceAng"))
        hole.addprevious(first)
    first.set("val", "0")
    as_percent = not spec.get("number_format")
    number_format = "0%" if as_percent else str(spec["number_format"])
    series = plot.series[0]
    slices = [str(c) for c in CHARTS["doughnut"]["slice_palette"]]
    for i, point in enumerate(series.points):
        fill = point_colours[i] if point_colours else slices[i % len(slices)]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(fill)
        point.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        point.format.line.width = Pt(1)
        if not ring.values[i]:
            _delete_point_label(point)
            continue
        label = point.data_label
        _set_font(label.font, "chart_data_label", color_hex=str(label_text_color(fill)))
        # No position: a doughnut draws its labels on the ring, and PowerPoint does
        # not accept dLblPos on a doughnut at all (its PDF export hung on one). Each
        # slice names itself where it has room: a colour-keyed legend fails #22.
        named = ring.names[i] is not None
        _point_label_shows(label, number_format, percent=as_percent, category=named)
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.show_percentage = as_percent
    labels.show_value = not as_percent
    labels.show_category_name = any(n is not None for n in ring.names)
    labels.show_series_name = False
    labels.show_legend_key = False
    _set_font(labels.font, "chart_data_label")


def add_chart(
    slide: Any,
    spec: dict[str, Any],
    box: tuple[float, float, float, float],
    lang: str,
    alt_text: str | None = None,
    *,
    legend: bool | None = None,
    layout: AxisLayout | None = None,
    warn: Callable[[str, str, str], None] | None = None,
) -> Any:
    """Draw spec (a deck.schema.json `chart`) into box (x, y, w, h in inches).

    legend overrides the spec's show_legend (False when a logo legend replaces it);
    layout places the plot area exactly, for logos along the category axis; warn
    receives (path in the chart spec, message, fix) for what the chart had to give
    up, such as a doughnut slice with no room for its name."""
    ctype = spec["type"]
    data = spec["data"]
    categories = [str(c) for c in data["categories"]]
    series_specs = data["series"]
    values = _all_values(series_specs)
    number_format = str(spec.get("number_format") or default_number_format(values))
    chart_data = CategoryChartData(number_format=number_format)
    ring = None
    if ctype == "doughnut":
        ring = ring_plan(spec, box, spec.get("show_legend") if legend is None else legend)
        if ring.unnamed and warn is not None and legend is None:
            _warn_unnamed(warn, ring)
    # A doughnut label shows its category, so a name broken to fit carries its breaks.
    chart_data.categories = [
        (ring.names[i] if ring and ring.names[i] else c) for i, c in enumerate(categories)
    ]
    for s in series_specs:
        chart_data.add_series(str(s.get("name", "")), list(s["values"]))
    x, y, w, h = box
    frame = slide.shapes.add_chart(
        XL_TYPES[ctype], Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    )
    chart = frame.chart
    _chart_space(chart)
    series_colours, point_colours = bank_colours(spec)
    # Only a multi-series bar chart needs a legend: a doughnut names its slices and a
    # multi-series line names each line at its end (Standard #22, E2E-OUTPUT-05).
    lines = ctype in ("line", "area_line")
    default_legend = ctype in ("bar", "bar_stacked", "bar_horizontal") and len(series_specs) > 1
    shown = bool(spec.get("show_legend", default_legend)) if legend is None else legend
    if ring is not None:
        shown, layout = ring.legend, AxisLayout(ring.inner, [], 0.0)
    _legend(chart, shown)
    if lines and len(series_specs) > 1 and layout is None:
        layout = end_label_layout(spec, box)
    stack = None
    if ctype == "bar_stacked" and layout is None:
        stack = stack_plan(spec, box, shown)
        layout = AxisLayout(stack.inner, [], 0.0)
        if stack.dropped and warn is not None:
            _warn_dropped(warn, spec, stack)
    if ctype in ("bar", "bar_stacked", "bar_horizontal"):
        _style_bars(chart, spec, number_format, series_colours, point_colours, stack)
    elif ctype in ("line", "area_line"):
        _style_lines(chart, spec, number_format, series_colours)
    else:
        assert ring is not None
        _style_doughnut(chart, spec, point_colours, ring)
    if layout is not None:
        _manual_layout(chart, layout.inner)
    set_alt_text(frame, alt_text or chart_alt_text(ctype, categories, series_specs, lang))
    return frame


# ---------------------------------------------------------------- waterfall


def waterfall_segments(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Each bar as {label, value, kind, start, end}. First and last are totals by default."""
    out = []
    running = 0.0
    last = len(items) - 1
    for i, item in enumerate(items):
        value = float(item["value"])
        total = bool(item.get("total", i in (0, last)))
        if total:
            start, end, kind = 0.0, value, "total"
        else:
            start, end = running, running + value
            kind = "increase" if value >= 0 else "decrease"
        running = end
        out.append(
            {"label": str(item["label"]), "value": value, "kind": kind, "start": start, "end": end}
        )
    return out


def _split(start: float, end: float) -> tuple[float, float, float]:
    """(invisible base, visible positive part, visible negative part) of one bar.

    A stacked column stacks positive values up from zero and negative values down
    from zero, each in series order, so a floating bar that crosses zero needs its
    two halves in two series and no base at all.
    """
    lo, hi = min(start, end), max(start, end)
    if lo >= 0:
        return lo, hi - lo, 0.0
    if hi <= 0:
        return hi, 0.0, lo - hi
    return 0.0, hi, lo


def _format_delta(value: float, kind: str, number_format: str, lang: str = "en") -> str:
    """A bar's label in the chart's own number format (BUILDER-CODE-01: the labels took
    only its decimals, so 0.00% printed 0.02 and a quoted unit vanished), signed: a
    step always, a total only when negative."""
    text = format_value(abs(value), number_format, lang)
    if kind == "total":
        return f"-{text}" if value < 0 else text
    return f"+{text}" if value >= 0 else f"-{text}"


def add_waterfall(
    slide: Any,
    spec: dict[str, Any],
    box: tuple[float, float, float, float],
    lang: str,
    alt_text: str | None = None,
) -> Any:
    """A bridge drawn as a stacked column with an invisible base (python-pptx has no
    native waterfall). Totals dark teal, increases cyan, decreases alert red; every
    label is the signed contribution, written as text so it reads right in any viewer.

    Four series: the invisible base, the part of each bar above zero, the part below
    zero, and an empty label carrier. Each visible part takes its kind's colour point
    by point, so the chart never looks like more than the one exhibit it is.

    A step's label sits just above its bar, on the carrier stacked on top: inside a
    thin bar a centred label runs over both edges. Totals, and steps below zero, keep
    a centred label in the bar."""
    cfg = CHARTS["waterfall"]
    segments = waterfall_segments(spec["data"]["items"])
    values = [s["value"] for s in segments]
    number_format = str(spec.get("number_format") or default_number_format(values))
    colours = {
        "total": hexcolor(cfg["total_color"]),
        "increase": hexcolor(cfg["increase_color"]),
        "decrease": hexcolor(cfg["decrease_color"]),
    }
    parts = [_split(seg["start"], seg["end"]) for seg in segments]
    chart_data = CategoryChartData(number_format=number_format)
    chart_data.categories = [s["label"] for s in segments]
    chart_data.add_series("Base", [b for b, _, _ in parts])
    chart_data.add_series("Above zero", [pos for _, pos, _ in parts])
    chart_data.add_series("Below zero", [neg for _, _, neg in parts])
    chart_data.add_series("Labels", [0.0] * len(segments))
    x, y, w, h = box
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    )
    chart = frame.chart
    _chart_space(chart)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = int(cfg["gap_width"])
    plot.overlap = 100
    base_series, above_zero, below_zero, carrier = plot.series
    for invisible in (base_series, carrier):
        invisible.format.fill.background()
        _no_line(invisible.format)
    for series in (above_zero, below_zero):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(colours["total"])
        _no_line(series.format)
        series.invert_if_negative = False
    for index, (seg, (_, pos, neg)) in enumerate(zip(segments, parts, strict=True)):
        for series, part in ((above_zero, pos), (below_zero, neg)):
            if part:
                point = series.points[index]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(colours[seg["kind"]])
                _no_line(point.format)
    above = hexcolor(str(nbg_tokens.get("type.chart_data_label.color")))
    for index, (seg, (_, pos, neg)) in enumerate(zip(segments, parts, strict=True)):
        floating_above_zero = seg["kind"] != "total" and neg == 0
        if floating_above_zero:
            label = carrier.points[index].data_label
            position, colour = XL_LABEL_POSITION.INSIDE_BASE, above
        else:
            holder = above_zero if pos != 0 or neg == 0 else below_zero
            label = holder.points[index].data_label
            position, colour = XL_LABEL_POSITION.CENTER, str(label_text_color(colours[seg["kind"]]))
        tf = label.text_frame
        tf.text = _format_delta(seg["value"], seg["kind"], number_format, lang)
        for run in tf.paragraphs[0].runs:
            _set_font(run.font, "chart_data_label", color_hex=colour)
        label.position = position
    _category_axis(chart)
    _value_axis(
        chart,
        visible=False,
        number_format=number_format,
        from_zero=all(min(s["start"], s["end"]) >= 0 for s in segments),
    )
    words = ALT["el" if lang == "el" else "en"]
    default_alt = (
        f"{words['waterfall']}. {words['from']} {segments[0]['label']} "
        f"{_format_delta(segments[0]['value'], 'total', number_format, lang)} {words['to']} "
        f"{segments[-1]['label']} "
        f"{_format_delta(segments[-1]['value'], 'total', number_format, lang)}, "
        f"{len(segments)} {words['steps']}."
    )
    set_alt_text(frame, alt_text or default_alt)
    return frame


def set_alt_text(shape: Any, text: str) -> None:
    """Write OOXML alt text (p:cNvPr/@descr). python-pptx exposes no API for it."""
    nv = shape._element.find(f".//{qn('p:cNvPr')}")
    if nv is not None:
        nv.set("descr", text)
