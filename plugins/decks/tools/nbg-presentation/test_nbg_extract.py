"""Tests for nbg_extract.py: a built deck read back as markdown, the way
/redesign-deck reads a colleague's deck."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import nbg_build
import nbg_extract
import pytest
from PIL import Image
from testkit import EXAMPLES_DIR, deck, passing_validator, write_spec

SCRIPT = Path(__file__).parent / "nbg_extract.py"
SOURCE = {"name": "Management accounts", "as_of": "30 June 2026"}


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> Path:
    tmp: Path = tmp_path_factory.mktemp("extract")
    spec = deck(
        [
            {
                "type": "content",
                "content": {
                    "title": "Three things moved",
                    "points": ["One", {"text": "Two", "level": 2}],
                },
                "notes": "Say the three things.",
            },
            {
                "type": "table",
                "content": {"title": "Fees grew", "source": SOURCE},
                "table": {"headers": ["Line", "2026"], "rows": [["Cards", "EUR 12M"]]},
            },
            {
                "type": "chart",
                "content": {"title": "Volumes rose", "source": SOURCE},
                "chart": {
                    "type": "bar",
                    "data": {
                        "categories": ["Q1", "Q2"],
                        "series": [{"name": "2026", "values": [3, 4]}],
                    },
                },
            },
            {
                "type": "waterfall",
                "content": {"title": "The bridge", "source": SOURCE},
                "chart": {
                    "type": "waterfall",
                    "data": {
                        "items": [
                            {"label": "Start", "value": 10},
                            {"label": "Up", "value": 5},
                            {"label": "End", "value": 15},
                        ]
                    },
                },
            },
            {
                "type": "image",
                "content": {"title": "A picture"},
                "image": {
                    "path": "illustrations/Growth.png",
                    "alt_text": "A rising savings balance",
                },
            },
        ]
    )
    out = tmp / "deck.pptx"
    nbg_build.build_presentation(write_spec(tmp, spec), out, validate=False)
    return out


def test_titles_bullets_notes_tables_charts_and_images(built):
    text = nbg_extract.extract(built)
    assert "## Slide 1: Test deck" in text
    assert "## Slide 2: Three things moved" in text
    assert "- One" in text and "  - Two" in text
    assert "Notes: Say the three things." in text
    assert "| Line | 2026 |" in text and "| Cards | EUR 12M |" in text
    assert "Chart (bar):" in text and "| Q1 | 3 |" in text
    assert "Chart (waterfall):" in text and "| Up | +5 |" in text
    assert "[image: A rising savings balance]" in text
    assert "Source: Management accounts, as of 30 June 2026" in text


def test_the_title_and_the_page_number_are_not_repeated_as_body_text(built):
    text = nbg_extract.extract(built)
    assert text.count("Three things moved") == 1
    assert "\n2\n" not in text and "\n3\n" not in text


def test_logos_are_not_reported_as_content(built):
    assert "[image: ]" not in nbg_extract.extract(built)


def test_every_example_extracts(tmp_path, monkeypatch):
    passing_validator(monkeypatch, nbg_build)
    for spec in sorted(EXAMPLES_DIR.glob("*.yaml")):
        out = tmp_path / f"{spec.stem}.pptx"
        nbg_build.build_presentation(spec, out)
        text = nbg_extract.extract(out)
        assert text.count("## Slide ") >= 5, spec.name


def test_a_pdf_extracts_page_by_page(tmp_path):
    pdf = tmp_path / "doc.pdf"
    pages = [Image.new("RGB", (200, 100), "white") for _ in range(2)]
    pages[0].save(pdf, save_all=True, append_images=pages[1:])
    text = nbg_extract.extract(pdf)
    assert "## Page 1" in text and "## Page 2" in text


def test_docx_is_refused_with_the_way_round_it(tmp_path, capsys):
    doc = tmp_path / "memo.docx"
    doc.write_bytes(b"PK")
    assert nbg_extract.main([str(doc)]) == 2
    assert "save the document as pdf" in capsys.readouterr().err.lower()


def test_a_missing_or_damaged_file_exits_2(tmp_path, capsys):
    assert nbg_extract.main([str(tmp_path / "nope.pptx")]) == 2
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip")
    assert nbg_extract.main([str(broken)]) == 2


def test_every_picture_gets_a_marker_even_without_alt_text_or_in_a_placeholder(tmp_path):
    """PROMPTS-CONTRACTS-05: a picture without alt text, or inside a picture placeholder,
    vanished from the markdown, and redesign asks for image files by those markers.
    A picture its author marked decorative (the builder's logos) stays out."""
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    png = tmp_path / "photo.png"
    Image.new("RGB", (60, 40), "teal").save(png)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    described = slide.shapes.add_picture(str(png), Inches(1), Inches(1))
    described._element.find(".//{*}cNvPr").set("descr", "The branch at night")
    bare = slide.shapes.add_picture(str(png), Inches(3), Inches(1))
    bare._element.find(".//{*}cNvPr").set("descr", "")  # as PowerPoint leaves it
    with_placeholder = prs.slides.add_slide(prs.slide_layouts[8])  # "Picture with Caption"
    placeholder = next(
        p for p in with_placeholder.placeholders if p.placeholder_format.type == 18
    )  # PP_PLACEHOLDER.PICTURE
    placed = placeholder.insert_picture(str(png))
    placed._element.find(".//{*}cNvPr").set("descr", "")
    logo = slide.shapes.add_picture(str(png), Inches(5), Inches(1))
    nbg_build.mark_decorative(logo)
    deck = tmp_path / "pictures.pptx"
    prs.save(str(deck))
    text = nbg_extract.extract(deck)
    assert "[image: The branch at night]" in text
    markers = [line for line in text.splitlines() if line.startswith("[image: no alt text")]
    assert len(markers) == 2, text  # the undescribed picture and the placeholder picture
    assert text.count("[image:") == 3, "the decorative logo is not content"


def _two_slide_deck(tmp_path, name, dress):
    """A deck whose first slide `dress` fills, and a second with plain text, so a test
    also sees whether extraction survives past slide 1."""
    from pptx import Presentation

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "Market share"
    dress(first)
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "Next steps"
    second.placeholders[1].text = "Launch the product"
    path = tmp_path / name
    prs.save(str(path))
    return path


def test_a_3d_chart_is_read_from_its_cached_values(tmp_path):
    """BUILDER-CODE-09: python-pptx cannot read a 3-D chart, and extract crashed with
    a traceback on the first one, so nothing of the deck came out."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    def pie_3d(slide):
        data = CategoryChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Share", [50, 30, 20])
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(1), Inches(1.5), Inches(6), Inches(4), data
        )
        pie = frame.chart.part._element.find(".//{*}pieChart")
        pie.tag = pie.tag.replace("pieChart", "pie3DChart")
        for child in list(pie):
            if child.tag.endswith("firstSliceAng"):
                pie.remove(child)

    text = nbg_extract.extract(_two_slide_deck(tmp_path, "pie3d.pptx", pie_3d))
    assert "Chart (3-D pie):" in text
    assert "| A | 50 |" in text and "| C | 20 |" in text
    assert "Launch the product" in text


def test_an_embedded_workbook_gets_a_marker(tmp_path):
    """BUILDER-CODE-09: an embedded Excel object vanished from the markdown."""
    import io

    from pptx.enum.shapes import PROG_ID
    from pptx.util import Inches

    def workbook(slide):
        slide.shapes.add_ole_object(
            io.BytesIO(b"not really a workbook"), PROG_ID.XLSX, Inches(1), Inches(2)
        )

    text = nbg_extract.extract(_two_slide_deck(tmp_path, "ole.pptx", workbook))
    assert "[embedded object: Excel.Sheet.12]" in text
    assert "Launch the product" in text


def test_smartart_gets_a_marker_with_its_text(tmp_path):
    """BUILDER-CODE-09: SmartArt vanished; its words live in the diagram data part."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    dgm = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"

    def smartart(slide):
        points = "".join(
            f'<dgm:pt modelId="{i}"><dgm:t><a:bodyPr/><a:p><a:r><a:t>{word}</a:t></a:r></a:p>'
            "</dgm:t></dgm:pt>"
            for i, word in enumerate(["Plan", "Build", "Launch"], 1)
        )
        blob = f'<dgm:dataModel xmlns:dgm="{dgm}" xmlns:a="{a}"><dgm:ptLst>{points}</dgm:ptLst></dgm:dataModel>'
        part = Part(
            PackURI("/ppt/diagrams/data1.xml"),
            "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
            slide.part.package,
            blob.encode(),
        )
        rid = slide.part.relate_to(part, RT.DIAGRAM_DATA)
        frame = etree.fromstring(
            '<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
            f' xmlns:a="{a}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<p:nvGraphicFramePr><p:cNvPr id="40" name="Diagram 1"/><p:cNvGraphicFramePr/><p:nvPr/>'
            '</p:nvGraphicFramePr><p:xfrm><a:off x="914400" y="1828800"/><a:ext cx="5486400" cy="1828800"/>'
            f'</p:xfrm><a:graphic><a:graphicData uri="{dgm}"><dgm:relIds xmlns:dgm="{dgm}"'
            f' r:dm="{rid}" r:lo="" r:qs="" r:cs=""/></a:graphicData></a:graphic></p:graphicFrame>'
        )
        slide.shapes._spTree.append(frame)

    text = nbg_extract.extract(_two_slide_deck(tmp_path, "smartart.pptx", smartart))
    assert "[SmartArt: Plan; Build; Launch]" in text
    assert "Launch the product" in text


def test_a_zip_bomb_is_refused_before_python_pptx_opens_it(tmp_path, capsys):
    """SECURITY-PUBLIC-2: a 406 KB crafted deck drove one extract to 583 MB."""
    import zipfile

    bomb = tmp_path / "bomb.pptx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/slides/slide1.xml", " " * (3 * 2**20))
    assert nbg_extract.main([str(bomb)]) == 2
    assert "zip bomb" in capsys.readouterr().err


def test_the_cli_writes_greek_as_utf8(tmp_path):
    out = tmp_path / "greek.pptx"
    nbg_build.build_presentation(EXAMPLES_DIR / "greek-deck.yaml", out, validate=False)
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), str(out)],
        capture_output=True,
        env={"PYTHONIOENCODING": "cp1252", "PATH": "/usr/bin:/bin"},
    )
    assert proc.returncode == 0, proc.stderr
    assert "Περιεχόμενα" in proc.stdout.decode("utf-8")
