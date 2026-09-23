#!/usr/bin/env python3
"""NBG presentation builder: renders a deck spec to .pptx, then validates it.

    decks-py build deck.yaml deck.pptx     (nbg_build.py deck.yaml deck.pptx)
    decks-py check deck.yaml [--format json]   (nbg_build.py --check deck.yaml)

The spec is tools/nbg-presentation/deck.schema.json. Every slide type in it is
drawn here, from scratch, with python-pptx: there is no template file. Every brand
value (colours, type scale, geometry, chart styling) comes from
shared/brand-system/tokens.yaml through nbg_tokens.py, so a brand change is one
edit in that file.

A build runs in three stages and stops at the first that fails:

1. check: nbg_spec.py normalises legacy names, validates the schema and the
   semantics, then this module lays every slide out in a dry run, which catches
   text that cannot fit. Nothing is written when a check fails (exit 1).
2. render: the same layout, for real, written to the output path.
3. validate: nbg_validate.py runs on the file in UTF-8 mode. Exit 1 on a brand
   violation, 2 when the validator could not run at all.

What PowerPoint shows is what the tests check. The earlier builder appended
bullet properties out of schema order (PowerPoint dropped every bullet), coloured
line series through a fill a line ignores (Office blue and red), inherited the
Office 2007 theme and python-pptx's 2013 metadata, and drew fixed-size boxes that
clipped or overprinted. None of that was visible to LibreOffice or to the
validator, which is why the tests now read the OOXML itself.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))

try:
    import yaml  # noqa: F401  (imported first so a missing environment says so plainly)
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError as _missing:  # pragma: no cover - exercised only without the environment
    sys.stderr.write(
        f"nbg_build.py: {_missing}. Run it through the launcher, which installs its "
        'requirements: bash "${CLAUDE_PLUGIN_ROOT}/bin/decks-py" build deck.yaml deck.pptx\n'
    )
    sys.exit(2)

import nbg_chart  # noqa: E402
import nbg_spec  # noqa: E402
import nbg_tokens  # noqa: E402
from nbg_chart import LABEL_DARK_HEX, label_text_color, set_alt_text  # noqa: E402,F401
from nbg_spec import CannotRun, Issue, Report  # noqa: E402
from nbg_text import ASCENT_EM, caps, format_number, metrics  # noqa: E402

# ---------------------------------------------------------------- brand tokens

TOKENS = nbg_tokens.load()
GEO = TOKENS["geometry"]
COMP = TOKENS["components"]
FONT = str(TOKENS["fonts"]["primary"])
FONT_BULLET = str(COMP["bullets"]["glyph_font"])
GUTTER = float(GEO["gutter"])
CONTENT_W = float(GEO["content_w"])
RIGHT = float(GEO["right_boundary"])
SLIDE_WIDTH = Emu(int(GEO["slide"]["w_emu"]))
SLIDE_HEIGHT = Emu(int(GEO["slide"]["h_emu"]))
SOURCE_H = float(GEO["source_line"]["h"])
SOURCE_Y = float(GEO["source_line"]["bottom"]) - SOURCE_H
BODY_TOP = float(GEO["body_top"])
BODY_BOTTOM = float(GEO["body_bottom"])
GRID_GAP = float(GEO["grid_gap"])
TITLE_SPACING = float(TOKENS["line_spacing"]["title"])
BODY_SPACING = float(TOKENS["line_spacing"]["body"])

ASSETS_DIR = nbg_spec.ASSETS_DIR
LOGO_PNG = ASSETS_DIR / "nbg-logo-gr.png"  # Standard #10: the Greek wordmark, always
BACK_COVER_LOGO_PNG = ASSETS_DIR / "nbg-back-cover-logo.png"

# Measured text is compared against 97% of the box: Pillow and PowerPoint agree on
# Aptos advances to well under that, and the margin keeps a fit decision on the
# safe side of whatever rounding PowerPoint does.
FIT = 0.97

LANG_TAG = {"en": "en-GB", "el": "el-GR"}
WORDS = {
    "en": {
        "contents": "Contents",
        "source": "Source",
        "as_of": "as of",
        "recommended": "Recommended",
    },
    "el": {
        "contents": "Περιεχόμενα",
        "source": "Πηγή",
        "as_of": "στοιχεία",
        "recommended": "Προτείνεται",
    },
}
SOURCE_PREFIXES = ("source:", "sources:", "πηγή:", "πηγές:")

ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
DECORATIVE_EXT_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
SLIDENUM_FIELD_ID = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"


def hexc(token: str) -> str:
    return str(nbg_tokens.color(token))


def rgb(token_or_hex: str) -> RGBColor:
    return RGBColor.from_string(nbg_tokens.color(token_or_hex))


C_WHITE = rgb("white")
C_DARK_TEAL = rgb("dark_teal")
C_TEAL = rgb("teal")
C_DARK_TEXT = rgb("body_text")
C_MEDIUM_GRAY = rgb("muted_grey")


@dataclass(frozen=True)
class Style:
    size: float
    bold: bool
    color: str  # hex


def style(role: str, **override: Any) -> Style:
    s = nbg_tokens.type_style(role)
    return Style(
        float(override.get("size", s["size"])),
        bool(override.get("bold", s["bold"])),
        str(override.get("color", s["color"])),
    )


# ---------------------------------------------------------------- errors and context


class FitError(ValueError):
    """Content that cannot be drawn faithfully. Refused rather than drawn clipped."""

    def __init__(self, path: str, message: str, fix: str = "") -> None:
        super().__init__(message)
        self.path = path
        self.message = message
        self.fix = fix


class SlideCrash(RuntimeError):
    """A slide raised something the builder did not expect: a builder defect, so the
    CLI exits 2 (could not run), naming the slide, rather than blaming the spec."""


class SpecInvalid(ValueError):
    """The spec failed its check; nothing was written."""

    def __init__(self, report: Report) -> None:
        self.report = report
        lines = [i.format() for i in report.errors]
        super().__init__("the deck spec has errors:\n" + "\n".join(lines))


@dataclass
class Frame:
    x: float
    y: float
    w: float
    h: float

    @property
    def bottom(self) -> float:
        return self.y + self.h

    def box(self) -> tuple[float, float, float, float]:
        return (self.x, self.y, self.w, self.h)


@dataclass
class Deck:
    """What every slide renderer needs: the presentation, language and location."""

    prs: Any
    lang: str = "en"
    spec_dir: Path = field(default_factory=Path.cwd)
    total: int = 0
    base: str = "slides"
    index: int = 0
    slide_spec: dict[str, Any] = field(default_factory=dict)
    warnings: list[Issue] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)
    require_title: bool = True

    def path(self, rel: str = "") -> str:
        head = f"{self.base}[{self.index}]"
        return f"{head}.{rel}" if rel else head

    def issue(self, level: str, rel: str, message: str, fix: str = "") -> Issue:
        sid = self.slide_spec.get("id")
        stype = self.slide_spec.get("type")
        return Issue(
            level,
            message,
            fix,
            self.path(rel),
            self.index + 1,
            sid if isinstance(sid, str) else None,
            stype if isinstance(stype, str) else None,
        )

    def warn(self, rel: str, message: str, fix: str = "") -> None:
        self.warnings.append(self.issue("warning", rel, message, fix))

    def error(self, rel: str, message: str, fix: str = "") -> None:
        """An error that leaves the slide drawable, so the rest of it is still checked."""
        self.errors.append(self.issue("error", rel, message, fix))

    def fit(self, rel: str, message: str, fix: str) -> FitError:
        return FitError(self.path(rel), message, fix)

    @property
    def words(self) -> dict[str, str]:
        return WORDS[self.lang]


# ---------------------------------------------------------------- text primitives


def _style_run(run: Any, s: Style, lang: str, font: str = FONT) -> None:
    f = run.font
    f.name = font
    f.size = Pt(s.size)
    f.bold = s.bold
    f.color.rgb = RGBColor.from_string(s.color)
    run._r.get_or_add_rPr().set("lang", LANG_TAG.get(lang, "en-GB"))


_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def _frame_basics(
    tf: Any, *, anchor: str = "top", wrap: bool = True, insets: tuple = (0, 0, 0, 0)
) -> None:
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(v) for v in insets)
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = _ANCHOR[anchor]


def _write(
    tf: Any,
    paragraphs: list[str] | str,
    s: Style,
    lang: str,
    *,
    align: str = "left",
    spacing: float | None = None,
) -> None:
    """Replace a text frame's text with `paragraphs`, each one styled run."""
    items = [paragraphs] if isinstance(paragraphs, str) else list(paragraphs)
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        if spacing is not None:
            p.line_spacing = spacing
        run = p.add_run()
        run.text = str(text)
        _style_run(run, s, lang)


def add_text(
    slide: Any,
    box: tuple[float, float, float, float],
    text: list[str] | str,
    s: Style,
    lang: str,
    *,
    align: str = "left",
    anchor: str = "top",
    spacing: float | None = None,
) -> Any:
    """A zero-inset, top-anchored text box (dimensions.md Text Box Rules)."""
    x, y, w, h = box
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _frame_basics(shape.text_frame, anchor=anchor)
    _write(shape.text_frame, text, s, lang, align=align, spacing=spacing)
    return shape


def lines_of(text: str, width: float, s: Style) -> list[str]:
    return list(metrics().wrap(str(text), width * FIT, s.size, s.bold))


def text_height(n_lines: int, s: Style, spacing: float = 1.0) -> float:
    return float(n_lines * metrics().line_height(s.size, spacing))


def _chars_that_fit(text: str, width: float, s: Style) -> int:
    """Roughly how many of text's characters fit on one line, for a fix message."""
    measured = metrics().width(text, s.size, s.bold)
    return max(1, int(len(text) * width * FIT / measured)) if measured else len(text)


def _strip_style(shape: Any) -> None:
    """Remove python-pptx's p:style, whose effectRef idx=2 is the theme's drop shadow.

    The brand forbids shadows on any shape (DOCS-MEMORY-8). The spPr carries an
    explicit fill and line on every shape this module draws, so nothing else is lost."""
    style_el = shape._element.find(qn("p:style"))
    if style_el is not None:
        shape._element.remove(style_el)


_SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON,
    "arrow_right": MSO_SHAPE.RIGHT_ARROW,
}


def add_shape(
    slide: Any,
    kind: str,
    box: tuple[float, float, float, float],
    *,
    fill: str | None,
    border: str | None = None,
    border_pt: float = 1.0,
    radius_in: float | None = None,
) -> Any:
    x, y, w, h = box
    shape = slide.shapes.add_shape(_SHAPES[kind], Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(shape)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if border:
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(border_pt)
    else:
        shape.line.fill.background()
    if kind == "rounded_rect" and radius_in is not None:
        # Tokens give the radius in inches; the preset takes it as a share of the
        # shorter side, capped at 0.5 (a fully round end).
        shape.adjustments[0] = max(0.0, min(0.5, radius_in / max(min(w, h), 1e-6)))
    return shape


def set_title(
    deck: Deck,
    slide: Any,
    text: str,
    box: tuple[float, float, float, float],
    s: Style,
    *,
    spacing: float = TITLE_SPACING,
) -> Any:
    """Write the slide's real title placeholder: screen readers navigate by it and the
    Accessibility Checker looks for it (BUILD-CODE-13). It is first in reading order."""
    title = slide.shapes.title
    if title is None:
        return add_text(slide, box, text, s, deck.lang, spacing=spacing)
    x, y, w, h = box
    title.left, title.top, title.width, title.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = title.text_frame
    _frame_basics(tf)
    _write(tf, text, s, deck.lang, spacing=spacing)
    return title


def _drop_title_placeholder(slide: Any) -> None:
    title = slide.shapes.title
    if title is not None:
        title._element.getparent().remove(title._element)


# ---------------------------------------------------------------- chrome


def new_slide(deck: Deck, *, titled: bool = True) -> Any:
    layout = deck.prs.slide_layouts[5 if titled else 6]  # "Title Only" / "Blank"
    slide = deck.prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C_WHITE
    return slide


def mark_decorative(picture: Any) -> None:
    """Empty alt text plus Office's decorative flag: a logo on every slide conveys
    nothing, and a screen reader announcing its file name on each one is noise."""
    nv = picture._element.find(qn("p:nvPicPr")).find(qn("p:cNvPr"))
    nv.set("descr", "")
    ext_lst = nv.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(nv, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", DECORATIVE_EXT_URI)
    flag = etree.SubElement(ext, f"{{{ADEC_NS}}}decorative", nsmap={"adec": ADEC_NS})
    flag.set("val", "1")


def add_logo(slide: Any, kind: str = "small") -> Any:
    """Width from tokens, height from the PNG's own aspect (Standard #4)."""
    pos = GEO["logo_back" if kind == "back_cover" else f"logo_{kind}"]
    path = BACK_COVER_LOGO_PNG if kind == "back_cover" else LOGO_PNG
    if not path.exists():
        raise CannotRun(f"missing brand asset {path}; the plugin install is incomplete")
    picture = slide.shapes.add_picture(
        str(path), Inches(pos["x"]), Inches(pos["y"]), width=Inches(pos["w"])
    )
    mark_decorative(picture)
    return picture


def add_page_number(deck: Deck, slide: Any) -> Any:
    """A slide-number field showing the physical position, as PowerPoint's own does,
    so the number stays right when someone inserts or moves a slide."""
    pos = GEO["page_number"]
    s = style("page_number")
    shape = slide.shapes.add_textbox(
        Inches(pos["x"]), Inches(pos["y"]), Inches(pos["w"]), Inches(pos["h"])
    )
    _frame_basics(shape.text_frame, anchor="middle", wrap=False)
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", SLIDENUM_FIELD_ID)
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", LANG_TAG[deck.lang])
    rpr.set("sz", str(int(s.size * 100)))
    rpr.set("b", "0")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", s.color)
    etree.SubElement(rpr, qn("a:latin")).set("typeface", FONT)
    etree.SubElement(fld, qn("a:t")).text = str(deck.index + 1)
    return shape


def add_pill(deck: Deck, slide: Any, text: str) -> Any:
    """The section pill, sized to its text (BUILD-CODE-6): never a fixed width."""
    comp = COMP["pill"]
    s = style("pill")
    label = caps(text, deck.lang)
    pad = float(comp["pad_x"])
    width = max(float(comp["min_w"]), metrics().width(label, s.size, True) / FIT + 2 * pad)
    x, y, h = float(comp["x"]), float(comp["y"]), float(comp["h"])
    if x + width > RIGHT:
        raise deck.fit(
            "content.bumper", f"the pill text needs {width:.2f} in", "shorten the bumper"
        )
    shape = add_shape(
        slide,
        "rounded_rect",
        (x, y, width, h),
        fill=comp["fill"],
        radius_in=float(comp["radius_in"]),
    )
    tf = shape.text_frame
    _frame_basics(tf, anchor="middle", wrap=bool(comp["word_wrap"]), insets=(pad, 0, pad, 0))
    _write(tf, label, s, deck.lang, align="center")
    return shape


def source_text(source: Any, lang: str) -> str | None:
    """The footnote under an exhibit: where the figures come from, and when they were true."""
    if not source:
        return None
    words = WORDS[lang]
    if isinstance(source, str):
        text = source.strip()
    else:
        name = str(source.get("name") or "").strip()
        as_of = str(source.get("as_of") or "").strip()
        basis = str(source.get("basis") or "").strip()
        if not name:
            return None
        text = f"{name}, {words['as_of']} {as_of}" if as_of else name
        if basis:
            text += f"; {basis}"
    if not text:
        return None
    return text if text.lower().startswith(SOURCE_PREFIXES) else f"{words['source']}: {text}"


def body_bottom(content: dict[str, Any]) -> float:
    """Where the body ends: above the takeaway strip and the source line."""
    bottom = nbg_spec.body_bottom(bool(content.get("source")))
    if content.get("takeaway"):
        strip = COMP["takeaway_strip"]
        bottom -= float(strip["h"]) + float(strip["gap_above"])
    return bottom


def draw_header(deck: Deck, slide: Any, content: dict[str, Any]) -> float:
    """Pill, action title and caption. Returns where the body starts."""
    bumper = content.get("bumper")
    if bumper:
        add_pill(deck, slide, str(bumper))
    title = content.get("title")
    title_y = float(GEO["title"]["y_with_pill" if bumper else "y"])
    top = BODY_TOP
    if title:
        s = style("title")
        lines = lines_of(str(title), CONTENT_W, s)
        if len(lines) > 2:
            raise deck.fit(
                "content.title",
                f"the title needs {len(lines)} lines at {s.size:g}pt",
                "Standard #11: one line, about 80 characters; say it shorter",
            )
        if len(lines) == 2:
            deck.warn(
                "content.title",
                "the title wraps to two lines",
                "Standard #11 wants one line; shorten it to about 80 characters",
            )
        h = max(float(GEO["title"]["h"]), text_height(len(lines), s, TITLE_SPACING) + 0.04)
        set_title(deck, slide, str(title), (GUTTER, title_y, CONTENT_W, h), s)
        top = max(BODY_TOP, title_y + h + 0.15)
    elif deck.require_title:
        raise deck.fit("content.title", "the slide has no title", "add content.title")
    else:
        _drop_title_placeholder(slide)
    description = content.get("description")
    if description:
        s = style("caption")
        lines = lines_of(str(description), CONTENT_W, s)
        h = max(float(GEO["caption"]["h"]), text_height(len(lines), s))
        add_text(slide, (GUTTER, top, CONTENT_W, h), str(description), s, deck.lang)
        top += h + float(GEO["caption"]["gap_below"])
    return top


def draw_footer(deck: Deck, slide: Any, content: dict[str, Any], *, numbered: bool = True) -> None:
    """Takeaway strip, source line, small logo and page number, drawn after the body so
    a screen reader reaches them last."""
    takeaway = content.get("takeaway")
    text = source_text(content.get("source"), deck.lang)
    if takeaway:
        strip = COMP["takeaway_strip"]
        s = Style(float(strip["size"]), bool(strip["bold"]), hexc(strip["text_color"]))
        pad = float(strip["pad_x"])
        if len(lines_of(str(takeaway), CONTENT_W - 2 * pad, s)) > 1:
            raise deck.fit(
                "content.takeaway",
                "the takeaway runs past one line",
                "cut it to one short sentence",
            )
        h = float(strip["h"])
        limit = (SOURCE_Y - 0.1) if text else BODY_BOTTOM
        shape = add_shape(slide, "rect", (GUTTER, limit - h, CONTENT_W, h), fill=strip["fill"])
        _frame_basics(shape.text_frame, anchor="middle", insets=(pad, 0, pad, 0))
        _write(shape.text_frame, str(takeaway), s, deck.lang)
    if text:
        s = style("source")
        if len(lines_of(text, CONTENT_W, s)) > 1:
            raise deck.fit(
                "content.source",
                "the source line runs past one line",
                "shorten source.name or source.basis",
            )
        add_text(slide, (GUTTER, SOURCE_Y, CONTENT_W, SOURCE_H), text, s, deck.lang)
    add_logo(slide, "small")
    if numbered:
        add_page_number(deck, slide)


def titled_frame(deck: Deck, slide: Any, content: dict[str, Any]) -> Frame:
    top = draw_header(deck, slide, content)
    bottom = body_bottom(content)
    if bottom - top < 0.6:
        raise deck.fit(
            "content",
            f"the title, caption and footer leave {max(0.0, bottom - top):.2f} in for the body",
            "drop the caption or the takeaway, or shorten the title",
        )
    return Frame(GUTTER, top, CONTENT_W, bottom - top)


# ---------------------------------------------------------------- bullets and text blocks


def _points(points: list[Any]) -> list[tuple[str, int]]:
    out = []
    for p in points:
        if isinstance(p, dict):
            out.append((str(p.get("text", "")), 2 if p.get("level") == 2 else 1))
        else:
            out.append((str(p), 1))
    return out


def _bullet_height(items: list[tuple[str, int]], width: float, size: float) -> float:
    comp = COMP["bullets"]
    indent = int(comp["indent_emu"]) / 914400
    s = Style(size, False, hexc("body_text"))
    total = 0.0
    for i, (text, level) in enumerate(items):
        n = len(lines_of(text, width - indent * level, s))
        total += text_height(n, s, BODY_SPACING)
        if i:
            total += float(comp["space_before_pt" if level == 1 else "space_before_l2_pt"]) / 72
    return total


def add_bullets(deck: Deck, slide: Any, points: list[Any], frame: Frame, rel: str) -> Any:
    """Bullets with a cyan glyph and a hanging indent, schema order in every a:pPr."""
    items = _points(points)
    body = nbg_tokens.get("type.body")
    size = float(body["size"])
    floor = float(body.get("min_size", size))
    while _bullet_height(items, frame.w, size) > frame.h and size > floor:
        size -= 1
    needed = _bullet_height(items, frame.w, size)
    if needed > frame.h:
        raise deck.fit(
            rel,
            f"{len(items)} bullet(s) need {needed:.2f} in at {size:g}pt; the body has {frame.h:.2f} in",
            "cut or shorten bullets, or split the slide (14pt is the floor, Standard #11)",
        )
    comp = COMP["bullets"]
    indent = int(comp["indent_emu"])
    s = Style(size, False, hexc("body_text"))
    shape = slide.shapes.add_textbox(
        Inches(frame.x), Inches(frame.y), Inches(frame.w), Inches(frame.h)
    )
    tf = shape.text_frame
    _frame_basics(tf)
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = BODY_SPACING
        if i:
            key = "space_before_pt" if level == 1 else "space_before_l2_pt"
            p.space_before = Pt(float(comp[key]))
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(indent * level))
        ppr.set("indent", str(-indent))
        if level == 2:
            ppr.set("lvl", "1")
        _bullet_glyph(ppr, str(comp["glyph" if level == 1 else "glyph_l2"]))
        run = p.add_run()
        run.text = text
        _style_run(run, s, deck.lang)
    return shape


def _bullet_glyph(ppr: Any, glyph: str) -> None:
    """buClr, buFont, buChar inserted in CT_TextParagraphProperties order.

    Appending them after python-pptx's own children put them out of schema order,
    and PowerPoint silently dropped every bullet (BUILD-CODE-1)."""
    comp = COMP["bullets"]
    bu_clr = etree.Element(qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", hexc(comp["glyph_color"]))
    ppr.insert_element_before(
        bu_clr,
        "a:buSzTx", "a:buSzPct", "a:buSzPts", "a:buFontTx", "a:buFont",
        "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst",
    )  # fmt: skip
    bu_font = etree.Element(qn("a:buFont"))
    bu_font.set("typeface", FONT_BULLET)
    ppr.insert_element_before(
        bu_font,
        "a:buNone",
        "a:buAutoNum",
        "a:buChar",
        "a:buBlip",
        "a:tabLst",
        "a:defRPr",
        "a:extLst",
    )
    bu_char = etree.Element(qn("a:buChar"))
    bu_char.set("char", glyph)
    ppr.insert_element_before(bu_char, "a:tabLst", "a:defRPr", "a:extLst")


def add_paragraph_block(
    deck: Deck, slide: Any, text: str, frame: Frame, rel: str, role: str = "body"
) -> Any:
    """Plain text that must fit its frame, shrinking from the role size to its floor."""
    base = nbg_tokens.get(f"type.{role}")
    size = float(base["size"])
    floor = float(base.get("min_size", size))
    s = style(role, size=size)
    while text_height(len(lines_of(text, frame.w, s)), s, BODY_SPACING) > frame.h and size > floor:
        size -= 1
        s = style(role, size=size)
    n = len(lines_of(text, frame.w, s))
    if text_height(n, s, BODY_SPACING) > frame.h:
        raise deck.fit(
            rel, f"the text needs {n} lines at {size:g}pt and does not fit", "shorten it"
        )
    return add_text(slide, frame.box(), text, s, deck.lang, spacing=BODY_SPACING)


# ---------------------------------------------------------------- images


def _image_stream(
    deck: Deck, raw: str, width_in: float, rel: str
) -> tuple[Any, tuple[int, int], bool]:
    """(PNG or JPEG stream, its pixel size, whether it came from a vector file)."""
    path = nbg_spec.resolve_asset(raw, deck.spec_dir)
    if path is None:
        raise deck.fit(rel, f"image not found: {raw}", "check the path")
    from PIL import Image

    vector = path.suffix.lower() == ".svg"
    if vector:
        try:
            import resvg_py
        except ImportError as e:
            raise deck.fit(
                rel, "SVG needs resvg-py", "run through bin/decks-py, or use a PNG"
            ) from e
    try:
        if vector:
            px = int(min(4000, max(600, width_in * 300)))
            stream: Any = io.BytesIO(bytes(resvg_py.svg_to_bytes(svg_path=str(path), width=px)))
        else:
            stream = io.BytesIO(path.read_bytes())
        with Image.open(stream) as img:
            size = img.size
    except Exception as e:  # noqa: BLE001 - a bad file is the spec's problem, not the environment's
        raise deck.fit(
            rel,
            f"'{path.name}' exists but cannot be read ({type(e).__name__}: {e})",
            "re-export it as a PNG, JPEG or valid SVG",
        ) from e
    stream.seek(0)
    return stream, size, vector


def _tinted(deck: Deck, stream: Any, colour: str, rel: str) -> Any:
    """The icon's shape in one colour: every pixel `colour`, the original alpha kept.
    Only a picture with transparency has a shape to keep; an opaque one is left as is."""
    from PIL import Image

    with Image.open(stream) as img:
        rgba = img.convert("RGBA")
    alpha = rgba.getchannel("A")
    if alpha.getextrema()[0] == 255:
        deck.warn(
            rel,
            "the icon has no transparent background, so it cannot be drawn white",
            "use an SVG or a transparent PNG",
        )
        stream.seek(0)
        return stream
    solid = Image.new("RGBA", rgba.size, "#" + colour)
    solid.putalpha(alpha)
    out = io.BytesIO()
    solid.save(out, format="PNG")
    out.seek(0)
    return out


def add_image(
    deck: Deck,
    slide: Any,
    raw: str,
    frame: Frame,
    *,
    alt: str,
    fit: str = "contain",
    rel: str = "image.path",
    tint: str | None = None,
) -> Any:
    """A picture inside frame: contain keeps all of it, cover fills the frame and crops.
    tint (a hex colour) redraws an icon's shape in that one colour. A raster is never
    enlarged past components.image.min_dpi: contain draws it at that size, centred,
    and says the image is too small for its slot; cover, which must fill, warns."""
    stream, (px_w, px_h), vector = _image_stream(deck, raw, frame.w, rel)
    if tint:
        stream = _tinted(deck, stream, tint, rel)
    aspect = px_w / px_h if px_h else 1.0
    dpi = float(COMP["image"]["min_dpi"])
    if not vector:
        needed = (
            max(frame.w, frame.h * aspect) if fit == "cover" else min(frame.w, frame.h * aspect)
        )
        if needed * dpi > px_w + 0.5:
            deck.warn(
                rel,
                f"the image is {px_w} x {px_h} px, too small for its {frame.w:.1f} x {frame.h:.1f} in "
                f"slot at {dpi:g} DPI",
                f"use an image at least {math.ceil(needed * dpi)} px wide, or an SVG",
            )
            if fit != "cover":
                frame = Frame(
                    frame.x + (frame.w - px_w / dpi) / 2,
                    frame.y + (frame.h - px_h / dpi) / 2,
                    px_w / dpi,
                    px_h / dpi,
                )
    if fit == "cover":
        picture = slide.shapes.add_picture(
            stream, Inches(frame.x), Inches(frame.y), Inches(frame.w), Inches(frame.h)
        )
        target = frame.w / frame.h
        if aspect > target:
            excess = 1 - target / aspect
            picture.crop_left = picture.crop_right = excess / 2
        elif aspect < target:
            excess = 1 - aspect / target
            picture.crop_top = picture.crop_bottom = excess / 2
    else:
        if frame.w / frame.h > aspect:
            h = frame.h
            w = h * aspect
        else:
            w = frame.w
            h = w / aspect
        x = frame.x + (frame.w - w) / 2
        y = frame.y + (frame.h - h) / 2
        picture = slide.shapes.add_picture(stream, Inches(x), Inches(y), Inches(w), Inches(h))
    set_alt_text(picture, alt)
    return picture


# ---------------------------------------------------------------- tables

_CURRENCY = r"(?:EUR|USD|GBP|CHF|€|\$|£)"
_SCALE = r"(?:bn|[KkMmBb]|εκατ\.?|δισ\.?|χιλ\.?)"
# A figure with its trimmings: sign, currency code or symbol before or after, digits
# with separators, a percent or a scale suffix, parentheses for a negative. "EUR 42M"
# and "€ 18,5 εκατ." are numbers and right-align like "+14%" (E2E-SMOKE-15).
_NUMERIC_CELL = re.compile(
    rf"^\(?[+\-−]?\s*{_CURRENCY}?\s*[+\-−]?\d[\d.,\s]*\s*(?:%|{_SCALE})?\s*{_CURRENCY}?\)?$",
    re.IGNORECASE,
)
_BLANKISH = {"", "-", "–", "n/a", "na", "n.a.", chr(0x2014)}


def is_numeric_cell(text: str) -> bool:
    return bool(_NUMERIC_CELL.match(text.strip()))


def numeric_columns(rows: list[list[str]], width: int) -> set[int]:
    out = set()
    for col in range(width):
        cells = [r[col].strip() for r in rows if col < len(r)]
        cells = [c for c in cells if c.casefold() not in _BLANKISH]
        if cells and all(is_numeric_cell(c) for c in cells):
            out.add(col)
    return out


def _is_figure(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def _decimals(value: Any) -> int:
    """How many decimals a number carries as written, at most two."""
    if not isinstance(value, float) or value.is_integer():
        return 0
    return min(2, len(f"{value:.6f}".rstrip("0").split(".")[1]))


def _cell_text(value: Any, lang: str = "en", decimals: int | None = None) -> str:
    """A cell as text. A number takes its column's precision (E2E-OUTPUT-07) and the
    deck language's separators (1.234,5 in Greek)."""
    if value is None:
        return ""
    if _is_figure(value):
        places = _decimals(value) if decimals is None else decimals
        return str(format_number(float(value), places, lang))
    return str(value)


def _cell_border(tcpr: Any) -> None:
    """White 1pt borders (tokens table.border), inserted ahead of the fill."""
    border = COMP["table"]["border"]
    width = str(int(Pt(float(border["width_pt"]))))
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        ln.set("w", width)
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hexc(border["color"]))
        tcpr.insert(i, ln)


def add_table(deck: Deck, slide: Any, spec: dict[str, Any], frame: Frame, rel: str) -> Any:
    comp = COMP["table"]
    headers = [str(h) for h in spec.get("headers") or []]
    width = max([len(headers)] + [len(r) for r in spec["rows"]])
    # One precision per column: the most decimals any figure in it carries, capped at two.
    places = [
        max((_decimals(r[c]) for r in spec["rows"] if c < len(r) and _is_figure(r[c])), default=0)
        for c in range(width)
    ]
    rows = [
        [_cell_text(v, deck.lang, places[c]) for c, v in enumerate(r)] + [""] * (width - len(r))
        for r in spec["rows"]
    ]
    headers += [""] * (width - len(headers))
    numeric = numeric_columns(rows, width)
    aligns = list(spec.get("column_align") or [])
    highlight = spec.get("highlight_column")
    inset = float(comp["cell_inset"])
    head_s = style("table_header")
    body_s = style("table_body")
    bold_s = style("table_body", bold=True)
    highlight_s = style("table_body", bold=True, color=hexc("teal"))

    def col_style(r: int, c: int) -> Style:
        if r < 0:
            return head_s
        if highlight is not None and c == highlight:
            return highlight_s
        return bold_s if c == 0 else body_s

    # Column widths from the widest text in each column. Spare width goes half to the
    # label column and half across the figures, so numbers stay near their labels;
    # too little width shrinks every column alike and the cells wrap.
    natural = []
    for c in range(width):
        texts = [(headers[c], head_s)] + [(r[c], col_style(i, c)) for i, r in enumerate(rows)]
        widest = max(metrics().width(t, s.size, s.bold) for t, s in texts)
        natural.append(widest / FIT + 2 * inset + 0.05)
    spare = frame.w - sum(natural)
    if spare >= 0:
        widths = list(natural)
        widths[0] += spare if width == 1 else spare / 2
        for c in range(1, width):
            widths[c] += spare / 2 / (width - 1)
    else:
        widths = [n * frame.w / sum(natural) for n in natural]

    def row_height(cells: list[str], r: int, minimum: float) -> float:
        n = max(
            len(lines_of(t, widths[c] - 2 * inset, col_style(r, c))) for c, t in enumerate(cells)
        )
        return max(minimum, text_height(n, head_s if r < 0 else body_s) + 0.1)

    header_h = row_height(headers, -1, float(comp["header_h"])) if any(headers) else 0.0
    body_hs = [row_height(r, i, float(comp["row_h"])) for i, r in enumerate(rows)]
    total = header_h + sum(body_hs)
    if total > frame.h + 1e-6:
        raise deck.fit(
            rel,
            f"the table needs {total:.2f} in ({len(rows)} rows); the body has {frame.h:.2f} in",
            "split the table across two slides or cut rows (12pt is the floor for cells)",
        )
    n_rows = len(rows) + (1 if header_h else 0)
    gf = slide.shapes.add_table(
        n_rows, width, Inches(frame.x), Inches(frame.y), Inches(frame.w), Inches(total)
    )
    table = gf.table
    tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
    style_id.text = NO_STYLE_TABLE
    table.first_row = False
    table.horz_banding = False
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)

    def align_of(c: int) -> str:
        explicit = aligns[c] if c < len(aligns) else "auto"
        if explicit != "auto":
            return str(explicit)
        return "right" if c in numeric else "left"

    def fill_cell(cell: Any, text: str, s: Style, fill: str, align: str) -> None:
        _cell_border(cell._tc.get_or_add_tcPr())
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.margin_left = cell.margin_right = Inches(inset)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = _ALIGN[align]
        run = p.add_run()
        run.text = text
        _style_run(run, s, deck.lang)

    r0 = 0
    if header_h:
        table.rows[0].height = Inches(header_h)
        for c in range(width):
            fill_cell(table.cell(0, c), headers[c], head_s, comp["header_fill"], align_of(c))
        r0 = 1
    for i, row in enumerate(rows):
        table.rows[r0 + i].height = Inches(body_hs[i])
        fill = comp["zebra_fill"] if i % 2 else "white"
        for c in range(width):
            fill_cell(table.cell(r0 + i, c), row[c], col_style(i, c), fill, align_of(c))
    set_alt_text(gf, spec.get("alt_text") or _table_alt_text(headers, rows, deck.lang))
    return gf


def _table_alt_text(headers: list[str], rows: list[list[str]], lang: str) -> str:
    named = ", ".join(h for h in headers if h.strip())
    if lang == "el":
        cols = f"{len(headers)} στήλες" + (f" ({named})" if named else "")
        return f"Πίνακας. {cols}. {len(rows)} γραμμές δεδομένων."
    cols = f"{len(headers)} columns" + (f" ({named})" if named else "")
    return f"Table. {cols}. {len(rows)} data rows."


# ---------------------------------------------------------------- slide types


def _zero_pad(value: Any) -> str:
    text = str(value)
    return text.zfill(2) if text.isdigit() else text


def render_cover(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    comp = COMP["cover"]
    title = str(content.get("title", ""))
    base = nbg_tokens.get("type.cover_title")
    size = float(base["size"])
    floor = float(base.get("min_size", size))
    box_w = float(comp["title"]["w"])
    s = style("cover_title", size=size)
    # Standard #13: one line. Shrink from 48 to 44pt; past that, shorten.
    while len(lines_of(title, box_w, s)) > 1 and size > floor:
        size -= 1
        s = style("cover_title", size=size)
    lines = lines_of(title, box_w, s)
    if len(lines) > 1:
        deck.error(
            "content.title",
            f"the cover title needs {len(lines)} lines even at {floor:g}pt",
            f"Standard #13: shorten it to one line, about {_chars_that_fit(title, box_w, s)} characters",
        )
    title_h = text_height(len(lines), s, TITLE_SPACING) + 0.1
    title_y = float(comp["title"]["y"])
    set_title(deck, slide, title, (GUTTER, title_y, box_w, title_h), s)
    cursor = title_y + title_h + float(comp["subtitle"]["gap_below_title"])
    subtitle = content.get("subtitle")
    if subtitle:
        ss = style("cover_subtitle")
        sw = float(comp["subtitle"]["w"])
        sub_lines = lines_of(str(subtitle), sw, ss)
        if len(sub_lines) > 1:
            deck.error(
                "content.subtitle",
                f"the subtitle needs {len(sub_lines)} lines at {ss.size:g}pt",
                "Standard #13: shorten it to one line, about "
                f"{_chars_that_fit(str(subtitle), sw, ss)} characters",
            )
        sub_h = text_height(len(sub_lines), ss) + 0.05
        add_text(slide, (GUTTER, cursor, sw, sub_h), str(subtitle), ss, deck.lang)
        cursor += sub_h
    y = cursor + 0.3
    for key, role in (("location", "cover_location"), ("date", "cover_date")):
        value = content.get(key)
        if not value:
            continue
        box = comp[key]
        ls = style(role)
        n = len(lines_of(str(value), float(box["w"]), ls))
        y = max(y, float(box["y"]))
        h = max(float(box["h"]), text_height(n, ls) + 0.05)
        add_text(slide, (GUTTER, y, float(box["w"]), h), str(value), ls, deck.lang)
        y += text_height(n, ls) + 0.15
    if y > float(GEO["logo_large"]["y"]) - 0.1:
        raise deck.fit(
            "content", "the cover text runs into the logo", "shorten the title or subtitle"
        )
    add_logo(slide, "large")
    return slide


def render_divider(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    comp = COMP["divider"]
    number = _zero_pad(content.get("number", ""))
    ns = style("divider_number")
    ts = style("divider_title")
    nbox = comp["number"]
    tbox = comp["title"]
    if metrics().width(number, ns.size) / FIT > float(nbox["w"]):
        raise deck.fit(
            "content.number", f"'{number}' is too wide for the number box", "use two digits"
        )
    tx = float(tbox["x"])
    tw = min(float(tbox["w"]), RIGHT - tx)
    title = str(content.get("title", ""))
    lines = lines_of(title, tw, ts)
    if len(lines) > 2:
        raise deck.fit("content.title", f"the section title needs {len(lines)} lines", "shorten it")
    if len(lines) == 2:
        deck.warn("content.title", "the section title wraps to two lines", "shorten it to one line")
    ny = float(nbox["y"])
    # Standard #3: one baseline. A zero-inset box puts its first baseline one ascent
    # below its top, so the smaller title sits lower by the difference in ascents.
    ty = ny + (ns.size - ts.size) * ASCENT_EM / 72
    th = max(float(tbox["h"]), text_height(len(lines), ts) + 0.05)
    set_title(deck, slide, title, (tx, ty, tw, th), ts, spacing=1.0)
    add_text(
        slide,
        (float(nbox["x"]), ny, float(nbox["w"]), float(nbox["h"])),
        number,
        ns,
        deck.lang,
        spacing=1.0,
    )
    add_logo(slide, "large")
    return slide


def render_contents(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    comp = COMP["contents"]
    header = str(content.get("title") or deck.words["contents"])
    hb = comp["header"]
    set_title(
        deck,
        slide,
        header,
        (float(hb["x"]), float(hb["y"]), float(hb["w"]), float(hb["h"])),
        style("contents_header"),
    )
    sections = content.get("sections") or []
    ns, ts, ds = style("contents_number"), style("contents_title"), style("contents_description")
    number_w = float(comp["number_w"])
    title_x = GUTTER + number_w
    title_w = min(float(comp["title_w"]), RIGHT - title_x)
    rows = []
    for i, section in enumerate(sections):
        t_lines = len(lines_of(str(section.get("title", "")), title_w, ts))
        desc = section.get("description")
        d_lines = len(lines_of(str(desc), title_w, ds)) if desc else 0
        height = text_height(t_lines, ts) + (text_height(d_lines, ds) + 0.04 if desc else 0)
        rows.append((_zero_pad(section.get("number", i + 1)), section, t_lines, d_lines, height))
    any_desc = any(r[3] for r in rows)
    step = float(comp["row_step" if any_desc else "row_step_compact"])
    tallest = max((r[4] for r in rows), default=0.0)
    step = max(step, tallest + 0.15)
    first = float(comp["first_row_y"])
    last_bottom = first + step * (len(rows) - 1) + (rows[-1][4] if rows else 0)
    if last_bottom > BODY_BOTTOM + 1e-6:
        fits = int((BODY_BOTTOM - first - tallest) // step) + 1
        raise deck.fit(
            "content.sections",
            f"{len(rows)} contents rows need {last_bottom - first:.2f} in; {fits} fit",
            "split the agenda, drop the teasers, or keep the main sections only (Standard #18)",
        )
    for i, (number, section, t_lines, d_lines, _height) in enumerate(rows):
        y = first + i * step
        add_text(slide, (GUTTER, y, number_w, text_height(1, ns) + 0.05), number, ns, deck.lang)
        th = text_height(t_lines, ts) + 0.02
        add_text(slide, (title_x, y, title_w, th), str(section.get("title", "")), ts, deck.lang)
        if section.get("description"):
            dh = text_height(d_lines, ds) + 0.02
            add_text(
                slide,
                (title_x, y + th + 0.02, title_w, dh),
                str(section["description"]),
                ds,
                deck.lang,
            )
    add_logo(slide, "small")
    if deck.total >= int(comp["page_number_from_slides"]):
        add_page_number(deck, slide)
    return slide


def render_content(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    add_bullets(deck, slide, content.get("points") or [], frame, "content.points")
    draw_footer(deck, slide, content)
    return slide


def _require_series(deck: Deck, chart: dict[str, Any], rel: str) -> None:
    data = chart.get("data") or {}
    if not data.get("series") or not data.get("categories"):
        raise deck.fit(
            f"{rel}.data",
            "the chart has no data",
            "add categories and series; a blank plot is not an exhibit",
        )
    ctype = chart.get("type")
    if ctype not in nbg_chart.XL_TYPES:
        raise deck.fit(
            f"{rel}.type",
            f"unknown chart type '{ctype}'",
            f"use one of: {', '.join(nbg_chart.XL_TYPES)}",
        )


def _bank_alt(deck: Deck, bank: str) -> str:
    name = str(nbg_chart.BANKS[bank]["name"])
    return f"Λογότυπο {name}" if deck.lang == "el" else f"{name} logo"


def _bank_logo(
    deck: Deck, slide: Any, bank: str, centre: tuple[float, float], height: float
) -> Any:
    """A bank's logo at its own aspect (Standard #4), centred on `centre`."""
    path = nbg_chart.logo_path(bank)
    if not path.exists():
        raise CannotRun(f"missing brand asset {path}; the plugin install is incomplete")
    picture = slide.shapes.add_picture(str(path), 0, 0, height=Inches(height))
    cx, cy = centre
    picture.left = Inches(cx) - picture.width // 2
    picture.top = Inches(cy - height / 2)
    set_alt_text(picture, _bank_alt(deck, bank))
    return picture


@dataclass(frozen=True)
class _LegendEntry:
    label: str
    colour: str
    bank: str | None
    width: float


def _bank_legend_rows(
    entries: list[tuple[str, str, str | None]], width: float
) -> tuple[list[list[_LegendEntry]], float]:
    """Legend entries (swatch, logo, name) packed into centred rows, and the band's height."""
    cfg = COMP["bank_logos"]
    logo_h, gap = float(cfg["h"]), float(cfg["gap"])
    swatch, entry_gap = float(cfg["swatch"]), float(cfg["entry_gap"])
    ls = style("chart_legend")
    packed: list[list[_LegendEntry]] = [[]]
    used = 0.0
    for label, colour, bank in entries:
        w = swatch + gap + metrics().width(label, ls.size, ls.bold) / FIT + 0.02
        if bank:
            w += nbg_chart.logo_width(bank, logo_h) + gap
        need = w + (entry_gap if packed[-1] else 0.0)
        if packed[-1] and used + need > width:
            packed.append([])
            used, need = 0.0, w
        packed[-1].append(_LegendEntry(label, colour, bank, w))
        used += need
    row_h = max(logo_h, text_height(1, ls))
    return packed, 2 * gap + len(packed) * row_h + (len(packed) - 1) * gap


def _draw_bank_legend(deck: Deck, slide: Any, rows: list[list[_LegendEntry]], frame: Frame) -> None:
    cfg = COMP["bank_logos"]
    logo_h, gap = float(cfg["h"]), float(cfg["gap"])
    swatch, entry_gap = float(cfg["swatch"]), float(cfg["entry_gap"])
    ls = style("chart_legend")
    th = text_height(1, ls)
    row_h = max(logo_h, th)
    y = frame.y + 2 * gap
    for row in rows:
        row_w = sum(e.width for e in row) + entry_gap * (len(row) - 1)
        x = frame.x + (frame.w - row_w) / 2
        mid = y + row_h / 2
        for entry in row:
            start = x
            add_shape(slide, "rect", (x, mid - swatch / 2, swatch, swatch), fill=entry.colour)
            x += swatch + gap
            if entry.bank:
                lw = nbg_chart.logo_width(entry.bank, logo_h)
                _bank_logo(deck, slide, entry.bank, (x + lw / 2, mid), logo_h)
                x += lw + gap
            add_text(
                slide, (x, mid - th / 2, start + entry.width - x, th), entry.label, ls, deck.lang
            )
            x = start + entry.width + entry_gap
        y += row_h + gap


def draw_chart(deck: Deck, slide: Any, chart: dict[str, Any], frame: Frame, rel: str) -> Any:
    """One native chart in frame. A peer-bank comparison (tokens.yaml banks) also gets
    its logos: under or beside the bars when the banks are a bar chart's categories,
    otherwise in a legend row of swatch, logo and name that replaces the chart's own."""
    _require_series(deck, chart, rel)
    alt = chart.get("alt_text")
    plan = nbg_spec.bank_plan(chart)
    if plan is None or chart.get("bank_logos") is False:
        return nbg_chart.add_chart(slide, chart, frame.box(), deck.lang, alt)
    mode, banks = plan
    if mode == "categories" and chart["type"] in ("bar", "bar_horizontal"):
        try:
            layout = nbg_chart.axis_layout(chart, frame.box())
        except ValueError as e:
            raise deck.fit(rel, str(e), "give the chart more room, or compare fewer banks") from e
        shape = nbg_chart.add_chart(slide, chart, frame.box(), deck.lang, alt, layout=layout)
        for bank, centre in zip(banks, layout.anchors, strict=True):
            if bank:
                _bank_logo(deck, slide, bank, centre, layout.logo_h)
        return shape
    series_colours, point_colours = nbg_chart.bank_colours(chart)
    data = chart["data"]
    if mode == "series":
        labels, colours = [str(s["name"]) for s in data["series"]], series_colours
    else:
        labels, colours = [str(c) for c in data["categories"]], point_colours or series_colours
    rows, band = _bank_legend_rows(list(zip(labels, colours, banks, strict=True)), frame.w)
    chart_frame = Frame(frame.x, frame.y, frame.w, frame.h - band)
    if chart_frame.h < 1.0:
        raise deck.fit(
            rel, "no room for the chart above its logo legend", "give the chart more room"
        )
    shape = nbg_chart.add_chart(slide, chart, chart_frame.box(), deck.lang, alt, legend=False)
    _draw_bank_legend(deck, slide, rows, Frame(frame.x, chart_frame.bottom, frame.w, band))
    return shape


def render_chart(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    draw_chart(deck, slide, spec.get("chart") or {}, frame, "chart")
    draw_footer(deck, slide, content)
    return slide


def render_waterfall(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    chart = spec.get("chart") or {}
    items = (chart.get("data") or {}).get("items") or []
    if len(items) < 2:
        raise deck.fit(
            "chart.data.items",
            "a waterfall needs at least two items",
            "add the opening total and a step",
        )
    nbg_chart.add_waterfall(slide, chart, frame.box(), deck.lang, chart.get("alt_text"))
    draw_footer(deck, slide, content)
    return slide


def render_table(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    table = spec.get("table") or {}
    if not table.get("rows"):
        raise deck.fit("table.rows", "the table has no rows", "add rows")
    add_table(deck, slide, table, frame, "table")
    draw_footer(deck, slide, content)
    return slide


def _kpi_tiles(
    deck: Deck,
    slide: Any,
    kpis: list[dict[str, Any]],
    frame: Frame,
    rel: str,
    *,
    vertical: bool = False,
) -> None:
    comp = COMP["kpi"]
    n = len(kpis)
    gap = GRID_GAP
    if vertical:
        tile_w = frame.w
        tile_h = min(float(comp["h"]), (frame.h - gap * (n - 1)) / n)
    else:
        tile_w = (frame.w - gap * (n - 1)) / n
        tile_h = min(float(comp["h"]), frame.h)
    pad = float(comp["pad"])
    inner = tile_w - 2 * pad
    value_base = nbg_tokens.get("type.kpi_value")
    ls = style("kpi_label")
    # Standard #20, parallel comparison: one value size for the row (the largest that
    # fits every tile) and the values and captions on shared lines.
    size = float(value_base["size"])
    floor = float(value_base["min_size"])
    values = [str(k["value"]) for k in kpis]
    while size > floor and any(metrics().width(v, size, True) > inner * FIT for v in values):
        size -= 1
    for i, value in enumerate(values):
        if metrics().width(value, size, True) > inner * FIT:
            raise deck.fit(
                f"{rel}[{i}].value",
                f"'{value}' does not fit its tile even at {size:g}pt",
                "shorten the value (3.3M, not 3,300,000), or use fewer tiles",
            )
    vs = style("kpi_value", size=size)
    value_h = text_height(1, vs)
    labels = [lines_of(str(k["label"]), inner, ls) for k in kpis]
    delta_h = text_height(1, style("kpi_delta"))
    stacks = [
        value_h + 0.08 + text_height(len(lines), ls) + (0.06 + delta_h if k.get("delta") else 0.0)
        for lines, k in zip(labels, kpis, strict=True)
    ]
    for i, stack in enumerate(stacks):
        if stack > tile_h - 2 * pad + 1e-6:
            raise deck.fit(
                f"{rel}[{i}].label",
                "the tile's value, label and delta do not fit",
                "shorten the label",
            )
    for i, kpi in enumerate(kpis):
        x = frame.x + (0 if vertical else i * (tile_w + gap))
        y = frame.y + (i * (tile_h + gap) if vertical else (frame.h - tile_h) / 2)
        add_shape(
            slide,
            "rounded_rect",
            (x, y, tile_w, tile_h),
            fill=comp["fill"],
            radius_in=float(comp["radius_in"]),
        )
        value = values[i]
        label = str(kpi["label"])
        label_h = text_height(len(labels[i]), ls)
        delta = kpi.get("delta")
        sentiment = kpi.get("sentiment", "neutral")
        ds = style("kpi_delta", color=hexc(comp["delta"][sentiment]))
        # A row shares the tallest tile's lines (value, caption, delta); a stacked column
        # centres each tile on its own.
        stack = stacks[i] if vertical else max(stacks)
        cy = y + (tile_h - stack) / 2
        add_text(slide, (x + pad, cy, inner, value_h), value, vs, deck.lang, align="center")
        cy += value_h + 0.08
        add_text(slide, (x + pad, cy, inner, label_h), label, ls, deck.lang, align="center")
        tallest = label_h if vertical else max(text_height(len(lines), ls) for lines in labels)
        cy += tallest + 0.06
        if delta:
            add_text(
                slide, (x + pad, cy, inner, delta_h), str(delta), ds, deck.lang, align="center"
            )


def render_kpi(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    _kpi_tiles(deck, slide, spec["kpis"], frame, "kpis")
    draw_footer(deck, slide, content)
    return slide


def _card_grid(n: int, layout: str | None) -> tuple[int, int]:
    if layout == "row" or (layout is None and n <= 4):
        return n, 1
    cols = 2 if n == 4 else 3
    return cols, math.ceil(n / cols)


def _icon_alt(deck: Deck, title: str) -> str:
    return f"Εικονίδιο: {title}" if deck.lang == "el" else f"{title} icon"


def _recommended_tab(deck: Deck, slide: Any, x: float, y: float) -> None:
    tab = COMP["card"]["recommended_tab"]
    label = caps(deck.words["recommended"], deck.lang)
    s = style("recommended_tab")
    pad = float(tab["pad_x"])
    h = float(tab["h"])
    w = metrics().width(label, s.size, True) / FIT + 2 * pad
    shape = add_shape(
        slide, "rounded_rect", (x, y - h / 2, w, h), fill=tab["fill"], radius_in=h / 2
    )
    _frame_basics(shape.text_frame, anchor="middle", wrap=False, insets=(pad, 0, pad, 0))
    _write(shape.text_frame, label, s, deck.lang, align="center")


def _card_body(
    deck: Deck, slide: Any, text: str, box: tuple[float, float, float, float], rel: str, what: str
) -> None:
    x, cursor, inner, room = box
    bs = style("card_body")
    needed = text_height(len(lines_of(text, inner, bs)), bs, BODY_SPACING)
    if needed > room + 1e-6:
        raise deck.fit(
            rel,
            f"{what} needs {needed:.2f} in; it has {max(room, 0):.2f} in",
            "shorten it, use fewer items, or change the layout (14pt is the floor)",
        )
    add_text(slide, (x, cursor, inner, room), text, bs, deck.lang, spacing=BODY_SPACING)


def render_cards(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    cards = spec["cards"]
    comp = COMP["card"]
    cols, rows = _card_grid(len(cards), spec.get("layout"))
    has_tab = any(c.get("recommended") for c in cards)
    offset = float(comp["recommended_tab"]["h"]) / 2 if has_tab else 0.0
    gap = GRID_GAP
    card_w = (frame.w - gap * (cols - 1)) / cols
    pad = float(comp["pad"])
    inner = card_w - 2 * pad
    ts, bs = style("card_title"), style("card_body")
    inner_gap = float(comp["gap"])

    def content_h(card: dict[str, Any]) -> float:
        h = 2 * pad + text_height(len(lines_of(str(card["title"]), inner, ts)), ts)
        if card.get("number") is not None:
            h += float(COMP["badge"]["size_in"]) + inner_gap
        elif card.get("icon"):
            h += float(comp["icon"]) + inner_gap
        if card.get("body"):
            body = lines_of(str(card["body"]), inner, bs)
            h += inner_gap + text_height(len(body), bs, BODY_SPACING)
        return h

    # One height for every card, from the fullest one (never under min_h), and the
    # grid centred in the body: short text no longer sits in a stretched card.
    room = (frame.h - offset - gap * (rows - 1)) / rows
    needs = [content_h(c) for c in cards]
    card_h = max(float(comp["min_h"]), max(needs))
    if card_h > room + 1e-6:
        worst = needs.index(max(needs))
        raise deck.fit(
            f"cards[{worst}].body",
            f"card {worst + 1} needs {max(needs):.2f} in; each card has {room:.2f} in",
            "shorten it, use fewer cards, or change the layout (14pt is the floor)",
        )
    top = frame.y + (frame.h - (offset + rows * card_h + gap * (rows - 1))) / 2
    for i, card in enumerate(cards):
        r, c = divmod(i, cols)
        x = frame.x + c * (card_w + gap)
        y = top + offset + r * (card_h + gap)
        fill, border, border_pt = comp["fill"], None, 1.0
        if card.get("recommended"):
            rec = comp["recommended"]
            fill, border, border_pt = rec["fill"], rec["border"], float(rec["border_pt"])
        elif card.get("highlight"):
            fill = comp["highlight_fill"]
        add_shape(
            slide,
            "rounded_rect",
            (x, y, card_w, card_h),
            fill=fill,
            border=border,
            border_pt=border_pt,
            radius_in=float(comp["radius_in"]),
        )
        cursor = y + pad
        if card.get("number") is not None:
            badge = COMP["badge"]
            size = float(badge["size_in"])
            shape = add_shape(slide, "oval", (x + pad, cursor, size, size), fill=badge["fill"])
            _frame_basics(shape.text_frame, anchor="middle", wrap=False)
            _write(
                shape.text_frame,
                str(card["number"]),
                style("badge_number"),
                deck.lang,
                align="center",
            )
            cursor += size + float(comp["gap"])
        elif card.get("icon"):
            icon = float(comp["icon"])
            add_image(
                deck,
                slide,
                str(card["icon"]),
                Frame(x + pad, cursor, icon, icon),
                alt=_icon_alt(deck, str(card["title"])),
                rel=f"cards[{i}].icon",
            )
            cursor += icon + float(comp["gap"])
        th = text_height(len(lines_of(str(card["title"]), inner, ts)), ts)
        add_text(slide, (x + pad, cursor, inner, th + 0.02), str(card["title"]), ts, deck.lang)
        cursor += th + float(comp["gap"])
        if card.get("body"):
            room = y + card_h - pad - cursor
            _card_body(
                deck,
                slide,
                str(card["body"]),
                (x + pad, cursor, inner, room),
                f"cards[{i}].body",
                f"card {i + 1}'s body",
            )
        if card.get("recommended"):
            _recommended_tab(deck, slide, x + pad, y)
    draw_footer(deck, slide, content)
    return slide


def render_process(deck: Deck, spec: dict[str, Any]) -> Any:
    """Standard #20 process flow: a teal rounded-square tile per step holding its
    number (or its icon, drawn white), the title and one line under the tile, and a
    grey arrow from each tile to the next. The row is centred in the body."""
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    steps = spec["steps"]
    comp = COMP["process"]
    tile_cfg, arrow = comp["tile"], comp["arrow"]
    n = len(steps)
    lane = float(arrow["w"]) + 2 * float(comp["gap"])
    col_w = (frame.w - lane * (n - 1)) / n
    tile = min(float(tile_cfg["size"]), col_w)
    ns = style("step_number", color=hexc(tile_cfg["text_color"]))
    ts, bs = style("card_title"), style("card_body")
    titles = [lines_of(str(s["title"]), col_w, ts) for s in steps]
    bodies = [lines_of(str(s["body"]), col_w, bs) if s.get("body") else [] for s in steps]
    title_h = text_height(max(len(t) for t in titles), ts)
    body_h = max(text_height(len(b), bs, BODY_SPACING) for b in bodies)
    block = tile + float(comp["label_gap"]) + title_h
    if body_h:
        block += float(comp["text_gap"]) + body_h
    if block > frame.h + 1e-6:
        raise deck.fit(
            "steps",
            f"the steps need {block:.2f} in and the body has {frame.h:.2f} in",
            "shorten the step text, or use fewer steps",
        )
    top = frame.y + (frame.h - block) / 2
    for i, step in enumerate(steps):
        col_x = frame.x + i * (col_w + lane)
        tile_x = col_x + (col_w - tile) / 2
        shape = add_shape(
            slide,
            "rounded_rect",
            (tile_x, top, tile, tile),
            fill=tile_cfg["fill"],
            radius_in=float(tile_cfg["radius_in"]),
        )
        if step.get("icon"):
            icon = float(tile_cfg["icon"])
            add_image(
                deck,
                slide,
                str(step["icon"]),
                Frame(tile_x + (tile - icon) / 2, top + (tile - icon) / 2, icon, icon),
                alt=_icon_alt(deck, str(step["title"])),
                rel=f"steps[{i}].icon",
                tint=hexc(tile_cfg["text_color"]),
            )
        else:
            _frame_basics(shape.text_frame, anchor="middle", wrap=False)
            _write(shape.text_frame, str(i + 1).zfill(2), ns, deck.lang, align="center")
        cursor = top + tile + float(comp["label_gap"])
        th = text_height(len(titles[i]), ts)
        add_text(
            slide,
            (col_x, cursor, col_w, th + 0.02),
            str(step["title"]),
            ts,
            deck.lang,
            align="center",
        )
        if bodies[i]:
            cursor += title_h + float(comp["text_gap"])
            bh = text_height(len(bodies[i]), bs, BODY_SPACING)
            add_text(
                slide,
                (col_x, cursor, col_w, bh + 0.02),
                str(step["body"]),
                bs,
                deck.lang,
                align="center",
                spacing=BODY_SPACING,
            )
        if i < n - 1:
            gap_left = tile_x + tile
            gap_right = col_x + col_w + lane + (col_w - tile) / 2
            ax = (gap_left + gap_right) / 2 - float(arrow["w"]) / 2
            ay = top + tile / 2 - float(arrow["h"]) / 2
            add_shape(
                slide,
                "arrow_right",
                (ax, ay, float(arrow["w"]), float(arrow["h"])),
                fill=arrow["color"],
            )
    draw_footer(deck, slide, content)
    return slide


_SPLITS = {"50/50": 0.5, "40/60": 0.4, "60/40": 0.6}


def _image_block(deck: Deck, slide: Any, image: dict[str, Any], frame: Frame, rel: str) -> None:
    caption = image.get("caption")
    fit = str(image.get("fit", "contain"))
    alt = str(image["alt_text"])
    if caption:
        cs = style("caption")
        ch = text_height(len(lines_of(str(caption), frame.w, cs)), cs) + 0.02
        gap = float(COMP["image"]["caption_gap"])
        picture_frame = Frame(frame.x, frame.y, frame.w, frame.h - ch - gap)
        add_image(
            deck, slide, str(image["path"]), picture_frame, alt=alt, fit=fit, rel=f"{rel}.path"
        )
        add_text(slide, (frame.x, frame.bottom - ch, frame.w, ch), str(caption), cs, deck.lang)
    else:
        add_image(deck, slide, str(image["path"]), frame, alt=alt, fit=fit, rel=f"{rel}.path")


def _column(deck: Deck, slide: Any, column: dict[str, Any], frame: Frame, rel: str) -> None:
    heading = column.get("heading")
    if heading:
        hs = style("card_title")
        hh = text_height(len(lines_of(str(heading), frame.w, hs)), hs)
        add_text(slide, (frame.x, frame.y, frame.w, hh + 0.02), str(heading), hs, deck.lang)
        gap = hh + float(COMP["two_column"]["heading_gap"])
        frame = Frame(frame.x, frame.y + gap, frame.w, frame.h - gap)
    kind = column["kind"]
    if kind == "bullets":
        add_bullets(deck, slide, column["points"], frame, f"{rel}.points")
    elif kind == "text":
        add_paragraph_block(deck, slide, str(column["text"]), frame, f"{rel}.text")
    elif kind == "chart":
        draw_chart(deck, slide, column["chart"], frame, f"{rel}.chart")
    elif kind == "table":
        add_table(deck, slide, column["table"], frame, f"{rel}.table")
    elif kind == "image":
        _image_block(deck, slide, column["image"], frame, f"{rel}.image")
    elif kind == "kpis":
        _kpi_tiles(deck, slide, column["kpis"], frame, f"{rel}.kpis", vertical=True)


def render_two_column(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    share = _SPLITS[spec.get("split", "50/50")]
    gap = float(COMP["two_column"]["gap"])
    left_w = (frame.w - gap) * share
    right_w = frame.w - gap - left_w
    _column(deck, slide, spec["left"], Frame(frame.x, frame.y, left_w, frame.h), "left")
    _column(
        deck,
        slide,
        spec["right"],
        Frame(frame.x + left_w + gap, frame.y, right_w, frame.h),
        "right",
    )
    draw_footer(deck, slide, content)
    return slide


def render_image(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    frame = titled_frame(deck, slide, content)
    _image_block(deck, slide, spec["image"], frame, "image")
    draw_footer(deck, slide, content)
    return slide


def _element(deck: Deck, slide: Any, el: dict[str, Any], rel: str) -> None:
    kind = el["kind"]
    frame = Frame(float(el["x"]), float(el["y"]), float(el["w"]), float(el["h"]))
    role = el.get("role", "body")
    base = nbg_tokens.type_style(role)
    colour = hexc(el["text_color"]) if el.get("text_color") else str(base["color"])
    s = Style(float(el.get("size", base["size"])), bool(el.get("bold", base["bold"])), colour)
    spacing = BODY_SPACING if role in ("body", "body_dense", "card_body") else 1.0
    if kind == "text":
        n = len(lines_of(str(el["text"]), frame.w, s))
        if text_height(n, s, spacing) > frame.h + 1e-6:
            raise deck.fit(
                f"{rel}.text",
                f"the text needs {n} line(s) and does not fit its box",
                "enlarge the box or shorten the text",
            )
        add_text(
            slide,
            frame.box(),
            str(el["text"]),
            s,
            deck.lang,
            align=el.get("align", "left"),
            anchor=el.get("anchor", "top"),
            spacing=spacing,
        )
    elif kind == "bullets":
        add_bullets(deck, slide, el["points"], frame, f"{rel}.points")
    elif kind == "shape":
        fill = el.get("fill")
        shape = add_shape(
            slide,
            el["shape"],
            frame.box(),
            fill=fill,
            border=el.get("border"),
            border_pt=float(el.get("border_pt", 1)),
            radius_in=float(COMP["card"]["radius_in"]),
        )
        if el.get("text"):
            if fill and not el.get("text_color"):
                on_fill = str(label_text_color(hexc(fill)))
                colour = hexc("white") if on_fill == "FFFFFF" else hexc("body_text")
            _frame_basics(
                shape.text_frame, anchor=el.get("anchor", "middle"), insets=(0.1, 0.05, 0.1, 0.05)
            )
            _write(
                shape.text_frame,
                str(el["text"]),
                Style(s.size, s.bold, colour),
                deck.lang,
                align=el.get("align", "center"),
            )
    elif kind == "image":
        add_image(deck, slide, str(el["path"]), frame, alt=str(el["alt_text"]), rel=f"{rel}.path")
    elif kind == "chart":
        draw_chart(deck, slide, el["chart"], frame, f"{rel}.chart")
    elif kind == "table":
        add_table(deck, slide, el["table"], frame, f"{rel}.table")
    elif kind == "line":
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.x + frame.w),
            Inches(frame.y + frame.h),
        )
        line.line.color.rgb = rgb(el.get("border", "light_grey"))
        line.line.width = Pt(float(el.get("border_pt", 1)))


def render_custom(deck: Deck, spec: dict[str, Any]) -> Any:
    slide = new_slide(deck)
    content = spec.get("content") or {}
    top = draw_header(deck, slide, content)
    bottom = body_bottom(content)
    above = "the title and its caption" if content.get("description") else "the title"
    below = "the takeaway strip" if content.get("takeaway") else "the source line"
    slack = 0.02  # rounding in the header geometry; an overprint is tenths of an inch
    for e, el in enumerate(spec["elements"]):
        # E2E-OUTPUT-04: the body a custom slide really has, not the fixed 1.3-6.5 in.
        y, h = float(el["y"]), float(el["h"])
        if y < top - slack:
            deck.error(
                f"elements[{e}]",
                f"the element starts at y {y:.2f} in, under {above}, which end at {top:.2f} in",
                f"move it to y {top:.2f} or lower",
            )
        if y + h > bottom + slack:
            deck.error(
                f"elements[{e}]",
                f"the element ends at {y + h:.2f} in, over {below}, which starts at {bottom:.2f} in",
                f"end it by {bottom:.2f} in",
            )
        _element(deck, slide, el, f"elements[{e}]")
    draw_footer(deck, slide, content)
    return slide


def render_back_cover(deck: Deck, spec: dict[str, Any]) -> Any:
    """Standard #19: the centred oval emblem and nothing else."""
    slide = new_slide(deck, titled=False)
    add_logo(slide, "back_cover")
    return slide


RENDERERS = {
    "cover": render_cover,
    "contents": render_contents,
    "divider": render_divider,
    "content": render_content,
    "chart": render_chart,
    "waterfall": render_waterfall,
    "table": render_table,
    "kpi": render_kpi,
    "cards": render_cards,
    "process": render_process,
    "two_column": render_two_column,
    "image": render_image,
    "custom": render_custom,
    "back_cover": render_back_cover,
}


# ---------------------------------------------------------------- the package


def _theme(prs: Any) -> None:
    """Rewrite theme1.xml: NBG colours and Aptos, no effect styles.

    Anything that inherits from the theme (a text box or chart a colleague adds in
    PowerPoint, a chart series with no explicit colour) follows it (BRAND-SSOT-6)."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "NBG")
    palette = nbg_tokens.chart_palette()
    links = TOKENS["extended_palettes"]["links"]
    slots = {
        "dk1": hexc("body_text"),
        "lt1": hexc("white"),
        "dk2": hexc("dark_teal"),
        "lt2": hexc("off_white"),
        **{f"accent{i + 1}": palette[i] for i in range(6)},
        "hlink": str(links["hyperlink"]).upper(),
        "folHlink": str(links["followed"]).upper(),
    }
    for name, value in slots.items():
        slot = scheme.find(f"a:{name}", ns)
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, qn("a:srgbClr")).set("val", value)
    fonts = root.find(".//a:fontScheme", ns)
    fonts.set("name", "NBG")
    for which in ("a:majorFont", "a:minorFont"):
        fonts.find(which, ns).find("a:latin", ns).set("typeface", FONT)
    for effect in root.findall(".//a:effectStyleLst/a:effectStyle", ns):
        for child in list(effect):
            effect.remove(child)
        etree.SubElement(effect, qn("a:effectLst"))
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _widen_layouts(prs: Any, from_width: int) -> None:
    """The python-pptx template's master and layouts are laid out for 4:3; stretch
    every placeholder to the 16:9 slide so a slide added in PowerPoint lands right."""
    k = int(GEO["slide"]["w_emu"]) / from_width
    for part in [prs.slide_master, *prs.slide_layouts]:
        for xfrm in part._element.iter(qn("a:xfrm")):
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            if off is not None:
                off.set("x", str(round(int(off.get("x")) * k)))
            if ext is not None:
                ext.set("cx", str(round(int(ext.get("cx")) * k)))


def new_presentation() -> Any:
    """A blank Presentation at 16:9 with the NBG theme: what every build starts from."""
    prs = Presentation()
    original = int(prs.slide_width)
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    size = prs.part._element.find(qn("p:sldSz"))
    if size is not None and "type" in size.attrib:
        del size.attrib["type"]
    if original != int(SLIDE_WIDTH):
        _widen_layouts(prs, original)
    _theme(prs)
    return prs


def _core_properties(prs: Any, spec: dict[str, Any], lang: str) -> None:
    """Title and author from the spec, dates of this build; none of python-pptx's
    "Steve Canny, 2013, generated using python-pptx" survives into a board pack."""
    raw = spec.get("presentation")
    pres: dict[str, Any] = raw if isinstance(raw, dict) else {}
    slides = nbg_spec.slide_list(spec)
    cover = next((s for s in slides if isinstance(s, dict) and s.get("type") == "cover"), {})
    title = pres.get("title") or (cover.get("content") or {}).get("title") or ""
    author = str(pres.get("author") or "")
    props = prs.core_properties
    props.title = str(title)
    props.subject = str(pres.get("purpose") or "")
    props.author = author
    props.last_modified_by = author
    props.comments = ""
    props.keywords = ""
    props.category = ""
    props.revision = 1
    props.language = LANG_TAG[lang]
    now = datetime.now(UTC).replace(microsecond=0, tzinfo=None)
    props.created = now
    props.modified = now
    package = prs.part.package
    for rel_id, rel in list(package._rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            package.drop_rel(rel_id)


def render(
    spec: dict[str, Any], spec_dir: Path, *, skip: set[int] | None = None
) -> tuple[Any, list[Issue], list[Issue]]:
    """Lay out every slide. Returns (presentation, errors, warnings); a slide whose
    content cannot fit becomes an error naming it, and the other slides still render."""
    lang = nbg_spec.language(spec)
    prs = new_presentation()
    slides = nbg_spec.slide_list(spec)
    deck = Deck(
        prs=prs, lang=lang, spec_dir=spec_dir, total=len(slides), base=nbg_spec.slides_path(spec)
    )
    errors: list[Issue] = []
    for index, slide_spec in enumerate(slides):
        if skip and (index + 1) in skip:
            continue
        deck.index = index
        deck.slide_spec = slide_spec
        renderer = RENDERERS.get(slide_spec.get("type"))
        if renderer is None:
            errors.append(
                deck.issue(
                    "error",
                    "type",
                    f"unknown slide type '{slide_spec.get('type')}'",
                    "run decks-py check",
                )
            )
            continue
        sid = slide_spec.get("id")
        try:
            slide = renderer(deck, slide_spec)
        except FitError as e:
            errors.append(
                Issue(
                    "error",
                    e.message,
                    e.fix,
                    e.path,
                    index + 1,
                    sid if isinstance(sid, str) else None,
                    slide_spec.get("type"),
                )
            )
            continue
        except (CannotRun, SlideCrash):
            raise
        except Exception as e:  # noqa: BLE001 - anything else is a builder defect on this slide
            label = ", ".join(str(p) for p in (sid, slide_spec.get("type")) if p)
            raise SlideCrash(
                f"slide {index + 1} ({label}): the builder failed on this slide with "
                f"{type(e).__name__}: {e}. The spec passed its schema, so this is a builder "
                "defect: report it with the spec"
            ) from e
        notes = slide_spec.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)
    _core_properties(prs, spec, lang)
    return prs, sorted(errors + deck.errors, key=lambda i: i.slide or 0), deck.warnings


def check(spec_path: Path | str) -> Report:
    """The full check: nbg_spec's checks, then a dry-run layout of every clean slide."""
    report = nbg_spec.check_file(spec_path)
    if report.spec is None or any(i.slide is None for i in report.errors):
        return report
    _, errors, warnings = render(report.spec, report.spec_path.parent, skip=report.error_slides())
    report.issues += errors + warnings
    return report


def _run_validator(pptx_path: Path) -> subprocess.CompletedProcess:
    validate_script = HERE / "nbg_validate.py"
    if not validate_script.exists():
        raise RuntimeError(f"validator missing: {validate_script}; the deck is UNVALIDATED")
    # -X utf8 and an explicit encoding: on Windows the validator's Greek and its tick
    # marks hit a cp1252 pipe, crashed, and every build reported UNVALIDATED.
    return subprocess.run(
        [sys.executable, "-X", "utf8", str(validate_script), str(pptx_path)],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def build_presentation(
    outline_path: Path | str, output_path: Path | str, *, validate: bool = True
) -> Report:
    """Check, render and validate. Raises SpecInvalid (exit 1) before writing anything
    when the spec fails its check, ValueError (exit 1) on a brand violation and
    RuntimeError (exit 2) when the validator could not run."""
    output_path = Path(output_path).expanduser().resolve()
    report = check(outline_path)
    for warning in report.warnings:
        print(warning.format())
    if not report.ok or report.spec is None:
        raise SpecInvalid(report)
    prs, errors, _ = render(report.spec, report.spec_path.parent)
    if errors:  # the dry run and the real render lay out identically; belt and braces
        report.issues += errors
        raise SpecInvalid(report)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"Built {len(prs.slides)} slide(s): {output_path}")
    if not validate:
        return report
    proc = _run_validator(output_path)
    print(proc.stdout, end="")
    if proc.stderr:
        print(proc.stderr, end="", file=sys.stderr)
    # nbg_validate exits 0 clean, 1 when a check failed, 2 when it could not finish.
    # An import-time crash never reaches its own handler and exits 1 with a
    # traceback; that is "could not run", not a failed check.
    crashed = proc.returncode not in (0, 1) or "Traceback (most recent call last)" in (
        proc.stderr or ""
    )
    if crashed:
        raise RuntimeError(
            f"nbg_validate.py could not run (exit {proc.returncode}); "
            f"the deck at {output_path} is UNVALIDATED"
        )
    if proc.returncode != 0:
        raise ValueError(f"nbg_validate.py found brand violations in {output_path}")
    print("Validation passed.")
    return report


# ---------------------------------------------------------------- helpers for direct callers
#
# Tests and scripts that assemble a deck slide by slide (rather than from a spec)
# call these. They render exactly as a spec would, in English, and leave the title
# out when none is given so a test can build a deliberately untitled slide.


def _direct(prs: Any, spec: dict[str, Any]) -> Any:
    deck = Deck(prs=prs, total=len(prs.slides) + 1, require_title=False)
    deck.index = len(prs.slides)
    deck.slide_spec = spec
    return RENDERERS[spec["type"]](deck, spec)


def create_cover_slide(prs: Any, content: dict[str, Any]) -> Any:
    return _direct(prs, {"type": "cover", "content": content})


def create_divider_slide(prs: Any, content: dict[str, Any]) -> Any:
    return _direct(prs, {"type": "divider", "content": content})


def create_contents_slide(
    prs: Any, sections: list[dict[str, Any]], page_number: int | None = None
) -> Any:
    return _direct(prs, {"type": "contents", "content": {"sections": sections}})


def create_content_slide(prs: Any, content: dict[str, Any], page_number: int | None = None) -> Any:
    return _direct(prs, {"type": "content", "content": content})


def create_chart_slide(
    prs: Any,
    content: dict[str, Any],
    page_number: int | None = None,
    chart_type: str = "bar",
    categories: list[Any] | None = None,
    series_list: list[dict[str, Any]] | None = None,
) -> Any:
    ctype = nbg_spec.resolve_chart_type(chart_type) or chart_type
    chart = {
        "type": ctype,
        "data": {"categories": list(categories or []), "series": list(series_list or [])},
    }
    return _direct(prs, {"type": "chart", "content": content, "chart": chart})


def create_table_slide(
    prs: Any, content: dict[str, Any], table_spec: dict[str, Any], page_number: int | None = None
) -> Any:
    return _direct(prs, {"type": "table", "content": content, "table": table_spec})


def create_waterfall_slide(
    prs: Any, content: dict[str, Any], page_number: int | None = None
) -> Any:
    content = dict(content)
    items = content.pop("waterfall_items", [])
    chart = {"type": "waterfall", "data": {"items": items}}
    return _direct(prs, {"type": "waterfall", "content": content, "chart": chart})


def create_back_cover_slide(prs: Any) -> Any:
    return _direct(prs, {"type": "back_cover"})


def _add_textbox(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    font_size: float = 14,
    color: RGBColor | None = None,
    bold: bool = False,
    align: Any = None,
    font_name: str = FONT,
) -> Any:
    """A zero-inset text box with one styled run (kept for direct callers and tests)."""
    colour = str(color) if color is not None else hexc("body_text")
    shape = add_text(slide, (x, y, w, h), text, Style(float(font_size), bool(bold), colour), "en")
    if align is not None:
        shape.text_frame.paragraphs[0].alignment = align
    if font_name != FONT:
        shape.text_frame.paragraphs[0].runs[0].font.name = font_name
    return shape


_label_text_color = label_text_color


# ---------------------------------------------------------------- command line


def _print_report(report: Report, fmt: str, stream: Any) -> None:
    if fmt == "json":
        stream.write(json.dumps(report.as_dict(), ensure_ascii=False, indent=2) + "\n")
    else:
        stream.write(report.format_text() + "\n")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="nbg_build.py",
        description="Build an NBG deck from a deck spec (deck.schema.json), or check one.",
    )
    parser.add_argument("spec", help="deck spec (YAML)")
    parser.add_argument("output", nargs="?", help="output .pptx path")
    parser.add_argument("-o", "--output-file", dest="output_file", help="output .pptx path")
    parser.add_argument("--check", action="store_true", help="check the spec only; write nothing")
    parser.add_argument(
        "--format", choices=("text", "json"), default="text", help="check output format"
    )
    args = parser.parse_args(argv)

    if args.check:
        try:
            report = check(args.spec)
        except (CannotRun, RuntimeError) as e:
            print(f"nbg_build.py could not run: {e}", file=sys.stderr)
            sys.exit(2)
        except Exception as e:  # noqa: BLE001 - a crash is "could not run", never "spec wrong"
            print(f"nbg_build.py could not run: {type(e).__name__}: {e}", file=sys.stderr)
            sys.exit(2)
        _print_report(report, args.format, sys.stdout)
        sys.exit(0 if report.ok else 1)

    if args.output_file and args.output:
        parser.error(f"two output paths given: {args.output!r} and {args.output_file!r}")
    output = args.output_file or args.output
    if not output:
        parser.error("an output path is required, as the second positional or with -o")
    try:
        build_presentation(args.spec, output)
    except SpecInvalid as e:
        _print_report(e.report, "text", sys.stderr)
        sys.exit(1)
    except (CannotRun, RuntimeError, OSError) as e:
        print(f"\nnbg_build.py could not run: {e}", file=sys.stderr)
        sys.exit(2)
    except ValueError as e:
        print(f"\nError: {e}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:  # noqa: BLE001 - a crash is "could not run", never "spec wrong"
        print(f"\nnbg_build.py could not run: {type(e).__name__}: {e}", file=sys.stderr)
        sys.exit(2)
    sys.exit(0)


if __name__ == "__main__":
    main()
