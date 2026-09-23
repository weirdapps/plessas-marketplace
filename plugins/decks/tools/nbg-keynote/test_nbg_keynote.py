#!/usr/bin/env python3
"""Tests for the NBG keynote compositor.

Runs in CI under `pytest plugins`. Locally, without the dependencies installed:
    uv run --no-project --with-requirements requirements.txt --with pytest \
        python -m pytest test_nbg_keynote.py -q
"""

import copy
import importlib
import io
import re
import shutil
import sys
from pathlib import Path

import nbg_keynote as k
import pytest
import yaml
from PIL import Image, JpegImagePlugin
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).parent
LOGO = HERE.parents[1] / "assets" / "nbg-logo-gr.png"

# ---------------------------------------------------------------- greek_upper


@pytest.mark.parametrize(
    "given,expected",
    [
        ("η συναίνεση", "Η ΣΥΝΑΙΝΕΣΗ"),  # tonos dropped
        ("το πρόβλημα", "ΤΟ ΠΡΟΒΛΗΜΑ"),
        ("Άγιος Ιωάννης", "ΑΓΙΟΣ ΙΩΑΝΝΗΣ"),
        ("ευθύνη", "ΕΥΘΥΝΗ"),
        ("the consensus", "THE CONSENSUS"),  # Latin untouched
        ("ΕΤΕ και NBG", "ΕΤΕ ΚΑΙ NBG"),  # acronyms survive
    ],
)
def test_greek_upper(given, expected):
    assert k.greek_upper(given) == expected


def test_greek_upper_keeps_dialytika():
    # The tonos goes, the dialytika stays: they mean different things.
    assert k.greek_upper("προϊόν") == "ΠΡΟΪΟΝ"


def test_plain_upper_would_have_been_wrong():
    """The bug this function exists to prevent."""
    assert "η συναίνεση".upper() == "Η ΣΥΝΑΊΝΕΣΗ"
    assert k.greek_upper("η συναίνεση") != "η συναίνεση".upper()


# ------------------------------------------------------------------ contrast


def test_luminance_endpoints():
    assert k.luminance((255, 255, 255)) == pytest.approx(1.0, abs=1e-6)
    assert k.luminance((0, 0, 0)) == pytest.approx(0.0, abs=1e-6)


def test_contrast_ratio_white_on_black():
    assert k.contrast_ratio((255, 255, 255), 0.0) == pytest.approx(21.0, abs=0.01)


def test_ground_gradient_clears_accent_text():
    """The house gradient must carry cyan body text without any scrim help."""
    assert k.contrast_ratio(k.ACCENT, k.luminance(k.GROUND_BOT)) >= 4.5


def test_min_ratio_threshold_is_18pt():
    # 2560px spans 960pt, so 18pt == 48px.
    assert k._min_ratio(47) == 4.5
    assert k._min_ratio(48) == 3.0


def test_ensure_contrast_darkens_a_bright_patch():
    canvas = Image.new("RGBA", (k.W, k.H), (235, 235, 235, 255))
    box = (400, 400, 1200, 500)
    before = k._bg_luminance(canvas, box)
    k.ensure_contrast(canvas, box, k.INK, 4.5, "test")
    after = k._bg_luminance(canvas, box)
    assert after < before
    assert k.contrast_ratio(k.INK, after) >= 4.5


def test_ensure_contrast_leaves_a_dark_ground_alone():
    canvas = k.gradient_bg()
    box = (400, 400, 1200, 500)
    before = k._bg_luminance(canvas, box)
    k.ensure_contrast(canvas, box, k.INK, 4.5, "test")
    assert k._bg_luminance(canvas, box) == pytest.approx(before, abs=1e-9)


# ---------------------------------------------------------------- bar geometry


def _bar_rects(cats, vals, **kw):
    canvas = k.gradient_bg()
    base = canvas.copy()
    k.draw_bars(canvas, cats, vals, **kw)
    # columns where the render differs from the clean ground = bar coverage
    import numpy as np

    diff = np.abs(
        np.array(canvas.convert("RGB"), dtype=int) - np.array(base.convert("RGB"), dtype=int)
    ).sum(axis=2)
    # Sample just above the baseline (1170) so even the shortest bar is crossed.
    cols = np.where(diff[1150, :] > 30)[0]
    runs, start = [], cols[0]
    for a, b in zip(cols, cols[1:]):
        if b - a > 1:
            runs.append((start, a))
            start = b
    runs.append((start, cols[-1]))
    return runs


def test_two_bars_are_capped_and_centred():
    runs = _bar_rects(["2008", "2024"], [186, 105])
    assert len(runs) == 2
    widths = [b - a for a, b in runs]
    assert all(w <= k.BAR_W_CAP + 2 for w in widths), widths
    # the group is centred in the content band
    left_pad, right_pad = runs[0][0] - k.M, k.X_RIGHT - runs[-1][1]
    assert abs(left_pad - right_pad) < 5


def test_seven_bars_span_the_content_band():
    runs = _bar_rects(["A", "B", "C", "D", "E", "F", "G"], [56, 49, 45, 43, 42, 30, 17], unit="%")
    assert len(runs) == 7
    assert runs[0][0] == pytest.approx(k.M, abs=3)
    assert runs[-1][1] == pytest.approx(k.X_RIGHT, abs=3)


def test_bar_heights_are_proportional_to_a_true_zero_baseline():
    canvas = k.gradient_bg()
    k.draw_bars(canvas, ["a", "b"], [100, 50])
    import numpy as np

    a = np.array(canvas.convert("RGB"), dtype=int)
    ground = np.array(k.gradient_bg().convert("RGB"), dtype=int)
    diff = np.abs(a - ground).sum(axis=2) > 30
    tall = np.where(diff[:, int(k.M + (k.X_RIGHT - k.M) * 0.30)])[0]
    short = np.where(diff[:, int(k.M + (k.X_RIGHT - k.M) * 0.70)])[0]
    # exclude the value label above each bar: measure from the shared baseline up
    h_tall = 1170 - tall.min()
    h_short = 1170 - short.min()
    assert h_short / h_tall == pytest.approx(0.5, abs=0.06)


def test_too_many_bars_is_rejected():
    with pytest.raises(SystemExit):
        k.draw_bars(k.gradient_bg(), list("abcdefgh"), [1] * 8)


def test_mismatched_bar_arrays_are_rejected():
    with pytest.raises(SystemExit):
        k.draw_bars(k.gradient_bg(), ["a", "b"], [1])


# ------------------------------------------------------------------ validation


def _spec(*slides):
    return {"meta": {"title": "t"}, "slides": list(slides)}


def test_valid_spec_has_no_errors(tmp_path):
    spec = _spec(
        {"type": "statement", "text": "x", "notes": "n"},
        {"type": "back"},
    )
    assert k.validate(spec, tmp_path) == []


def test_missing_speaker_notes_is_an_error(tmp_path):
    spec = _spec({"type": "statement", "text": "x"})
    errors = k.validate(spec, tmp_path)
    assert any("speaker notes" in e for e in errors)


def test_back_cover_needs_no_notes(tmp_path):
    assert k.validate(_spec({"type": "back"}), tmp_path) == []


def test_missing_required_key_is_an_error(tmp_path):
    errors = k.validate(_spec({"type": "hero-stat", "notes": "n"}), tmp_path)
    assert any("'value'" in e for e in errors)


def test_unknown_slide_type_is_an_error(tmp_path):
    errors = k.validate(_spec({"type": "carousel", "notes": "n"}), tmp_path)
    assert any("unknown type" in e for e in errors)


def test_third_chart_is_rejected(tmp_path):
    chart = {"type": "bars", "cats": ["a"], "vals": [1], "notes": "n"}
    errors = k.validate(_spec(chart, dict(chart), dict(chart)), tmp_path)
    assert any("at most 2" in e for e in errors)


def test_two_charts_are_allowed(tmp_path):
    chart = {"type": "bars", "cats": ["a"], "vals": [1], "notes": "n"}
    assert k.validate(_spec(chart, dict(chart)), tmp_path) == []


def test_unknown_scrim_is_an_error(tmp_path):
    errors = k.validate(
        _spec({"type": "statement", "text": "x", "notes": "n", "scrim": "diagonal"}), tmp_path
    )
    assert any("unknown scrim" in e for e in errors)


def test_missing_image_is_an_error(tmp_path):
    errors = k.validate(
        _spec({"type": "statement", "text": "x", "notes": "n", "image": "nope.png"}), tmp_path
    )
    assert any("image not found" in e for e in errors)


def test_shipped_example_yaml_is_valid():
    """Its photographs are the author's to supply, so validate it the way its
    header documents: with --placeholder-images. Layout is measured too."""
    spec = yaml.safe_load((HERE / "example.yaml").read_text(encoding="utf-8"))
    assert k.validate(spec, HERE, placeholder_images=True) == []


def test_missing_image_is_a_warning_with_placeholder_images(tmp_path, capsys):
    spec = _spec({"type": "statement", "text": "x", "notes": "n", "image": "nope.png"})
    assert k.validate(spec, tmp_path, placeholder_images=True) == []
    assert "image not found" in capsys.readouterr().err


# ------------------------------------------------------------------ palette


def test_dark_tokens_match_the_documented_hex():
    assert k.GROUND_BOT == (0, 56, 65)  # #003841, the brand Dark Teal
    assert k.ACCENT == (0, 223, 248)  # #00DFF8
    assert k.NEGATIVE == (255, 82, 99)  # #FF5263


def test_named_colour_lookup():
    assert k._color({"color": "negative"}, "color", k.INK) == k.NEGATIVE
    assert k._color({"color": "#00DFF8"}, "color", k.INK) == k.ACCENT
    assert k._color({}, "color", k.INK) == k.INK


def test_unknown_colour_is_rejected():
    with pytest.raises(SystemExit):
        k._color({"color": "chartreuse"}, "color", k.INK)


# ------------------------------------------------------------------ determinism


def test_grain_is_deterministic_per_seed():
    a = k.grain(k.gradient_bg(), 42)
    b = k.grain(k.gradient_bg(), 42)
    c = k.grain(k.gradient_bg(), 43)
    assert a.tobytes() == b.tobytes()
    assert a.tobytes() != c.tobytes()


# ------------------------------------------- specs that used to clear --validate

# README.md sells `--validate` as the gate before a render. These four specs
# passed it and then failed the build; cases 1 and 3 came out as raw tracebacks.


def test_whitespace_only_text_is_rejected(tmp_path):
    errors = k.validate(_spec({"type": "statement", "text": "   ", "notes": "n"}), tmp_path)
    assert any("blank" in e for e in errors), errors


def test_whitespace_only_text_can_no_longer_crash_the_renderer():
    """wrap() returns no lines, and max() over no lines used to raise."""
    canvas = k.gradient_bg()
    y = k.para(canvas, k.M, 400, "   ", k.font("l", 50), k.INK, 1200, 60)
    assert y == 400


def test_mismatched_cats_and_vals_is_rejected(tmp_path):
    spec = _spec({"type": "bars", "cats": ["a", "b", "c"], "vals": [1, 2], "notes": "n"})
    errors = k.validate(spec, tmp_path)
    assert any("same length" in e for e in errors), errors


def test_all_zero_bar_values_are_rejected(tmp_path):
    spec = _spec({"type": "bars", "cats": ["a", "b"], "vals": [0, 0], "notes": "n"})
    errors = k.validate(spec, tmp_path)
    assert any("positive value" in e for e in errors), errors


def test_all_zero_bar_values_can_no_longer_divide_by_zero():
    """v / mx with mx == 0 raised ZeroDivisionError mid-render."""
    canvas = k.gradient_bg()
    k.draw_bars(canvas, ["a", "b"], [0, 0])  # must not raise


def test_non_numeric_bar_values_are_rejected(tmp_path):
    spec = _spec({"type": "bars", "cats": ["a"], "vals": ["x"], "notes": "n"})
    errors = k.validate(spec, tmp_path)
    assert any("non-numeric" in e for e in errors), errors


def test_unknown_colour_is_rejected_by_validate(tmp_path):
    spec = _spec({"type": "hero-stat", "value": "42", "color": "chartreuse", "notes": "n"})
    errors = k.validate(spec, tmp_path)
    assert any("chartreuse" in e for e in errors), errors


def test_unknown_colour_in_a_duo_stat_half_is_rejected(tmp_path):
    spec = _spec(
        {
            "type": "duo-stat",
            "left": {"value": "1", "color": "puce"},
            "right": {"value": "2"},
            "notes": "n",
        }
    )
    errors = k.validate(spec, tmp_path)
    assert any("puce" in e and "left.color" in e for e in errors), errors


def test_is_color_rejects_six_non_hex_characters():
    assert k.is_color("#00DFF8")
    assert k.is_color("accent")
    assert not k.is_color("zzzzzz")
    assert not k.is_color("chartreuse")
    assert not k.is_color(42)


def test_blank_point_is_rejected(tmp_path):
    spec = _spec({"type": "points", "points": ["real", "  "], "notes": "n"})
    errors = k.validate(spec, tmp_path)
    assert any("point 2 is blank" in e for e in errors), errors


def test_negative_bar_value_is_rejected(tmp_path):
    """One negative value used to clear --validate and then kill the build in
    Pillow ("y1 must be greater than or equal to y0")."""
    spec = _spec(
        {"type": "bars", "cats": list("abcd"), "vals": [6.2, 3.1, -1.4, 2.0], "notes": "n"}
    )
    errors = k.validate(spec, tmp_path)
    assert any("negative" in e for e in errors), errors


def test_negative_bar_value_never_reaches_pillow():
    with pytest.raises(SystemExit):
        k.draw_bars(k.gradient_bg(), list("abcd"), [6.2, 3.1, -1.4, 2.0])


# ------------------------------------------------------------------ fonts


def _real_ttf() -> Path:
    """Any loadable TrueType file on this machine, found with the default search.

    The lookup tests below exercise file names and directories only, so any real
    font stands in for Aptos or Calibri. The renderer tests above need one too.
    """
    for names in k.FONT_FILES.values():
        for name in names:
            p = k._find_font_file(name)
            if p is not None:
                return p
    raise AssertionError("no usable font on this machine")


def _use_font_dirs(monkeypatch, dirs):
    monkeypatch.setattr(k, "FONT_DIRS", list(dirs))
    monkeypatch.setattr(k, "_font_cache", {})
    monkeypatch.setattr(k, "_font_warned", set())
    monkeypatch.setattr(k, "_font_path_cache", {}, raising=False)


def _stock(tmp_path, names, sub="DFonts"):
    src = _real_ttf()
    d = tmp_path / sub
    d.mkdir(parents=True)
    for name in names:
        shutil.copyfile(src, d / name)
    return d


def test_default_font_dirs_cover_office_bundles_and_windows(monkeypatch):
    """Office for Mac keeps Aptos and Calibri inside each app bundle, and Windows
    keeps them in two font directories. None of them used to be searched, so a
    stock colleague machine aborted before slide 1."""
    monkeypatch.setenv("WINDIR", "C:\\Windows")
    monkeypatch.setenv("LOCALAPPDATA", "C:\\Users\\a\\AppData\\Local")
    dirs = [str(d).replace("\\", "/") for d in k.default_font_dirs()]
    for app in ("PowerPoint", "Word", "Excel"):
        assert f"/Applications/Microsoft {app}.app/Contents/Resources/DFonts" in dirs
    assert "C:/Windows/Fonts" in dirs
    assert "C:/Users/a/AppData/Local/Microsoft/Windows/Fonts" in dirs


def test_calibri_fallback_uses_the_file_names_office_ships(tmp_path, monkeypatch):
    """Office ships Calibri Bold as Calibrib.ttf; the lookup asked for 'Calibri Bold.ttf'."""
    d = _stock(tmp_path, ["Calibri.ttf", "Calibrib.ttf"])
    _use_font_dirs(monkeypatch, [d])
    assert Path(k.font("x", 20).path).name == "Calibrib.ttf"
    assert Path(k.font("r", 20).path).name == "Calibri.ttf"


def test_font_file_names_match_case_insensitively(tmp_path, monkeypatch):
    """Windows spells them calibrib.ttf; an installer may shout APTOS-LIGHT.TTF."""
    d = _stock(tmp_path, ["calibrib.ttf", "APTOS-LIGHT.TTF"], sub="Fonts")
    _use_font_dirs(monkeypatch, [d])
    assert Path(k.font("x", 20).path).name == "calibrib.ttf"
    assert Path(k.font("l", 20).path).name == "APTOS-LIGHT.TTF"


def test_office_bundled_aptos_wins_over_calibri(tmp_path, monkeypatch):
    names = ["Aptos.ttf", "Aptos-Light.ttf", "Aptos-SemiBold.ttf", "Aptos-ExtraBold.ttf"]
    d = _stock(tmp_path, [*names, "Calibri.ttf", "Calibrib.ttf"])
    _use_font_dirs(monkeypatch, [d])
    assert Path(k.font("x", 20).path).name == "Aptos-ExtraBold.ttf"
    assert Path(k.font("s", 20).path).name == "Aptos-SemiBold.ttf"
    assert Path(k.font("l", 20).path).name == "Aptos-Light.ttf"
    assert Path(k.font("r", 20).path).name == "Aptos.ttf"


def test_missing_fonts_error_says_where_it_looked(tmp_path, monkeypatch):
    _use_font_dirs(monkeypatch, [tmp_path])
    with pytest.raises(SystemExit) as e:
        k.font("x", 20)
    assert str(tmp_path) in str(e.value)


# ----------------------------------------------------------- contrast coverage

# Every text element must pass through the contrast guard. The guard used to
# skip the kicker, the cover title, hero and duo values, bar labels and the
# source note, which are exactly the elements a bright photograph breaks.

_FULL_SLIDES: dict[str, dict] = {
    "cover": {
        "title": "The headline",
        "subtitle": "What it is about",
        "speaker": {"name": "Speaker Name", "role": "Role"},
        "venue": "Venue",
    },
    "statement": {
        "kicker": "The consensus",
        "text": "One sentence.",
        "emphasis": "And more.",
        "coda": "The turn.",
        "source": "Source: publication, 2026.",
    },
    "hero-stat": {
        "kicker": "The evidence",
        "value": "68m",
        "caption": "what it counts",
        "support": "The detail.",
        "source": "Source: publication, 2026.",
    },
    "duo-stat": {
        "kicker": "The moat",
        "title": "Two numbers",
        "left": {"value": "19%", "caption": "a"},
        "right": {"value": "5%", "caption": "b"},
        "source": "Source: publication, 2026.",
    },
    "bars": {
        "kicker": "The retreat",
        "title": "A comparison",
        "cats": ["Austria", "Italy", "Spain"],
        "vals": [56, 49, 45],
        "highlight": 1,
        "unit": "%",
        "axis_label": "Share of customers",
        "source": "Source: publication, 2026.",
    },
    "points": {
        "kicker": "The public good",
        "title": "Three things",
        "points": ["One.", "Two.", "Three."],
        "takeaway": "The sentence.",
        "source": "Source: publication, 2026.",
    },
    "divider": {"kicker": "The playbook", "title": "Three moves"},
    "closing": {"kicker": "The talk", "text": "The line to repeat.", "coda": "The short version."},
}

_MUST_BE_GUARDED = {
    "cover": {"cover title", "cover subtitle", "speaker name", "speaker role", "venue"},
    "statement": {"kicker", "statement", "emphasis", "coda", "source"},
    "hero-stat": {"kicker", "hero value", "stat caption", "stat support", "source"},
    "duo-stat": {"kicker", "duo title", "duo value", "duo caption", "source"},
    "bars": {"kicker", "bars title", "bar value", "bar label", "axis label", "source"},
    "points": {"kicker", "points title", "point", "takeaway", "source"},
    "divider": {"kicker", "divider title"},
    "closing": {"kicker", "closing", "closing coda"},
}


@pytest.mark.parametrize("kind", sorted(_FULL_SLIDES))
def test_every_text_element_clears_contrast_over_a_bright_photo(tmp_path, monkeypatch, kind):
    Image.new("RGB", (1600, 900), (226, 233, 237)).save(tmp_path / "bright.jpg")
    slide = {"type": kind, "image": "bright.jpg", "scrim": "right", **_FULL_SLIDES[kind]}
    calls = []
    real = k.ensure_contrast

    def spy(canvas, box, fg, min_ratio, label=""):
        real(canvas, box, fg, min_ratio, label)
        calls.append((label, k.contrast_ratio(fg, k._bg_luminance(canvas, box)), min_ratio))

    monkeypatch.setattr(k, "ensure_contrast", spy)
    ground, on_image = k._ground(slide, tmp_path)
    ctx = {"logo": k.white_wordmark(LOGO), "greek": False}
    k.RENDERERS[kind](ground, slide, ctx, on_image)

    guarded = {label for label, _, _ in calls}
    missing = _MUST_BE_GUARDED[kind] - guarded
    assert not missing, f"drawn without the contrast guard: {sorted(missing)}"
    failing = [(label, round(r, 2), need) for label, r, need in calls if r < need]
    assert not failing, failing


# -------------------------------------------------------------------- text fit


def test_cover_title_over_the_line_budget_is_rejected(tmp_path):
    title = (
        "Payments at the speed of trust, and what it takes to get there, "
        "told in far more words than any cover can carry"
    )
    errors = k.validate(_spec({"type": "cover", "title": title, "notes": "n"}), tmp_path)
    assert any("cover title" in e and "lines" in e for e in errors), errors


def test_two_line_cover_title_wraps_inside_the_gutters(tmp_path):
    """It used to be drawn as one unwrapped line and cut off at the canvas edge."""
    slide = {
        "type": "cover",
        "title": "Payments at the speed of trust, and what it takes to get there",
        "subtitle": "The subordinate clause",
        "notes": "n",
    }
    assert k.validate(_spec(slide), tmp_path) == []
    boxes = k.text_boxes(slide)
    title = [b for label, b in boxes if label == "cover title"]
    assert len(title) == 2
    assert max(b[2] for _, b in boxes) <= k.X_RIGHT + 4
    subtitle = [b for label, b in boxes if label == "cover subtitle"]
    assert min(b[1] for b in subtitle) > max(b[3] for b in title)


def test_hero_value_past_the_right_gutter_is_rejected(tmp_path):
    spec = _spec(
        {"type": "hero-stat", "value": "€4.8bn", "align": "right", "caption": "c", "notes": "n"}
    )
    errors = k.validate(spec, tmp_path)
    assert any("hero value" in e and "gutter" in e for e in errors), errors


def test_divider_title_wraps_inside_the_right_gutter():
    slide = {
        "type": "divider",
        "title": "From here to there, in three deliberate and measurable moves",
        "notes": "n",
    }
    boxes = k.text_boxes(slide)
    assert max(b[2] for label, b in boxes if label == "divider title") <= k.X_RIGHT + 4


def test_text_running_into_the_footer_is_rejected(tmp_path):
    spec = _spec(
        {
            "type": "points",
            "title": "t",
            "notes": "n",
            "points": [f"Point number {i} carries one line." for i in range(1, 16)],
        }
    )
    errors = k.validate(spec, tmp_path)
    assert any("footer" in e for e in errors), errors


def test_overlapping_text_blocks_are_rejected(tmp_path):
    spec = _spec(
        {
            "type": "statement",
            "notes": "n",
            "coda": "The turn.",
            "coda_y": 520,
            "text": "A statement long enough to wrap onto a second and a third line of large type.",
        }
    )
    errors = k.validate(spec, tmp_path)
    assert any("overlaps" in e for e in errors), errors


# ------------------------------------------------------------------ bar widths


@pytest.mark.parametrize("n", range(1, 8))
def test_bar_width_cap_holds_for_every_bar_count(n):
    """The 380px cap applied only to three bars or fewer, so four bars came out
    528px wide, wider than three: the 'wall' the rule forbids."""
    runs = _bar_rects([f"c{i}" for i in range(n)], [10 + i for i in range(n)])
    assert len(runs) == n
    widths = [b - a for a, b in runs]
    assert max(widths) <= k.BAR_W_CAP + 2, widths
    left_pad, right_pad = runs[0][0] - k.M, k.X_RIGHT - runs[-1][1]
    assert abs(left_pad - right_pad) < 5


def _pixel(canvas, x, y):
    return canvas.convert("RGB").getpixel((int(x), int(y)))


def test_bars_without_a_highlight_use_the_default_bar_fill():
    canvas = k.gradient_bg()
    k.draw_bars(canvas, list("abcde"), [5, 5, 5, 5, 5])
    runs = _bar_rects(list("abcde"), [5, 5, 5, 5, 5])
    for a, b in runs:
        assert _pixel(canvas, (a + b) / 2, 1100) == k.ACCENT_BAR


def test_a_highlight_mutes_the_other_bars():
    canvas = k.gradient_bg()
    k.draw_bars(canvas, list("abc"), [5, 5, 5], highlight=1)
    runs = _bar_rects(list("abc"), [5, 5, 5], highlight=1)
    colours = [_pixel(canvas, (a + b) / 2, 1100) for a, b in runs]
    assert colours == [k.ACCENT_MUTE, k.ACCENT, k.ACCENT_MUTE]


# ------------------------------------------------------------------ PDF, PPTX


def _frames(tmp_path, n=2):
    out = []
    for i in range(n):
        p = tmp_path / f"frame_{i}.jpg"
        Image.new("RGB", (320, 180), (0, 56 + 40 * i, 65)).save(p, quality=92, subsampling=0)
        out.append(p)
    return out


def _first_pdf_jpeg(pdf: Path) -> Image.Image:
    data = pdf.read_bytes()
    # Image.open parses only the headers, so everything from the SOI marker on is enough.
    return Image.open(io.BytesIO(data[data.index(b"\xff\xd8\xff") :]))


def test_pdf_keeps_4_4_4_chroma(tmp_path):
    """Hard rule 7. The PDF is the file that goes on the AV laptop, and it was
    re-encoded at 4:2:0 while the frames were 4:4:4."""
    out = tmp_path / "deck.pdf"
    k.build_pdf(_frames(tmp_path), out)
    assert JpegImagePlugin.get_sampling(_first_pdf_jpeg(out)) == 0  # 0 = 4:4:4, 2 = 4:2:0


def _picture_descr(slide):
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
    return pics[0]._element.nvPicPr.cNvPr.get("descr", "")


def test_pptx_slides_carry_a_title_and_alt_text(tmp_path):
    """A screen reader used to announce every slide as 'slide_01.jpg'."""
    slides = [
        {
            "type": "statement",
            "kicker": "The consensus",
            "text": "One sentence.",
            "source": "Source: publication, 2026.",
            "notes": "n1",
        },
        {
            "type": "hero-stat",
            "value": "68m",
            "caption": "customers",
            "support": "The detail.",
            "notes": "n2",
        },
    ]
    out = tmp_path / "deck.pptx"
    k.build_pptx(_frames(tmp_path), [s["notes"] for s in slides], out, slides=slides)
    s1, s2 = Presentation(str(out)).slides
    assert s1.shapes.title.text == "One sentence."
    assert s2.shapes.title.text == "68m customers"
    d1, d2 = _picture_descr(s1), _picture_descr(s2)
    for part in ("The consensus", "One sentence.", "Source: publication, 2026."):
        assert part in d1
    for part in ("68m", "customers", "The detail."):
        assert part in d2
    assert ".jpg" not in d1 + d2


def test_bars_alt_text_reads_out_the_values():
    s = {
        "type": "bars",
        "title": "Share",
        "cats": ["Austria", "Italy"],
        "vals": [56, 49],
        "unit": "%",
    }
    text = k.slide_text(s)
    assert "Austria 56%" in text and "Italy 49%" in text


# ------------------------------------------------------------ end-to-end builds


def _build(tmp_path, monkeypatch, slides, *extra):
    spec = tmp_path / "talk.yaml"
    spec.write_text(
        yaml.safe_dump({"meta": {"title": "t"}, "slides": slides}, allow_unicode=True),
        encoding="utf-8",
    )
    argv = ["nbg_keynote.py", str(spec), "--out", str(tmp_path / "talk"), *extra]
    monkeypatch.setattr(sys, "argv", argv)
    assert k.main() == 0
    return tmp_path / "talk_frames"


def _stmt(name):
    return {"type": "statement", "text": f"Slide {name}", "notes": f"note {name}"}


def test_slides_pairs_frames_with_their_own_slide_after_an_insert(tmp_path, monkeypatch):
    """--slides reused frames by position: after inserting a slide it paired the
    stale image of C with the notes of B, and B disappeared."""
    frames = _build(tmp_path, monkeypatch, [_stmt("A"), _stmt("B"), _stmt("C")])
    v1 = {n: (frames / f"slide_{i:02d}.jpg").read_bytes() for i, n in enumerate("ABC", 1)}

    frames = _build(
        tmp_path, monkeypatch, [_stmt("A"), _stmt("X"), _stmt("B"), _stmt("C")], "--slides", "2"
    )
    assert (frames / "slide_01.jpg").read_bytes() == v1["A"]
    assert (frames / "slide_03.jpg").read_bytes() == v1["B"]
    assert (frames / "slide_04.jpg").read_bytes() == v1["C"]
    prs = Presentation(str(tmp_path / "talk.pptx"))
    assert [s.notes_slide.notes_text_frame.text for s in prs.slides] == [
        "note A",
        "note X",
        "note B",
        "note C",
    ]


def test_slides_never_reuses_the_frame_of_an_edited_slide(tmp_path, monkeypatch):
    frames = _build(tmp_path, monkeypatch, [_stmt("A"), _stmt("B")])
    before = (frames / "slide_02.jpg").read_bytes()
    edited = {**_stmt("B"), "text": "Slide B, edited"}
    frames = _build(tmp_path, monkeypatch, [_stmt("A"), edited], "--slides", "1")
    assert (frames / "slide_02.jpg").read_bytes() != before


def test_a_shorter_deck_leaves_no_stale_frame_behind(tmp_path, monkeypatch):
    _build(tmp_path, monkeypatch, [_stmt("A"), _stmt("B"), _stmt("C")])
    frames = _build(tmp_path, monkeypatch, [_stmt("A"), _stmt("B")])
    assert sorted(p.name for p in frames.glob("slide_*.jpg")) == ["slide_01.jpg", "slide_02.jpg"]


def test_shipped_example_builds_end_to_end(tmp_path, monkeypatch):
    """example.yaml could not pass its own first documented command without five
    photographs nobody has. --placeholder-images lays it out before they arrive."""
    spec = yaml.safe_load((HERE / "example.yaml").read_text(encoding="utf-8"))
    n = len(spec["slides"])
    out = tmp_path / "example"
    argv = ["nbg_keynote.py", str(HERE / "example.yaml"), "--out", str(out), "--placeholder-images"]
    monkeypatch.setattr(sys, "argv", argv)
    assert k.main() == 0

    assert len(list((tmp_path / "example_frames").glob("slide_*.jpg"))) == n
    prs = Presentation(str(out.with_suffix(".pptx")))
    assert len(prs.slides) == n
    for slide, s in zip(prs.slides, spec["slides"], strict=True):
        assert ".jpg" not in _picture_descr(slide)
        assert slide.shapes.title is not None and slide.shapes.title.text.strip()
        if s["type"] != "back":
            assert slide.notes_slide.notes_text_frame.text.strip()
    pdf = out.with_suffix(".pdf")
    assert len(re.findall(rb"/Type\s*/Page(?![a-z])", pdf.read_bytes())) == n
    assert JpegImagePlugin.get_sampling(_first_pdf_jpeg(pdf)) == 0


# ------------------------------------------------------------------ tokens


def test_dark_palette_and_gutter_are_read_from_tokens_yaml(monkeypatch):
    """shared/brand-system/tokens.yaml is the machine source for the dark tokens."""
    import nbg_tokens

    fake = copy.deepcopy(nbg_tokens.load())
    fake["keynote"]["dark"]["accent"] = "123456"
    fake["keynote"]["gutter_px"] = 160
    monkeypatch.setattr(nbg_tokens, "load", lambda: fake)
    try:
        importlib.reload(k)
        assert k.ACCENT == (0x12, 0x34, 0x56)
        assert k.M == 160
    finally:
        monkeypatch.undo()
        importlib.reload(k)
    assert k.M == 155


def test_dark_constants_equal_the_tokens():
    import nbg_tokens

    dark = nbg_tokens.get("keynote.dark")
    pairs = {
        "ink": k.INK,
        "ink_2": k.INK_2,
        "ink_3": k.INK_3,
        "accent": k.ACCENT,
        "accent_bar": k.ACCENT_BAR,
        "accent_mute": k.ACCENT_MUTE,
        "ground_from": k.GROUND_TOP,
        "ground_to": k.GROUND_BOT,
        "negative": k.NEGATIVE,
        "rule": k.RULE,
    }
    for name, rgb in pairs.items():
        assert "".join(f"{c:02X}" for c in rgb) == dark[name].upper(), name


# ------------------------------------------------------------------ CLI


def _write_spec(tmp_path, *slides):
    p = tmp_path / "talk.yaml"
    p.write_text(yaml.safe_dump(_spec(*slides), allow_unicode=True), encoding="utf-8")
    return p


def test_validate_flag_names_the_font_behind_each_weight(tmp_path, monkeypatch, capsys):
    """A fallback font must be visible before a render, not discovered on stage."""
    spec = _write_spec(tmp_path, {"type": "statement", "text": "x", "notes": "n"})
    monkeypatch.setattr(sys, "argv", ["nbg_keynote.py", str(spec), "--validate"])
    assert k.main() == 0
    out = capsys.readouterr().out
    for weight in ("ExtraBold", "SemiBold", "Light", "Regular"):
        assert weight in out


def test_slides_index_past_the_end_is_an_error(tmp_path, monkeypatch, capsys):
    spec = _write_spec(tmp_path, {"type": "statement", "text": "x", "notes": "n"})
    argv = ["nbg_keynote.py", str(spec), "--out", str(tmp_path / "t"), "--slides", "4"]
    monkeypatch.setattr(sys, "argv", argv)
    assert k.main() == 1
    assert "no slide 4" in capsys.readouterr().err
